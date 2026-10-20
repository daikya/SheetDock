"""Comparison CLI for the validated XLSX row/column operations."""
from datetime import datetime
from pathlib import Path
import argparse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from sheetdock.xlsx_gui import XlsxSession
from sheetdock.xlsx_io import snapshot

from sheetdock.xlsx_axis import ACTIONS, shifted

AxisSession = XlsxSession


def make_samples(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    targets = [folder / f'axis_{name}.xlsx' for name in ('basic', 'formula', 'merged', 'readonly')]
    if any(path.exists() for path in targets):
        raise FileExistsError('同名サンプルが存在します。既存ファイルは上書きしません')
    for name, target in zip(('basic', 'formula', 'merged', 'readonly'), targets):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Data'
        for row in (['コード', '数量', '日付', '確認'], ['00123', 3, datetime(2026, 10, 16), True],
                    ['00456', 8, datetime(2026, 10, 17), False], ['00789', 12, datetime(2026, 10, 18), True]):
            sheet.append(row)
        sheet['A2'].font = Font(bold=True, color='FF2563EB')
        sheet['A2'].fill = PatternFill('solid', fgColor='FFFDE68A')
        sheet.column_dimensions['A'].width = 24
        sheet.column_dimensions['B'].width = 14
        sheet.column_dimensions['C'].width = 22
        sheet.row_dimensions[2].height = 32
        notes = book.create_sheet('Notes')
        notes['A1'] = '他シートは変更しない'
        if name == 'formula':
            notes['A2'] = '=Data!B2'
        elif name == 'merged':
            sheet.merge_cells('A1:B1')
        elif name == 'readonly':
            rule = DataValidation(type='list', formula1='"はい,いいえ"')
            sheet.add_data_validation(rule)
            rule.add('D2')
        with target.open('xb') as handle:
            book.save(handle)
        book.close()
        print(target)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path', nargs='?')
    parser.add_argument('--make-samples', type=Path)
    parser.add_argument('--sheet')
    parser.add_argument('--action', choices=ACTIONS)
    parser.add_argument('--index', type=int)
    parser.add_argument('--count', type=int, default=1)
    parser.add_argument('--save-as', type=Path)
    args = parser.parse_args(argv)
    if args.make_samples:
        if args.path or args.action or args.save_as:
            parser.error('サンプル生成と操作は別々に実行してください')
    elif not args.path or not args.action or args.index is None or not args.save_as:
        parser.error('path、--action、--index、--save-asを指定してください')
    try:
        if args.make_samples:
            make_samples(args.make_samples)
            return 0
        session = AxisSession(args.path)
        try:
            if args.sheet:
                if args.sheet not in session.probe.book.sheetnames:
                    raise ValueError('指定したシートが存在しません')
                session.sheet_name = args.sheet
            before = snapshot(session.probe.book)
            changed = session.axis_operation(args.action, args.index, args.count)
            if not changed:
                print('保存対象の変更はありません。出力ファイルは作成しません。')
                return 0
            after = snapshot(session.probe.book)
            session.history()
            if snapshot(session.probe.book) != before:
                raise ValueError('Undoの検証に失敗しました')
            session.history(True)
            if snapshot(session.probe.book) != after:
                raise ValueError('Redoの検証に失敗しました')
            session.save_as(args.save_as)
            print(f'行列操作・Undo/Redo・保存再読込の検証完了：{args.save_as}')
            print('元ファイルは変更していません。本体GUIの履歴とは独立したCLI実行です。')
            return 0
        finally:
            session.probe.book.close()
    except Exception as exc:
        print(f'操作中止：{exc}')
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
