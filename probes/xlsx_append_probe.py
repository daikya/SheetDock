"""Comparison CLI for production trailing blank row/column operations."""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from sheetdock.xlsx_gui import XlsxSession
from sheetdock.xlsx_io import snapshot
from sheetdock.xlsx_extent import extent, extents

AppendSession = XlsxSession


def make_samples(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    paths = [folder / f'append_{name}.xlsx' for name in ('basic', 'empty', 'formula', 'merged', 'readonly')]
    if any(path.exists() for path in paths):
        raise FileExistsError('同名サンプルが存在します。別フォルダーを指定してください')
    for name, path in zip(('basic', 'empty', 'formula', 'merged', 'readonly'), paths):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Data'
        if name != 'empty':
            sheet.append(['コード', '数量'])
            sheet.append(['00123', 3])
            sheet.row_dimensions[2].height = 32
            sheet.column_dimensions['A'].width = 24
        book.create_sheet('Other')['A1'] = '他シート不変'
        if name == 'formula':
            book['Other']['A2'] = '=Data!B2'
        elif name == 'merged':
            sheet.merge_cells('A1:B1')
        elif name == 'readonly':
            rule = DataValidation(type='list', formula1='"はい,いいえ"')
            sheet.add_data_validation(rule)
            rule.add('B2')
        try:
            with path.open('xb') as stream:
                book.save(stream)
        finally:
            book.close()
        print(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path', nargs='?', type=Path)
    parser.add_argument('--make-samples', type=Path)
    parser.add_argument('--sheet')
    parser.add_argument('--axis', choices=['rows', 'columns'])
    parser.add_argument('--count', type=int, default=1)
    parser.add_argument('--save-as', type=Path)
    args = parser.parse_args(argv)
    if args.make_samples:
        if args.path or args.axis or args.save_as or args.sheet or args.count != 1:
            parser.error('サンプル生成と操作は別々に実行してください')
    elif not args.path or not args.axis or not args.save_as:
        parser.error('path、--axis、--save-asを指定してください')
    try:
        if args.make_samples:
            make_samples(args.make_samples)
            return 0
        # Check this before loading or changing a workbook. save_as rechecks atomically.
        if args.save_as.exists():
            raise FileExistsError('出力先が既に存在します。別の名前を指定してください')
        session = AppendSession(args.path)
        try:
            if args.sheet:
                if args.sheet not in session.probe.book.sheetnames:
                    raise ValueError('指定したシートが存在しません')
                session.sheet_name = args.sheet
            source = args.path.read_bytes()
            before, before_extent = snapshot(session.probe.book), extents(session.probe.book)
            session.append_axis(args.axis, args.count)
            after, after_extent = snapshot(session.probe.book), extents(session.probe.book)
            session.history()
            if snapshot(session.probe.book) != before or extents(session.probe.book) != before_extent or session.dirty:
                raise ValueError('Undoの検証に失敗しました')
            session.history(True)
            if snapshot(session.probe.book) != after or extents(session.probe.book) != after_extent:
                raise ValueError('Redoの検証に失敗しました')
            session.save_as(args.save_as)
            reopened = AppendSession(args.save_as)
            try:
                if snapshot(reopened.probe.book) != after or extents(reopened.probe.book) != after_extent:
                    raise ValueError('別名保存・再読込の検証に失敗しました')
                print('表示範囲（行×列）:', extent(reopened.probe.book[session.sheet_name]))
                print('openpyxlのセル範囲（行×列）:', (reopened.probe.book[session.sheet_name].max_row,
                      reopened.probe.book[session.sheet_name].max_column))
            finally:
                reopened.probe.book.close()
            if args.path.read_bytes() != source:
                raise ValueError('検証中に元ファイルが変更されました')
            print(f'末尾追加・Undo/Redo・保存再読込の検証完了：{args.save_as}')
            print('元ファイルは変更していません。本体GUIと同じ追加処理を使う独立CLIです。')
            return 0
        finally:
            session.probe.book.close()
    except Exception as exc:
        print(f'操作中止：{exc}')
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
