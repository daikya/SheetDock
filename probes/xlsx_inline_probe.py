"""Comparison GUI using the production XLSX inline-edit validator."""
from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from tksheet import Sheet

from sheetdock import dialogs
from sheetdock.xlsx_gui import XlsxSession, protection_reason, cell_state, display_value
from sheetdock.xlsx_display import format_value
from sheetdock.xlsx_extent import extent
from sheetdock.xlsx_axis import VIEW_ROWS, VIEW_COLUMNS


from sheetdock.xlsx_inline import InlineBridge


class InlineProbe:
    def __init__(self, root):
        self.root, self.session, self._flush_id = root, None, None
        root.title('SheetDock — XLSX直接編集 比較プローブ')
        root.geometry('1050x650')
        root.columnconfigure(0, weight=1)
        root.rowconfigure(2, weight=1)
        bar = ttk.Frame(root)
        bar.grid(row=0, column=0, sticky='ew', padx=8, pady=6)
        ttk.Button(bar, text='開く', command=self.open).pack(side='left')
        self.save_button = ttk.Button(bar, text='別名保存', command=self.save_as, state='disabled')
        self.save_button.pack(side='left', padx=4)
        self.undo_button = ttk.Button(bar, text='Undo', command=self.history, state='disabled')
        self.undo_button.pack(side='left')
        self.redo_button = ttk.Button(bar, text='Redo', command=lambda: self.history(True), state='disabled')
        self.redo_button.pack(side='left', padx=4)
        self.sheet_name = tk.StringVar()
        self.sheets = ttk.Combobox(bar, textvariable=self.sheet_name, state='readonly', width=18)
        self.sheets.pack(side='left', padx=8)
        self.sheets.bind('<<ComboboxSelected>>', self.switch_sheet)
        ttk.Label(root, text='ダブルクリック／F2で編集。Enter・Tab・別セルクリックで確定、Escで取消。\n'
                  '入力欄は元の値、表は表示書式。Ctrl+Z / Ctrl+Shift+Zは入力中は文字、確定後はブックのUndo/Redo。',
                  wraplength=1000).grid(row=1, column=0, sticky='ew', padx=8)
        self.sheet = Sheet(root, data=[['']], headers=['A'], cell_auto_resize_enabled=False)
        self.sheet.grid(row=2, column=0, sticky='nsew', padx=8, pady=4)
        self.sheet.enable_bindings('single_select', 'drag_select', 'row_select', 'column_select',
                                   'arrowkeys', 'copy', 'select_all', 'edit_cell')
        self.status = tk.StringVar(value='検証用XLSXを開いてください。保存は未使用名への別名保存のみです。')
        ttk.Label(root, textvariable=self.status, wraplength=1000).grid(row=3, column=0, sticky='ew', padx=8, pady=6)
        self.detail = tk.StringVar()
        ttk.Label(root, textvariable=self.detail, wraplength=1000).grid(row=4, column=0, sticky='ew', padx=8, pady=3)
        self.sheet.bind('<<SheetSelect>>', self.selection_detail)
        self.bridge = InlineBridge(None, self.queue_refresh, self.status.set)
        self.sheet.extra_bindings('begin_edit_cell', self.bridge.begin)
        self.sheet.edit_validation(self.bridge.validate)
        self.sheet.bind_key_text_editor('<Escape>', self.cancel_edit)
        self.sheet.bind_key_text_editor('<FocusOut>', self.focus_out)
        for key, redo in [('<Control-z>', False), ('<Control-Z>', False),
                          ('<Control-Shift-z>', True), ('<Control-Shift-Z>', True)]:
            self.sheet.bind_key_text_editor(key, lambda event, redo=redo: self.text_history(event, redo))
            root.bind(key, lambda event, redo=redo: self.history_key(event, redo))
        root.protocol('WM_DELETE_WINDOW', self.close)

    def cancel_edit(self, event=None):
        self.bridge.cancel()
        self.sheet.close_text_editor(set_data=False)
        return 'break'

    def focus_out(self, event):
        # The same Tk Text is reused. An older queued FocusOut must not close
        # a newly opened editor that already owns focus again.
        if event.widget.focus_get() is event.widget:
            return 'break'
        event.keysym = 'FocusOut'
        self.sheet.MT.close_text_editor(event)
        return 'break'

    def text_history(self, event, redo=False):
        try:
            event.widget.edit_redo() if redo else event.widget.edit_undo()
        except tk.TclError:
            pass
        return 'break'

    def selection_detail(self, event=None):
        selected = self.sheet.get_currently_selected()
        if self.session and selected:
            cell = self.session.sheet.cell(selected.row + 1, selected.column + 1)
            reason = protection_reason(self.session.probe, self.session.sheet, cell)
            self.detail.set(f'{cell.coordinate} / 型={cell.data_type} / {reason or "編集可能"}')
        else:
            self.detail.set('')

    def history_key(self, event, redo=False):
        if self.sheet.MT.text_editor.open:
            return 'break'
        self.history(redo)
        return 'break'

    def queue_refresh(self):
        if self._flush_id is None:
            self._flush_id = self.root.after_idle(self.flush)

    def flush(self):
        if self._flush_id is not None:
            self.root.after_cancel(self._flush_id)
            self._flush_id = None
        error, self.bridge.pending_error = self.bridge.pending_error, None
        # Validation already supplied the formatted display value to tksheet.
        # Avoid resetting a newer editor opened before this idle callback runs.
        self.refresh_state()
        if error:
            # The native editor has closed; avoid reentrant FocusOut validation.
            self.status.set('入力取消：' + error)
            dialogs.showerror('直接編集の入力エラー', error, parent=self.root)

    def settle(self):
        was_open = self.sheet.MT.text_editor.open
        if was_open:
            self.sheet.MT.text_editor.tktext.focus_set()
            self.sheet.close_text_editor()
        if self.sheet.MT.text_editor.open:
            self.status.set('先に入力を確定または取消してください')
            return False
        if self._flush_id is not None:
            self.flush()
        return not (was_open and self.bridge.last_error)

    def confirm_discard(self):
        return not self.session or not self.session.dirty or dialogs.askyesno(
            '未保存の変更', '編集内容を破棄して続けますか？', parent=self.root)

    def open(self):
        if not self.settle():
            return
        path = filedialog.askopenfilename(parent=self.root, filetypes=[('Excel workbook', '*.xlsx')])
        if path:
            self.open_path(path)

    def open_path(self, path):
        if not self.settle() or not self.confirm_discard():
            return
        try:
            session = XlsxSession(path)
            if self.session:
                self.session.probe.book.close()
            self.session, self.bridge.session = session, session
            self.bridge.cancel()
            self.sheets.configure(values=session.probe.book.sheetnames)
            self.render()
        except Exception as exc:
            dialogs.showerror('読込失敗', str(exc), parent=self.root)

    def switch_sheet(self, event=None):
        name = self.sheet_name.get()
        if not self.session:
            return
        if self.settle() and name in self.session.probe.book.sheetnames:
            self.session.sheet_name = name
            self.bridge.context = None
            self.render()
        else:
            self.sheet_name.set(self.session.sheet_name)

    def render(self):
        if not self.session:
            return
        selected = self.sheet.get_currently_selected()
        sheet = self.session.sheet
        rows, columns = extent(sheet)
        rows, columns = min(rows, VIEW_ROWS), min(columns, VIEW_COLUMNS)
        self.sheet.set_sheet_data([[format_value(sheet.cell(r, c).value, sheet.cell(r, c).number_format)[0]
            for c in range(1, columns + 1)] for r in range(1, rows + 1)], delete_options=True)
        self.sheet.headers([get_column_letter(c) for c in range(1, columns + 1)])
        self.sheet.set_column_widths(self.session.column_widths())
        self.sheet.set_row_heights(self.session.row_heights(float(self.root.winfo_fpixels('1p')),
                                                         self.sheet.MT.min_row_height))
        for r in range(rows):
            for c in range(columns):
                if protection_reason(self.session.probe, sheet, sheet.cell(r + 1, c + 1)):
                    self.sheet.highlight_cells(r, c, bg='#e5e7eb', fg='#374151')
        row, column = (selected.row, selected.column) if selected else (0, 0)
        self.sheet.select_cell(min(row, rows - 1), min(column, columns - 1))
        self.refresh_state()

    def refresh_state(self):
        if not self.session:
            return
        sheet = self.session.sheet
        self.sheet_name.set(sheet.title)
        editable = not self.session.probe.read_only
        self.save_button.configure(state='normal' if editable else 'disabled')
        self.undo_button.configure(state='normal' if editable and self.session.undo_stack else 'disabled')
        self.redo_button.configure(state='normal' if editable and self.session.redo_stack else 'disabled')
        self.status.set(f'{self.session.probe.path.name} / {sheet.title} / '
            f'{"読み取り専用" if not editable else "検証範囲内"} / 未保存：{"あり" if self.session.dirty else "なし"}')
        self.selection_detail()

    def history(self, redo=False):
        if self.session and self.settle():
            try:
                self.session.history(redo)
                self.bridge.context = None
                self.render()
            except Exception as exc:
                dialogs.showerror('Undo/Redo失敗', str(exc), parent=self.root)

    def save_as(self):
        if not self.session or not self.settle() or self.session.probe.read_only:
            return
        target = filedialog.asksaveasfilename(parent=self.root, defaultextension='.xlsx',
                                             filetypes=[('Excel workbook', '*.xlsx')])
        if target:
            try:
                self.session.save_as(target)
                self.render()
            except Exception as exc:
                dialogs.showerror('別名保存失敗', str(exc), parent=self.root)

    def close(self):
        if self.settle() and self.confirm_discard():
            if self.session:
                self.session.probe.book.close()
            self.root.destroy()


def make_samples(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    paths = [folder / f'inline_{name}.xlsx' for name in ('basic', 'readonly')]
    if any(path.exists() for path in paths):
        raise FileExistsError('同名サンプルが存在します。別フォルダーを指定してください')
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = 'Data'
        sheet.append(['文字列', '数値', '真偽値', '日時', '数式', '結合', 'エラー'])
        sheet.append(['00123', 1234.5, True, datetime(2026, 10, 20, 8, 15), '=B2*2', '結合', '#DIV/0!'])
        sheet['A3'] = 'セル内\n改行'
        sheet['B3'] = 42
        sheet['B2'].number_format = '#,##0.00'
        sheet['D2'].number_format = 'yyyy/mm/dd hh:mm'
        sheet.merge_cells('F2:F3')
        sheet.row_dimensions[5].height = 15
        sheet.column_dimensions['H'].width = 12
        book.create_sheet('Other')['A1'] = '他シート'
        for path in paths:
            if path.stem.endswith('readonly'):
                rule = DataValidation(type='list', formula1='"はい,いいえ"')
                sheet.add_data_validation(rule)
                rule.add('A2')
            with path.open('xb') as stream:
                book.save(stream)
            print(path)
    finally:
        book.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path', nargs='?', type=Path)
    parser.add_argument('--make-samples', type=Path)
    args = parser.parse_args(argv)
    if args.path and args.make_samples:
        parser.error('サンプル生成とGUI起動は別々に実行してください')
    if args.make_samples:
        try:
            make_samples(args.make_samples)
            return 0
        except Exception as exc:
            print(f'操作中止：{exc}')
            return 1
    root = tk.Tk()
    app = InlineProbe(root)
    if args.path:
        app.open_path(args.path)
    root.mainloop()
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
