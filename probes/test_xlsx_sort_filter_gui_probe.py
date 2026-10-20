from copy import copy

import pytest
from openpyxl import load_workbook

from probes.xlsx_sort_filter_probe import make_samples
from probes.xlsx_sort_filter_gui_probe import XlsxProjection
from sheetdock.xlsx_extent import extents
from sheetdock.xlsx_io import snapshot


@pytest.fixture
def samples(tmp_path):
    make_samples(tmp_path)
    return tmp_path


@pytest.fixture
def model(samples):
    m = XlsxProjection(samples/'sort_basic.xlsx')
    try:
        yield m
    finally:
        m.session.probe.book.close()


def tokyo(model):
    model.session.set_filter(5, query='東京')
    return model.view()[3]


def test_filter_view_maps_cell_edit_and_preserves_hidden_rows(model):
    before = snapshot(model.session.probe.book)
    token = tokyo(model)
    assert model.mapping == [2, 4, 6]
    assert model.edit(1, 0, '00099', token)
    assert model.session.sheet['A4'].value == '00099'
    assert model.session.sheet['A3'].value == '00002' and model.session.sheet['A5'].value == '00004'
    assert model.session.history() and snapshot(model.session.probe.book) == before
    assert model.session.history(True) and model.session.sheet['A4'].value == '00099'


def test_filter_key_edit_disappears_and_history_reevaluates(model):
    token = tokyo(model)
    model.edit(1, 4, '北海道', token)
    assert model.view()[2] == [2, 6]
    assert model.session.history() and model.view()[2] == [2, 4, 6]
    assert model.session.history(True) and model.view()[2] == [2, 6]


def test_paste_resolves_all_visible_targets_before_filter_re_evaluation(model):
    token = tokyo(model)
    before = snapshot(model.session.probe.book)
    assert model.paste(0, 4, [['大阪', 'first'], ['京都', 'second']], token)
    sheet = model.session.sheet
    assert sheet['E2'].value == '大阪' and sheet['F2'].value == 'first'
    assert sheet['E4'].value == '京都' and sheet['F4'].value == 'second'
    assert sheet['E3'].value == '大阪' and sheet['F3'].value == '改行\n引用符"と,カンマ'
    assert model.view()[2] == [6] and len(model.session.undo_stack) == 1
    assert model.session.history() and snapshot(model.session.probe.book) == before
    assert not model.session.dirty


@pytest.mark.parametrize('row,column,values', [(2, 0, [['a'], ['b']]), (0, 5, [['a', 'b']]),
    (0, 0, [['a'], ['b', 'c']]), (0, 0, [[123]]), (0, -1, [['a']]), (0, 0, [])])
def test_bad_paste_is_atomic(model, row, column, values):
    token = tokyo(model)
    before = snapshot(model.session.probe.book)
    with pytest.raises((ValueError, TypeError, IndexError)):
        model.paste(row, column, values, token)
    assert snapshot(model.session.probe.book) == before and not model.session.undo_stack


@pytest.mark.parametrize('text', ['abc', 'NaN', 'Infinity'])
def test_late_numeric_validation_failure_rolls_back_earlier_cells(model, text):
    token = tokyo(model)
    before = snapshot(model.session.probe.book)
    with pytest.raises(ValueError):
        model.paste(0, 0, [['changed', text]], token)
    assert snapshot(model.session.probe.book) == before and not model.session.dirty
    assert not model.session.undo_stack and not model.session.redo_stack


def test_number_date_boolean_types_and_original_formats_are_maintained(model):
    token = tokyo(model)
    original_format = model.session.sheet['B2'].number_format
    assert model.paste(0, 0, [['00099', '2500.25', '2026-10-22T09:30:00', 'FALSE']], token)
    sheet = model.session.sheet
    assert [sheet.cell(2, c).data_type for c in range(1, 5)] == ['s', 'n', 'd', 'b']
    assert sheet['A2'].value == '00099' and sheet['B2'].value == 2500.25 and sheet['D2'].value is False
    assert sheet['B2'].number_format == original_format


@pytest.mark.parametrize('change', ['filter', 'sheet', 'edit', 'sort', 'render'])
def test_stale_selection_rejected_after_context_change(model, change):
    token = model.view()[3]
    if change == 'filter':
        model.session.set_filter(5, query='東京')
    elif change == 'sheet':
        model.session.sheet_name = 'Other'
    elif change == 'edit':
        model.session.edit(1, 0, 'changed')
    elif change == 'sort':
        model.session.sort_rows(2)
    else:
        model.view()
    before = snapshot(model.session.probe.book)
    with pytest.raises(ValueError, match='選び直し'):
        model.edit(0, 0, 'wrong target', token)
    assert snapshot(model.session.probe.book) == before


@pytest.mark.parametrize('name,column,row', [('readonly', 0, 0), ('merged', 4, 0), ('error', 5, 1)])
def test_protected_targets_cannot_be_edited_through_projection(samples, name, column, row):
    m = XlsxProjection(samples/f'sort_{name}.xlsx')
    try:
        token = m.view()[3]
        before = snapshot(m.session.probe.book)
        with pytest.raises(PermissionError):
            m.edit(row, column, 'wrong', token)
        assert snapshot(m.session.probe.book) == before and not m.session.undo_stack
    finally:
        m.session.probe.book.close()


def test_formula_in_other_sheet_and_merged_range_whole_area_remain_protected(samples):
    m = XlsxProjection(samples/'sort_formula.xlsx', header=False)
    try:
        m.session.sheet_name = 'Other'
        token = m.view()[3]
        with pytest.raises(PermissionError):
            m.edit(1, 0, '=1+1', token)
        assert m.session.sheet['A2'].value == '=Data!B2'
    finally:
        m.session.probe.book.close()
    m = XlsxProjection(samples/'sort_merged.xlsx')
    try:
        token = m.view()[3]
        with pytest.raises(PermissionError):
            m.paste(0, 3, [['FALSE', 'Tokyo', 'x']], token)
        assert m.session.sheet['D2'].value is True
    finally:
        m.session.probe.book.close()


def test_zero_matches_still_saves_full_book_and_filter_is_not_saved(model, samples):
    token = tokyo(model)
    model.edit(1, 0, '00099', token)
    model.session.set_filter(5, query='不存在')
    assert model.view()[1] == []
    before, bounds = snapshot(model.session.probe.book), extents(model.session.probe.book)
    target = samples/'full_copy.xlsx'
    model.session.save_as(target)
    assert not model.session.dirty and model.session.filters
    book = load_workbook(target)
    try:
        assert snapshot(book) == before and extents(book) == bounds
        assert book['Data']['A4'].value == '00099' and book['Data']['A3'].value == '00002'
        assert book['Other']['A1'].value == '他シート不変' and book['Data'].auto_filter.ref is None
    finally:
        book.close()
    reopened = XlsxProjection(target)
    try:
        assert reopened.view()[2] == [2, 3, 4, 5, 6] and not reopened.session.filters
    finally:
        reopened.session.probe.book.close()
    assert model.session.history() and model.session.dirty
    assert model.session.history(True) and not model.session.dirty


def test_save_failure_preserves_filter_edit_and_history(model, samples):
    token = tokyo(model)
    model.edit(0, 0, '00099', token)
    path = samples/'sort_basic.xlsx'
    before = path.read_bytes()
    with pytest.raises(FileExistsError):
        model.session.save_as(path)
    assert path.read_bytes() == before and model.session.dirty and model.session.filters
    assert model.session.history() and not model.session.dirty


def test_no_change_and_zero_match_edits_never_create_history(model):
    token = tokyo(model)
    assert not model.edit(0, 0, '00123', token) and not model.session.undo_stack
    model.session.set_filter(5, query='none')
    token = model.view()[3]
    with pytest.raises(IndexError):
        model.edit(0, 0, 'wrong', token)
    assert not model.session.dirty and not model.session.undo_stack


def test_literal_formula_input_rolls_back_entire_paste(model):
    token = tokyo(model)
    before = snapshot(model.session.probe.book)
    with pytest.raises(ValueError):
        model.paste(0, 4, [['大阪', '=1+1']], token)
    assert snapshot(model.session.probe.book) == before and not model.session.undo_stack


def test_mapped_edit_sort_and_another_edit_undo_in_exact_operation_order(model):
    before = snapshot(model.session.probe.book)
    token = tokyo(model)
    model.edit(1, 0, '00099', token)
    model.session.clear_filter()
    model.session.sort_rows(2)
    token = model.view()[3]
    model.edit(0, 0, '00088', token)
    after = snapshot(model.session.probe.book)
    for _ in range(3):
        assert model.session.history()
    assert snapshot(model.session.probe.book) == before and not model.session.dirty
    for _ in range(3):
        assert model.session.history(True)
    assert snapshot(model.session.probe.book) == after
