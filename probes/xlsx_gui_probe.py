"""Run: python probes/xlsx_gui_probe.py. Experimental XLSX GUI only."""
from __future__ import annotations

from datetime import date, datetime, time, timedelta
from copy import copy
import csv
from io import BytesIO, StringIO
import os
import tempfile
import re
from pathlib import Path
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk

if __package__ in (None, ''):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
from tksheet import Sheet
from probes.xlsx_probe import ProtectedCellError, XlsxProbe, snapshot
from sheetdock.icons import apply_icon
from sheetdock.storage import ExternalChangeError, fingerprint, save_lock
from sheetdock.csv_io import CsvDocument, CsvFormat, save_csv
from probes.xlsx_storage_probe import load_checked, save_checked


VIEW_ROWS, VIEW_COLUMNS = 1000, 50
HISTORY_LIMIT = 100
_DEFAULT_EXPECTED = object()
# All model changes go through the controlled editor, never widget mutations.
VIEW_BINDINGS = ('single_select', 'drag_select', 'row_select', 'column_select',
                 'arrowkeys', 'copy', 'select_all', 'column_width_resize',
                 'row_height_resize')


def display_value(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def parse_input(text, previous):
    """Preserve existing type; new/blank cells accept literal strings only."""
    if previous is None or isinstance(previous, str):
        return text
    if isinstance(previous, bool):
        if text.strip().upper() not in {'TRUE', 'FALSE'}:
            raise ValueError('真偽値はTRUEまたはFALSEで入力してください')
        return text.strip().upper() == 'TRUE'
    if isinstance(previous, (int, float)):
        try:
            return int(text) if text.strip().lstrip('+-').isdigit() else float(text)
        except ValueError:
            raise ValueError('数値を入力してください') from None
    if isinstance(previous, (datetime, date, time)):
        return type(previous).fromisoformat(text)
    if isinstance(previous, timedelta):
        raise ValueError('時間間隔の編集は次の段階で対応します')
    raise ValueError('この型の編集には対応していません')


def protection_reason(probe, sheet, cell):
    if probe.read_only:
        return 'ブックは読み取り専用です'
    if cell.data_type == 'f':
        return '数式セルは編集できません（再計算しません）'
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            return f'結合範囲 {merged} は編集できません'
    if cell.data_type == 'e':
        return '既存のエラー値は編集できません'
    if isinstance(cell.value, timedelta):
        return '時間間隔の編集は未対応です'
    return ''


def clipboard_rows(text):
    if len(text) > 1_000_000:
        raise ValueError('今回の貼り付けは100万文字までです')
    rows = list(csv.reader(StringIO(text, newline=''), delimiter='\t', strict=True))
    if not rows:
        return []
    if not rows[0] or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError('貼り付け範囲の列数をそろえてください')
    if len(rows) * len(rows[0]) > 50_000:
        raise ValueError('今回の貼り付けは5万セルまでです')
    return rows


def cell_state(cell):
    return cell.value, cell.data_type, copy(cell._style)


def restore_cell(cell, state):
    cell._value, cell.data_type, cell._style = state[0], state[1], copy(state[2])


class XlsxSession:
    def __init__(self, path):
        self.probe, self.stamp = load_checked(path)
        if sum(s.max_row * s.max_column for s in self.probe.book) > 100_000:
            self.probe.book.close()
            raise ValueError('今回の検証GUIはブック全体10万セルまでです')
        self.sheet_name = self.probe.book.active.title
        self.saved_values = self.values()
        self.undo_stack, self.redo_stack = [], []

    @property
    def sheet(self):
        return self.probe.book[self.sheet_name]

    def values(self):
        # Include structure and formats so deleting/recreating a sheet is tracked.
        return snapshot(self.probe.book)

    @property
    def dirty(self):
        return self.values() != self.saved_values

    def edit(self, row, column, text):
        return self.apply_rows(row, column, [[text]])

    def paste(self, row, column, text):
        return self.apply_rows(row, column, clipboard_rows(text))

    def apply_rows(self, row, column, rows):
        if self.probe.read_only:
            raise ProtectedCellError('ブックは読み取り専用です')
        if not rows:
            return False
        sheet = self.sheet
        if (row < 0 or column < 0 or row + len(rows) > min(sheet.max_row, VIEW_ROWS)
                or column + len(rows[0]) > min(sheet.max_column, VIEW_COLUMNS)):
            raise ValueError('貼り付け先は現在表示中の行列範囲内を指定してください')
        pending = []
        for r, values in enumerate(rows, row + 1):
            for c, text in enumerate(values, column + 1):
                cell = sheet.cell(r, c)
                reason = protection_reason(self.probe, sheet, cell)
                if reason:
                    raise ProtectedCellError(f'{cell.coordinate}: {reason}')
                pending.append((cell, cell_state(cell), parse_input(text, cell.value)))
        try:
            for cell, before, value in pending:
                self.probe.set_value(self.sheet_name, cell.coordinate, value)
        except Exception:
            for cell, before, value in pending:
                restore_cell(cell, before)
            raise
        changes = [(self.sheet_name, cell.coordinate, before, cell_state(cell))
                   for cell, before, value in pending if before != cell_state(cell)]
        if not changes:
            return False
        self.undo_stack.append((self.sheet_name, changes))
        del self.undo_stack[:-HISTORY_LIMIT]
        self.redo_stack.clear()
        return True

    def history(self, redo=False):
        if self.probe.read_only:
            raise ProtectedCellError('ブックは読み取り専用です')
        source, target = (self.redo_stack, self.undo_stack) if redo else (self.undo_stack, self.redo_stack)
        if not source:
            return False
        name, changes = source[-1]
        if isinstance(changes, dict):
            book = load_workbook(BytesIO(changes['after' if redo else 'before']), data_only=False)
            self.probe.book.close()
            self.probe.book = book
            self.sheet_name = changes['after_name' if redo else 'before_name']
        else:
            for sheet_name, address, before, after in changes:
                restore_cell(self.probe.book[sheet_name][address], after if redo else before)
            self.sheet_name = name
        source.pop()
        target.append((name, changes))
        return True

    def structure_reason(self, operation):
        if self.probe.read_only:
            return 'ブックは読み取り専用です'
        if operation in {'rename', 'delete'} and any(c.data_type == 'f' for s in self.probe.book for row in s for c in row):
            return '数式を含むブックでは参照保護のため削除・名前変更はできません'
        if operation == 'delete' and self.sheet.sheet_state == 'visible' and sum(
                s.sheet_state == 'visible' for s in self.probe.book) <= 1:
            return '最後の表示シートは削除できません'
        return ''

    def validate_sheet_name(self, name, *, current=None):
        if (not isinstance(name, str) or not name.strip() or len(name) > 31
                or re.search(r'[\\/\?\*\[\]:\x00-\x1f]', name)
                or name.startswith("'") or name.endswith("'") or name.casefold() == 'history'):
            raise ValueError('シート名は1～31文字で、禁止文字・先頭末尾の引用符・Historyは使用できません')
        if any(s.title.casefold() == name.casefold() and s.title != current for s in self.probe.book):
            raise ValueError('同じ名前のシートがあります（大文字小文字は区別しません）')

    def pack_book(self):
        stream = BytesIO()
        self.probe.book.save(stream)
        checked = load_workbook(BytesIO(stream.getvalue()), data_only=False)
        try:
            if snapshot(checked) != snapshot(self.probe.book):
                raise ValueError('シート操作の復元用データを検証できませんでした')
        finally:
            checked.close()
        return stream.getvalue()

    def sheet_operation(self, operation, name=None):
        if operation not in {'add', 'rename', 'delete'}:
            raise ValueError('未対応のシート操作です')
        reason = self.structure_reason(operation)
        if reason:
            raise ProtectedCellError(reason)
        if operation in {'add', 'rename'}:
            self.validate_sheet_name(name, current=self.sheet_name if operation == 'rename' else None)
        if operation == 'rename' and name == self.sheet_name:
            return False
        before, before_name = self.pack_book(), self.sheet_name
        try:
            if operation == 'add':
                self.probe.book.create_sheet(name)  # Append only; no reordering.
                self.sheet_name = name
            elif operation == 'rename':
                sheet = self.sheet
                if sheet.title.casefold() == name.casefold():
                    index = 1
                    existing = {s.title.casefold() for s in self.probe.book}
                    while f'_SheetDock_{index}'.casefold() in existing:
                        index += 1
                    sheet.title = f'_SheetDock_{index}'
                sheet.title = name
                self.sheet_name = name
            else:
                self.probe.book.remove(self.sheet)
                self.sheet_name = next(s.title for s in self.probe.book if s.sheet_state == 'visible')
            self.probe.book.active = self.probe.book.index(self.sheet)
            after = self.pack_book()
        except Exception:
            self.probe.book = load_workbook(BytesIO(before), data_only=False)
            self.sheet_name = before_name
            raise
        self.undo_stack.append((before_name, {'before': before, 'after': after,
                                             'before_name': before_name, 'after_name': self.sheet_name}))
        del self.undo_stack[:-HISTORY_LIMIT]
        self.redo_stack.clear()
        return True

    def find(self, query, *, all_sheets=True, include_formulas=False):
        if not query:
            raise ValueError('検索文字列を入力してください')
        result = []
        for sheet in self.probe.book if all_sheets else [self.sheet]:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'f' and not include_formulas:
                        continue
                    if query in display_value(cell.value):
                        result.append((sheet.title, cell.row - 1, cell.column - 1,
                                       display_value(cell.value)))
        return result

    def replace_all(self, query, replacement, *, all_sheets=True):
        if self.probe.read_only:
            raise ProtectedCellError('ブックは読み取り専用です')
        matches = self.find(query, all_sheets=all_sheets, include_formulas=True)
        pending, skipped, occurrences = [], 0, 0
        for name, row, column, value in matches:
            sheet = self.probe.book[name]
            cell = sheet.cell(row + 1, column + 1)
            if protection_reason(self.probe, sheet, cell) or not isinstance(cell.value, str):
                skipped += 1
                continue
            new = cell.value.replace(query, replacement)
            if new != cell.value:
                pending.append((name, cell, cell_state(cell), new))
                occurrences += cell.value.count(query)
        try:
            for name, cell, before, value in pending:
                self.probe.set_value(name, cell.coordinate, value)
        except Exception:
            for name, cell, before, value in pending:
                restore_cell(cell, before)
            raise
        if pending:
            changes = [(name, cell.coordinate, before, cell_state(cell))
                       for name, cell, before, value in pending]
            self.undo_stack.append((self.sheet_name, changes))
            del self.undo_stack[:-HISTORY_LIMIT]
            self.redo_stack.clear()
        return len(pending), occurrences, skipped

    def export_csv(self, path, *, formulas='reject'):
        if formulas not in {'reject', 'text'}:
            raise ValueError('数式の出力方法が不正です')
        sheet = self.sheet
        if any(c.data_type == 'f' and not isinstance(c.value, str) for row in sheet for c in row):
            raise ValueError('特殊数式のCSV出力は未対応です')
        if formulas == 'reject' and any(c.data_type == 'f' for row in sheet for c in row):
            raise ValueError('数式セルがあります。「数式文字列を出力」を明示的に選んでください')
        rows = [[display_value(c.value) for c in row] for row in sheet]
        target = Path(path).resolve()
        if target.suffix.lower() != '.csv':
            raise ValueError('出力先は.csvを指定してください')
        if target.exists():
            raise FileExistsError(target)
        temporary = None
        with save_lock(target):
            try:
                with tempfile.NamedTemporaryFile(dir=target.parent, prefix=f'.{target.name}.',
                                                 suffix='.csv', delete=False) as stream:
                    temporary = Path(stream.name)
                save_csv(CsvDocument(rows, CsvFormat('utf-8', True, '\r\n', ',')),
                         temporary, backup=False)
                os.link(temporary, target)
            finally:
                if temporary is not None:
                    temporary.unlink(missing_ok=True)

    def save_as(self, target):
        if Path(target).exists():
            raise FileExistsError(target)
        stamp = save_checked(self.probe, target, None, backup=False)
        self.probe.path = Path(target).resolve()
        self.stamp = stamp
        self.saved_values = self.values()

    def save(self, *, backup=True, expected=_DEFAULT_EXPECTED):
        stamp = save_checked(self.probe, self.probe.path,
                             self.stamp if expected is _DEFAULT_EXPECTED else expected,
                             backup=backup)
        self.stamp = stamp
        self.saved_values = self.values()


class XlsxGuiProbe:
    def __init__(self, root):
        self.root, self.session = root, None
        root.title('SheetDock — Phase 2 XLSX検証')
        root.geometry('1050x700')
        toolbar = ttk.Frame(root)
        toolbar.pack(fill='x', padx=8, pady=8)
        ttk.Button(toolbar, text='XLSXを開く', command=self.open_file).pack(side='left')
        self.edit_button = ttk.Button(toolbar, text='選択セルを編集', command=self.edit_cell, state='disabled')
        self.edit_button.pack(side='left', padx=5)
        self.save_button = ttk.Button(toolbar, text='別名保存', command=self.save_as, state='disabled')
        self.save_button.pack(side='left')
        self.overwrite_button = ttk.Button(toolbar, text='上書き保存', command=self.save, state='disabled')
        self.overwrite_button.pack(side='left', padx=5)
        self.backup = tk.BooleanVar(value=True)
        ttk.Checkbutton(toolbar, text='バックアップ', variable=self.backup).pack(side='left')
        self.names = ttk.Combobox(toolbar, state='readonly', width=25)
        self.names.pack(side='left', padx=8)
        self.names.bind('<<ComboboxSelected>>', self.switch_sheet)
        ttk.Button(toolbar, text='読取専用の理由', command=self.show_reasons).pack(side='left')
        actions = ttk.Frame(root)
        actions.pack(fill='x', padx=8, pady=3)
        self.undo_button = ttk.Button(actions, text='Undo', command=lambda: self.history(False), state='disabled')
        self.undo_button.pack(side='left')
        self.redo_button = ttk.Button(actions, text='Redo', command=lambda: self.history(True), state='disabled')
        self.redo_button.pack(side='left', padx=5)
        self.paste_button = ttk.Button(actions, text='選択セルから貼り付け', command=self.paste, state='disabled')
        self.paste_button.pack(side='left')
        ttk.Label(actions, text='範囲選択→Ctrl+Cでコピー / Ctrl+Vで貼付け / Ctrl+ZでUndo / Ctrl+Shift+ZでRedo').pack(side='left', padx=8)
        structure = ttk.Frame(root)
        structure.pack(fill='x', padx=8, pady=3)
        self.structure_buttons = {}
        for operation, label in [('add', 'シート追加'), ('rename', 'シート名変更'), ('delete', 'シート削除')]:
            button = ttk.Button(structure, text=label, command=lambda op=operation: self.sheet_operation(op), state='disabled')
            button.pack(side='left', padx=3)
            self.structure_buttons[operation] = button
        ttk.Label(structure, text='数式を含むブックの削除・名前変更は制限します。追加は末尾です。').pack(side='left')
        tools = ttk.Frame(root)
        tools.pack(fill='x', padx=8, pady=3)
        ttk.Label(tools, text='検索').pack(side='left')
        self.query = tk.StringVar()
        self.query_entry = ttk.Entry(tools, textvariable=self.query, width=16)
        self.query_entry.pack(side='left')
        self.all_sheets = tk.BooleanVar(value=True)
        ttk.Checkbutton(tools, text='全シート', variable=self.all_sheets).pack(side='left')
        self.formula_search = tk.BooleanVar(value=False)
        ttk.Checkbutton(tools, text='数式も検索', variable=self.formula_search).pack(side='left')
        ttk.Button(tools, text='次を検索', command=self.find_next).pack(side='left')
        self.replacement = tk.StringVar()
        ttk.Entry(tools, textvariable=self.replacement, width=12).pack(side='left', padx=4)
        self.replace_button = ttk.Button(tools, text='文字列セルを一括置換', command=self.replace_all, state='disabled')
        self.replace_button.pack(side='left')
        exporting = ttk.Frame(root)
        exporting.pack(fill='x', padx=8, pady=3)
        self.formula_export = ttk.Combobox(exporting, state='readonly', width=25,
                                         values=['数式があれば出力中止', '数式文字列を出力'])
        self.formula_export.current(0)
        self.formula_export.pack(side='left')
        self.export_button = ttk.Button(exporting, text='選択シートをCSV出力', command=self.export_csv, state='disabled')
        self.export_button.pack(side='left', padx=5)
        ttk.Label(exporting, text='UTF-8 BOM / CRLF / カンマ。書式・結合・他シートは出力しません。').pack(side='left')
        self.search_signature, self.search_cursor = None, -1
        self.status = tk.StringVar(value='XLSXファイルを開いてください')
        ttk.Label(root, textvariable=self.status, wraplength=1000).pack(fill='x', padx=8)
        self.sheet_widget = Sheet(root, data=[['']], headers=['A'])
        self.sheet_widget.pack(fill='both', expand=True, padx=8, pady=8)
        self.sheet_widget.enable_bindings(*VIEW_BINDINGS)
        self.sheet_widget.bind('<<SheetSelect>>', self.selection_changed)
        for keys in ('<Control-v>', '<Control-V>'):
            self.sheet_widget.bind(keys, self.paste)
        # Dialogs and toolbar actions can leave focus outside the table.
        # Bind history on the owning window so all its child widgets participate.
        # Tk dialogs have their own toplevel bindtag and keep their local editing.
        for keys, callback in [('<Control-z>', lambda e: self.history(False)),
                               ('<Control-Z>', lambda e: self.history(False)),
                               ('<Control-Shift-z>', lambda e: self.history(True)),
                               ('<Control-Shift-Z>', lambda e: self.history(True))]:
            root.bind(keys, callback)
        root.bind('<Control-s>', self.save)
        root.bind('<Control-S>', self.save)
        root.bind('<Control-f>', lambda e: self.query_entry.focus_set())
        self.detail = tk.StringVar(value='')
        ttk.Label(root, textvariable=self.detail, wraplength=1000).pack(fill='x', padx=8)
        ttk.Label(root, text='編集は「選択セルを編集」から行います。数式は文字列表示、結合範囲はセルを分けて表示します。\n'
                  '灰色＝保護セル。テーマ色・数値の表示書式・太字等の再現は未対応。先頭1000行×50列まで表示。').pack(fill='x', padx=8, pady=8)
        root.protocol('WM_DELETE_WINDOW', self.close)

    def confirm_discard(self):
        return not self.session or not self.session.dirty or messagebox.askyesno(
            '未保存の変更', '編集内容を破棄して続けますか？', parent=self.root)

    def open_file(self):
        if not self.confirm_discard():
            return
        path = filedialog.askopenfilename(parent=self.root, filetypes=[('Excel workbook', '*.xlsx')])
        if not path:
            return
        try:
            session = XlsxSession(path)
            self.session = session
            self.names.configure(values=session.probe.book.sheetnames)
            self.names.set(session.sheet_name)
            self.render()
        except Exception as exc:
            messagebox.showerror('読込失敗', str(exc), parent=self.root)

    def switch_sheet(self, event=None):
        if self.session:
            self.session.sheet_name = self.names.get()
            self.render()

    def render(self):
        session = self.session
        sheet = session.sheet
        rows, columns = min(sheet.max_row, VIEW_ROWS), min(sheet.max_column, VIEW_COLUMNS)
        data = [[display_value(sheet.cell(r, c).value) for c in range(1, columns + 1)]
                for r in range(1, rows + 1)]
        self.sheet_widget.set_sheet_data(data, delete_options=True)
        self.sheet_widget.headers([get_column_letter(c) for c in range(1, columns + 1)])
        for r in range(rows):
            for c in range(columns):
                cell = sheet.cell(r + 1, c + 1)
                protected = protection_reason(session.probe, sheet, cell)
                bg, fg = None, None
                if cell.fill.patternType == 'solid' and cell.fill.fgColor.type == 'rgb':
                    bg = '#' + cell.fill.fgColor.rgb[-6:]
                if cell.font.color and cell.font.color.type == 'rgb':
                    fg = '#' + cell.font.color.rgb[-6:]
                if protected:
                    bg, fg = '#e5e7eb', '#374151'
                if bg or fg:
                    self.sheet_widget.highlight_cells(r, c, bg=bg or False, fg=fg or False, redraw=False)
        self.sheet_widget.refresh()
        mode = '読み取り専用' if session.probe.read_only else '検証対象範囲内'
        self.status.set(f'{session.probe.path.name} / {sheet.title} / {mode} / '
                        f'未保存={session.dirty} / 表示 {rows}行×{columns}列（全体 {sheet.max_row}行×{sheet.max_column}列）')
        self.save_button.configure(state='disabled' if session.probe.read_only else 'normal')
        self.overwrite_button.configure(state='disabled' if session.probe.read_only else 'normal')
        self.replace_button.configure(state='disabled' if session.probe.read_only else 'normal')
        self.export_button.configure(state='normal')
        self.edit_button.configure(state='disabled')
        self.paste_button.configure(state='disabled')
        self.undo_button.configure(state='normal' if session.undo_stack else 'disabled')
        self.redo_button.configure(state='normal' if session.redo_stack else 'disabled')
        self.names.set(session.sheet_name)
        self.names.configure(values=session.probe.book.sheetnames)
        for operation, button in self.structure_buttons.items():
            button.configure(state='disabled' if session.structure_reason(operation) else 'normal')
        self.detail.set('セルを選択してください')

    def selected_cell(self):
        selected = self.sheet_widget.get_currently_selected()
        if not self.session or not selected:
            return None
        if selected.row >= min(self.session.sheet.max_row, VIEW_ROWS) or selected.column >= min(self.session.sheet.max_column, VIEW_COLUMNS):
            return None
        return self.session.sheet.cell(selected.row + 1, selected.column + 1)

    def selection_changed(self, event=None):
        cell = self.selected_cell()
        if cell is None:
            self.edit_button.configure(state='disabled')
            self.paste_button.configure(state='disabled')
            return
        reason = protection_reason(self.session.probe, self.session.sheet, cell)
        self.edit_button.configure(state='disabled' if reason else 'normal')
        self.paste_button.configure(state='disabled' if reason else 'normal')
        self.detail.set(f'{cell.coordinate} / 型={cell.data_type} / {reason or "編集可能（既存の型を維持）"}')

    def edit_cell(self):
        cell = self.selected_cell()
        if cell is None:
            return
        reason = protection_reason(self.session.probe, self.session.sheet, cell)
        if reason:
            messagebox.showinfo('編集制限', reason, parent=self.root)
            return
        text = simpledialog.askstring('セル編集', f'{cell.coordinate}（型={cell.data_type}）の新しい値',
                                      initialvalue=display_value(cell.value), parent=self.root)
        if text is None:
            return
        try:
            self.session.edit(cell.row - 1, cell.column - 1, text)
            self.render()
        except Exception as exc:
            messagebox.showerror('編集失敗', str(exc), parent=self.root)

    def save_as(self):
        if not self.session or self.session.probe.read_only:
            return
        target = filedialog.asksaveasfilename(parent=self.root, defaultextension='.xlsx',
                                             filetypes=[('Excel workbook', '*.xlsx')])
        if not target:
            return
        try:
            self.session.save_as(target)
            self.render()
            messagebox.showinfo('別名保存完了', 'コピーを保存しました。元ファイルは変更していません。', parent=self.root)
        except Exception as exc:
            messagebox.showerror('保存失敗', str(exc), parent=self.root)

    def paste(self, event=None):
        cell = self.selected_cell()
        if cell is None:
            return 'break'
        try:
            text = self.root.clipboard_get()
            self.session.paste(cell.row - 1, cell.column - 1, text)
            self.render()
        except Exception as exc:
            messagebox.showerror('貼り付け失敗', str(exc), parent=self.root)
        return 'break'

    def save(self, event=None):
        if not self.session or self.session.probe.read_only:
            return 'break'
        try:
            self.session.save(backup=self.backup.get())
            self.render()
        except ExternalChangeError:
            self.resolve_external_change()
        except Exception as exc:
            messagebox.showerror('保存失敗', str(exc), parent=self.root)
        return 'break'

    def resolve_external_change(self):
        try:
            observed = fingerprint(self.session.probe.path)
            answer = messagebox.askyesnocancel(
                '外部変更を検出',
                'ファイルが外部で変更または削除されています。\n'
                'はい：現在の編集内容で上書き\n'
                'いいえ：ファイルを再読込（編集内容は破棄）\n'
                'キャンセル：編集を続ける', parent=self.root)
            if answer is True:
                self.session.save(backup=self.backup.get(), expected=observed)
                self.render()
            elif answer is False and self.confirm_discard():
                candidate = XlsxSession(self.session.probe.path)
                self.session = candidate
                self.names.configure(values=candidate.probe.book.sheetnames)
                self.render()
        except Exception as exc:
            messagebox.showerror('保存・再読込失敗', str(exc), parent=self.root)

    def history(self, redo=False):
        if self.session:
            try:
                if self.session.history(redo):
                    self.render()
            except Exception as exc:
                messagebox.showerror('履歴操作失敗', str(exc), parent=self.root)
        return 'break'

    def find_next(self):
        if not self.session:
            return
        try:
            matches = self.session.find(self.query.get(), all_sheets=self.all_sheets.get(),
                                        include_formulas=self.formula_search.get())
            if not matches:
                messagebox.showinfo('検索', '見つかりませんでした', parent=self.root)
                return
            signature = (id(self.session), self.query.get(), self.all_sheets.get(),
                         self.formula_search.get(), tuple(matches))
            self.search_cursor = (self.search_cursor + 1) % len(matches) if signature == self.search_signature else 0
            self.search_signature = signature
            name, row, column, value = matches[self.search_cursor]
            self.session.sheet_name = name
            self.render()
            if row < VIEW_ROWS and column < VIEW_COLUMNS:
                self.sheet_widget.select_cell(row, column)
                self.sheet_widget.see(row, column)
                self.selection_changed()
            else:
                messagebox.showinfo('検索結果', f'{name}!{get_column_letter(column+1)}{row+1} は表示上限の範囲外です', parent=self.root)
            self.detail.set(self.detail.get() + f' / 検索 {self.search_cursor+1}/{len(matches)}')
        except Exception as exc:
            messagebox.showerror('検索失敗', str(exc), parent=self.root)

    def replace_all(self):
        if not self.session or self.session.probe.read_only:
            return
        if not messagebox.askyesno('一括置換', '文字列セルだけを置換します。保護セル・数値・日時・真偽値は除外します。実行しますか？', parent=self.root):
            return
        try:
            cells, occurrences, skipped = self.session.replace_all(
                self.query.get(), self.replacement.get(), all_sheets=self.all_sheets.get())
            self.render()
            messagebox.showinfo('置換完了', f'{cells}セル・{occurrences}箇所を置換 / {skipped}セルを除外', parent=self.root)
        except Exception as exc:
            messagebox.showerror('置換失敗', str(exc), parent=self.root)

    def export_csv(self):
        if not self.session:
            return
        target = filedialog.asksaveasfilename(parent=self.root, defaultextension='.csv',
                                             filetypes=[('CSV', '*.csv')])
        if not target:
            return
        try:
            mode = 'text' if self.formula_export.current() == 1 else 'reject'
            self.session.export_csv(target, formulas=mode)
            messagebox.showinfo('CSV出力完了', '選択シートの値をCSVへ出力しました。XLSXの保存状態は変わりません。\n数式文字列は、Excel等で開くと数式として解釈される場合があります。', parent=self.root)
        except Exception as exc:
            messagebox.showerror('CSV出力失敗', str(exc), parent=self.root)

    def show_reasons(self):
        if self.session:
            reasons = self.session.probe.reasons
            messagebox.showinfo('読取専用の理由', '\n'.join(reasons) if reasons else '現在の検証範囲で未対応機能は検出されていません。', parent=self.root)

    def sheet_operation(self, operation):
        if not self.session:
            return
        try:
            reason = self.session.structure_reason(operation)
            if reason:
                raise ProtectedCellError(reason)
            name = None
            if operation in {'add', 'rename'}:
                name = simpledialog.askstring('シート操作', 'シート名を入力してください', parent=self.root,
                    initialvalue=self.session.sheet_name if operation == 'rename' else '')
                if name is None:
                    return
            elif not messagebox.askyesno('シート削除', f'「{self.session.sheet_name}」を削除しますか？\nUndoで元に戻せます。', parent=self.root):
                return
            self.session.sheet_operation(operation, name)
            self.render()
        except Exception as exc:
            messagebox.showerror('シート操作失敗', str(exc), parent=self.root)

    def close(self):
        if self.confirm_discard():
            self.root.destroy()


def main():
    root = tk.Tk()
    apply_icon(root)
    XlsxGuiProbe(root)
    root.update_idletasks()
    root.geometry(f'+{max(0, (root.winfo_screenwidth()-1050)//2)}+{max(0, (root.winfo_screenheight()-700)//2)}')
    root.mainloop()


if __name__ == '__main__':
    main()
