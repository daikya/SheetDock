"""Phase 2 experimental XLSX backend. Not used by the production GUI."""
from __future__ import annotations

import argparse
from copy import copy
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path
import re
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.writer.theme import theme_xml


class ProtectedCellError(ValueError):
    pass


# Deliberately narrow probe coverage. Unknown elements/parts block writing.
_TAGS = set("""
worksheet sheetPr outlinePr pageSetUpPr dimension sheetViews sheetView selection
sheetFormatPr cols col sheetData row c v f is t mergeCells mergeCell pageMargins
workbook workbookPr workbookProtection bookViews workbookView sheets sheet
definedNames calcPr sst si
styleSheet numFmts numFmt fonts font name charset family b i strike outline shadow
condense extend color sz u vertAlign scheme fills fill patternFill fgColor bgColor
borders border start end left right top bottom diagonal vertical horizontal
cellStyleXfs cellXfs xf alignment protection cellStyles cellStyle colors indexedColors
rgbColor tableStyles
Properties Application AppVersion coreProperties creator lastModifiedBy created
modified title subject description keywords category revision version identifier
language contentStatus lastPrinted
Relationships Relationship Types Default Override
""".split())
_NAMESPACES = {
    'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties',
    'http://schemas.openxmlformats.org/package/2006/metadata/core-properties',
    'http://purl.org/dc/elements/1.1/', 'http://purl.org/dc/terms/',
    'http://schemas.openxmlformats.org/package/2006/relationships',
    'http://schemas.openxmlformats.org/package/2006/content-types',
}
_ATTR_NAMESPACES = {
    'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'http://www.w3.org/XML/1998/namespace',
    'http://www.w3.org/2001/XMLSchema-instance',
}
_RELATIONS = {
    'http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + suffix
    for suffix in ('officeDocument', 'extended-properties', 'worksheet', 'styles',
                   'theme', 'sharedStrings')
} | {'http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties'}


def inspect_package(data: bytes) -> tuple[str, ...]:
    """Report unsupported structures before openpyxl can discard them.

    This is a probe gate, not a complete OOXML compatibility/security validator.
    """
    reasons = set()
    with ZipFile(BytesIO(data)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            reasons.add('ZIP内の部品名が重複しています')
        for name in names:
            known = name in {
                '[Content_Types].xml', '_rels/.rels', 'docProps/app.xml',
                'docProps/core.xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels',
                'xl/styles.xml', 'xl/sharedStrings.xml', 'xl/theme/theme1.xml',
            } or re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name)
            if not known:
                reasons.add(f'未対応の部品: {name}')
                continue
            raw = archive.read(name)
            if name == 'xl/theme/theme1.xml':
                if raw != theme_xml.encode():
                    reasons.add('標準テーマ以外のテーマ')
                continue
            root = ET.fromstring(raw)
            for node in root.iter():
                namespace, _, tag = node.tag[1:].partition('}')
                if namespace not in _NAMESPACES or tag not in _TAGS:
                    reasons.add(f'未対応のXML要素: {tag}')
                if any(key.startswith('{') and key[1:].split('}')[0]
                       not in _ATTR_NAMESPACES for key in node.attrib):
                    reasons.add('未対応の属性名前空間')
                if tag == 'Relationship' and (
                    node.get('Type') not in _RELATIONS or node.get('TargetMode')
                ):
                    reasons.add('未対応または外部参照の関連付け')
                if tag in {'definedNames', 'workbookProtection'} and (len(node) or node.attrib):
                    reasons.add('名前定義またはブック保護')
                if tag == 'f' and node.attrib:
                    reasons.add('共有・配列などの特殊数式')
                if tag == 'r':
                    reasons.add('リッチテキスト')
    return tuple(sorted(reasons))


def snapshot(book):
    """Compare values, types, cell formats, merges and sheet order on reload."""
    return tuple((sheet.title, sheet.sheet_state,
                  tuple((key, tuple(sorted(dict(dim).items())))
                        for key, dim in sorted(sheet.column_dimensions.items())),
                  tuple((key, tuple(sorted(dict(dim).items())))
                        for key, dim in sorted(sheet.row_dimensions.items())),
                  tuple(sorted(str(r) for r in sheet.merged_cells.ranges)),
                  tuple((cell.coordinate, cell.value, cell.data_type,
                         copy(cell.font), copy(cell.fill), copy(cell.border),
                         copy(cell.alignment), cell.number_format, copy(cell.protection))
                        for row in sheet for cell in row
                        if cell.value is not None or cell.has_style)) for sheet in book)


class XlsxProbe:
    def __init__(self, path):
        self.path = Path(path).resolve()
        if self.path.suffix.lower() != '.xlsx':
            raise ValueError('.xlsxだけを対象にしています')
        data = self.path.read_bytes()
        self.reasons = list(inspect_package(data))
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            self.book = load_workbook(BytesIO(data), data_only=False, keep_links=True)
        if caught:
            self.reasons.append('openpyxlが読込時に警告を出しました')

    @property
    def read_only(self):
        return bool(self.reasons)

    def set_value(self, sheet_name, address, value):
        if self.read_only:
            raise PermissionError('読み取り専用です')
        sheet = self.book[sheet_name]
        cell = sheet[address]
        if cell.data_type == 'f' or any(cell.coordinate in r for r in sheet.merged_cells.ranges):
            raise ProtectedCellError('数式セルと結合範囲は編集できません')
        if isinstance(value, str) and value.startswith('='):
            raise ProtectedCellError('数式の入力には対応していません')
        if isinstance(value, str) and (len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value)):
            raise ValueError('Excelで保持できない文字列です')
        old = cell.value
        # Excel/openpyxl serializes a literal empty string as an empty cell.
        # Treat explicit empty text input as clearing the cell before validation.
        if value == '' and isinstance(value, str):
            value = None
        def family(v):
            if v is None:
                return 'blank'
            if isinstance(v, bool):
                return 'bool'
            if isinstance(v, (int, float)):
                return 'number'
            if isinstance(v, (date, datetime, time, timedelta)):
                return type(v)
            return type(v)
        if old is not None and value is not None and family(old) != family(value):
            raise TypeError('既存セルの型を変更できません')
        if isinstance(value, float):
            import math
            if not math.isfinite(value):
                raise ValueError('有限の数値を指定してください')
        if value is not None and not isinstance(value, (str, bool, int, float, date, datetime, time, timedelta)):
            raise TypeError('未対応の値です')
        if cell.data_type == 'e':
            raise ProtectedCellError('既存のエラー値は編集できません')
        cell.value = value
        if isinstance(value, str):
            cell.data_type = 's'  # '#N/A' remains a literal string.

    def save_as(self, destination):
        """New files only; never overwrite the source or another existing file."""
        if self.read_only:
            raise PermissionError('未対応機能を含むため保存できません')
        destination = Path(destination)
        if destination.suffix.lower() != '.xlsx':
            raise ValueError('保存先は.xlsxを指定してください')
        if destination.exists():
            raise FileExistsError(destination)
        output = BytesIO()
        self.book.save(output)
        output.seek(0)
        checked = load_workbook(output, data_only=False)
        try:
            if snapshot(self.book) != snapshot(checked):
                raise ValueError('再読込の検証で差異が見つかりました')
        finally:
            checked.close()
        # Exclusive creation closes the exists()/write race; source stays untouched.
        with destination.open('xb') as handle:
            try:
                handle.write(output.getvalue())
                handle.flush()
                import os
                os.fsync(handle.fileno())
            except BaseException:
                handle.close()
                destination.unlink(missing_ok=True)
                raise


def main():
    parser = argparse.ArgumentParser(description='SheetDock Phase 2 XLSX技術検証（GUI未接続）')
    parser.add_argument('source', type=Path)
    parser.add_argument('--save-as', type=Path)
    args = parser.parse_args()
    probe = XlsxProbe(args.source)
    print('シート:', ', '.join(probe.book.sheetnames))
    print('判定:', '読み取り専用' if probe.read_only else '検証対象範囲内（互換性保証ではありません）')
    for reason in probe.reasons:
        print(' -', reason)
    if args.save_as:
        probe.save_as(args.save_as)
        print('別名保存と再読込検証が完了しました:', args.save_as)


if __name__ == '__main__':
    main()
