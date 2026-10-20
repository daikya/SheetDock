"""Independent typed XLSX row sort and non-destructive filter validation."""
from __future__ import annotations

import argparse
from copy import copy
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from sheetdock.xlsx_axis import HISTORY_LIMIT, VIEW_ROWS, VIEW_COLUMNS
from sheetdock.xlsx_extent import extent, extents
from sheetdock.xlsx_gui import XlsxSession, display_value
from sheetdock.xlsx_io import ProtectedCellError, inspect_package, snapshot


def typed_key(cell):
    if cell is None or cell.value is None or cell.value == '':
        return None, None
    if cell.data_type in {'f', 'e'}:
        raise ProtectedCellError('数式・エラー値は並べ替えキーにできません')
    value = cell.value
    if type(value) is str:
        return 'text', value
    if type(value) is bool:
        return 'boolean', value
    if type(value) in (int, float):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError('有限でない数値は比較できません')
        return 'number', Decimal(str(value))
    if type(value) in (datetime, time):
        if value.tzinfo is not None:
            raise ValueError('タイムゾーン付きの日時は対象外です')
        return type(value).__name__, value
    if type(value) in (date, timedelta):
        return type(value).__name__, value
    raise TypeError('この型の並べ替えには対応していません')


def relocated_cells(records, mapping):
    moved = []
    for record in records:
        row, column = coordinate_to_tuple(record[0])
        target = mapping.get(row, row)
        moved.append((target, column, (f'{get_column_letter(column)}{target}', *record[1:])))
    return tuple(record for _, _, record in sorted(moved, key=lambda item: item[:2]))


class SortFilterSession(XlsxSession):
    def __init__(self, path):
        super().__init__(path)
        self.filters = {}

    def check_column(self, column):
        rows, columns = extent(self.sheet)
        if rows > VIEW_ROWS or columns > VIEW_COLUMNS:
            raise ValueError('今回の検証は対象シート全体1000行×50列以内です')
        if type(column) is not int or not 1 <= column <= columns:
            raise ValueError('列番号は既存範囲内の1始まりで指定してください')

    def sort_rows(self, column, *, header=True, descending=False):
        self.check_column(column)
        if type(header) is not bool or type(descending) is not bool:
            raise TypeError('見出し・降順はboolで指定してください')
        if self.sheet_name in self.filters:
            raise ValueError('並べ替え前にフィルターを解除してください')
        reason = self.axis_reason()
        if reason:
            raise ProtectedCellError(reason)
        if self.sheet.sheet_format.zeroHeight or self.sheet.sheet_format.defaultRowHeight == 0:
            raise ProtectedCellError('標準行が非表示のシートは対象外です')
        if any(cell.data_type == 'e' for cell in self.sheet._cells.values()):
            raise ProtectedCellError('対象シートにエラー値があるため並べ替えできません')
        count, _ = extent(self.sheet)
        start = 2 if header else 1
        filled, blanks, family = [], [], None
        for row in range(start, count + 1):
            kind, key = typed_key(self.sheet._cells.get((row, column)))
            if kind is None:
                blanks.append(row)
            else:
                if family is not None and kind != family:
                    raise TypeError('キー列の型が混在しています。今回の並べ替えは同じ型だけが対象です')
                family = kind
                filled.append((key, row))
        filled.sort(key=lambda item: item[0], reverse=descending)
        order = [row for _, row in filled] + blanks
        if order == list(range(start, count + 1)):
            return False
        mapping = {source: target for target, source in enumerate(order, start)}
        before = self.pack_book()
        baseline, old_extents = snapshot(self.probe.book), extents(self.probe.book)
        candidate = load_workbook(BytesIO(before), data_only=False)
        try:
            sheet = candidate[self.sheet_name]
            old_cells = dict(sheet._cells)
            sheet._cells.clear()
            for (row, column), cell in old_cells.items():
                moved = copy(cell)
                moved.parent, moved.row, moved.column = sheet, mapping.get(row, row), column
                moved._style = copy(cell._style)
                sheet._cells[(moved.row, column)] = moved
            dimensions = list(sheet.row_dimensions.items())
            sheet.row_dimensions.clear()
            for index, dimension in dimensions:
                moved = copy(dimension)
                moved.parent, moved.index = sheet, mapping.get(index, index)
                moved._style = copy(dimension._style)
                sheet.row_dimensions[moved.index] = moved
            actual = snapshot(candidate)
            active = candidate.sheetnames.index(self.sheet_name)
            for i, (old, new) in enumerate(zip(baseline, actual)):
                if i != active:
                    if old != new:
                        raise ValueError('他シートが変化したため中止しました')
                    continue
                if any(old[j] != new[j] for j in (0, 1, 2, 4)):
                    raise ValueError('列定義・シート情報・結合状態が変化しました')
                expected_dimensions = tuple(sorted((mapping.get(index, index), value) for index, value in old[3]))
                if new[3] != expected_dimensions or new[5] != relocated_cells(old[5], mapping):
                    raise ValueError('値・型・書式・行定義の移動結果が期待と一致しません')
            if extents(candidate) != old_extents:
                raise ValueError('並べ替えで論理的な表示範囲が変化しました')
            stream = BytesIO()
            candidate.save(stream)
            after = stream.getvalue()
            if inspect_package(after):
                raise ValueError('保存後に未対応の構造を検出しました')
            checked = load_workbook(BytesIO(after), data_only=False)
            try:
                if snapshot(checked) != actual or extents(checked) != old_extents:
                    raise ValueError('保存・再読込後の内容が一致しません')
            finally:
                checked.close()
            self.probe.book.close()
            self.probe.book, candidate = candidate, None
            self.undo_stack.append((self.sheet_name, {'before': before, 'after': after,
                'before_name': self.sheet_name, 'after_name': self.sheet_name}))
            del self.undo_stack[:-HISTORY_LIMIT]
            self.redo_stack.clear()
            return True
        finally:
            if candidate is not None:
                candidate.close()

    def set_filter(self, column, *, operator='contains', query='', header=True):
        self.check_column(column)
        if self.sheet.sheet_format.zeroHeight or self.sheet.sheet_format.defaultRowHeight == 0:
            raise ValueError('標準行が非表示のシートは今回の表示検証対象外です')
        if operator not in {'contains', 'equals', 'empty', 'nonempty'} or not isinstance(query, str) or type(header) is not bool:
            raise ValueError('フィルター条件が不正です')
        self.filters[self.sheet_name] = (column, operator, query, header)

    def clear_filter(self):
        self.filters.pop(self.sheet_name, None)

    def visible_rows(self, *, header=True):
        if type(header) is not bool:
            raise TypeError('見出し指定はboolです')
        condition = self.filters.get(self.sheet_name)
        if condition:
            column, operator, query, header = condition
        count, _ = extent(self.sheet)
        if count > VIEW_ROWS or extent(self.sheet)[1] > VIEW_COLUMNS:
            raise ValueError('今回の表示検証は1000行×50列以内です')
        if self.sheet.sheet_format.zeroHeight or self.sheet.sheet_format.defaultRowHeight == 0:
            raise ValueError('標準行が非表示のシートは今回の表示検証対象外です')
        indices = [row for row in range(2 if header else 1, count + 1)
                   if not ((dimension := self.sheet.row_dimensions.get(row)) is not None
                           and (dimension.hidden or dimension.height == 0))]
        if not condition:
            return indices
        visible = []
        for row in indices:
            cell = self.sheet._cells.get((row, column))
            text = display_value(cell.value if cell is not None else None)
            if {'contains': lambda: query in text, 'equals': lambda: query == text,
                'empty': lambda: text == '', 'nonempty': lambda: text != ''}[operator]():
                visible.append(row)
        return visible


def make_samples(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    names = ('basic', 'formula', 'merged', 'readonly', 'mixed', 'hidden', 'error')
    paths = [folder / f'sort_{name}.xlsx' for name in names]
    if any(path.exists() for path in paths):
        raise FileExistsError('同名サンプルが存在します。別フォルダーを指定してください')
    for name, path in zip(names, paths):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Data'
        sheet.append(['コード', '数量', '日付', '有効', '区分', 'メモ'])
        for row, (code, number, place) in enumerate([
                ('00123', 10, '東京'), ('00002', 2, '大阪'), ('00003', 2, '東京'),
                ('00004', None, '京都'), ('00005', -1, '東京')], 2):
            sheet.append([code, number, datetime(2026, 10, row), row % 2 == 0, place, '改行\n引用符"と,カンマ'])
            sheet.row_dimensions[row].height = 18 + row * 3
            sheet.cell(row, 2).number_format = '#,##0.00'
            sheet.cell(row, 3).number_format = 'yyyy-mm-dd'
            sheet.cell(row, 1).font = Font(bold=row % 2 == 0, color='FF115599')
            sheet.cell(row, 5).fill = PatternFill('solid', fgColor='FFDDFFDD' if row % 2 else 'FFFFDDDD')
            sheet.cell(row, 6).alignment = Alignment(wrap_text=True)
        sheet.column_dimensions['A'].width = 18
        sheet.column_dimensions['F'].width = 30
        book.create_sheet('Other')['A1'] = '他シート不変'
        if name == 'formula':
            book['Other']['A2'] = '=Data!B2'
        elif name == 'merged':
            sheet.merge_cells('E2:F2')
        elif name == 'readonly':
            rule = DataValidation(type='list', formula1='"東京,大阪"')
            sheet.add_data_validation(rule)
            rule.add('E2')
        elif name == 'mixed':
            sheet['B3'] = '2'
        elif name == 'hidden':
            sheet.row_dimensions[3].hidden = True
        elif name == 'error':
            sheet['F3'] = '#N/A'
        try:
            with path.open('xb') as stream:
                book.save(stream)
        finally:
            book.close()
        print(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path', nargs='?', type=Path)
    parser.add_argument('--make-samples', type=Path)
    parser.add_argument('--sheet')
    parser.add_argument('--sort-column', type=int)
    parser.add_argument('--descending', action='store_true')
    parser.add_argument('--no-header', action='store_true')
    parser.add_argument('--filter-column', type=int)
    parser.add_argument('--operator', choices=('contains', 'equals', 'empty', 'nonempty'), default='contains')
    parser.add_argument('--query', default='')
    parser.add_argument('--save-as', type=Path)
    args = parser.parse_args(argv)
    if args.make_samples:
        if args.path or args.sort_column is not None or args.filter_column is not None or args.save_as or args.sheet:
            parser.error('サンプル生成と操作は別々に実行してください')
    elif not args.path:
        parser.error('XLSXパスが必要です')
    try:
        if args.make_samples:
            make_samples(args.make_samples)
            return 0
        if args.save_as and args.save_as.exists():
            raise FileExistsError('出力先が既に存在します。別の名前を指定してください')
        if args.save_as and args.filter_column is not None:
            raise ValueError('今回のCLIはフィルター結果を保存しません。--save-asを外してください')
        original = args.path.read_bytes()
        session = SortFilterSession(args.path)
        try:
            if args.sheet:
                if args.sheet not in session.probe.book.sheetnames:
                    raise ValueError('指定したシートが存在しません')
                session.sheet_name = args.sheet
            before, bounds = snapshot(session.probe.book), extents(session.probe.book)
            if args.sort_column is not None:
                changed = session.sort_rows(args.sort_column, header=not args.no_header, descending=args.descending)
                after = snapshot(session.probe.book)
                if changed:
                    if not session.history() or snapshot(session.probe.book) != before or extents(session.probe.book) != bounds:
                        raise ValueError('Undo検証に失敗しました')
                    if not session.history(True) or snapshot(session.probe.book) != after:
                        raise ValueError('Redo検証に失敗しました')
                print('並べ替え・Undo/Redoの検証完了' if changed else '並べ替え結果は同じです。履歴は追加していません')
            if args.filter_column is not None:
                baseline = snapshot(session.probe.book)
                dirty, undo, redo = session.dirty, list(session.undo_stack), list(session.redo_stack)
                session.set_filter(args.filter_column, operator=args.operator, query=args.query, header=not args.no_header)
                visible = session.visible_rows(header=not args.no_header)
                if (snapshot(session.probe.book) != baseline or session.dirty != dirty
                        or session.undo_stack != undo or session.redo_stack != redo):
                    raise ValueError('フィルターがブックか履歴を変更しました')
                print('フィルターの非破壊検証完了')
            else:
                visible = session.visible_rows(header=not args.no_header)
            print('表示対象の本文行（1始まり、見出し除外）:' if not args.no_header
                  else '表示対象のデータ行（1始まり、見出しなし）:', visible)
            print('未保存:', 'あり' if session.dirty else 'なし')
            for row in visible:
                cell = session.sheet._cells.get((row, 1))
                print(f'データ行{row}: {display_value(cell.value if cell else None)}')
            if args.save_as:
                expected = snapshot(session.probe.book)
                session.save_as(args.save_as)
                reopened = SortFilterSession(args.save_as)
                try:
                    if snapshot(reopened.probe.book) != expected or extents(reopened.probe.book) != bounds:
                        raise ValueError('別名保存・再読込検証に失敗しました')
                finally:
                    reopened.probe.book.close()
                print(f'別名保存・再読込検証完了: {args.save_as}')
            if args.path.read_bytes() != original:
                raise ValueError('元ファイルが変更されました')
            print('元ファイルは変更していません。本体GUIにはまだ接続していません。')
            return 0
        finally:
            session.probe.book.close()
    except Exception as error:
        print(f'操作中止：{error}')
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
