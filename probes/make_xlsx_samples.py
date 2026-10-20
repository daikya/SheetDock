"""Generate synthetic Phase 2 fixtures. Existing files are never overwritten."""
from datetime import datetime, time
from pathlib import Path
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


def create_samples(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = '売上'
    sheet.append(['商品コード', '数量', '単価', '合計', '日付', '確認'])
    sheet.append(['00123', 3, 120.5, '=B2*C2', datetime(2026, 10, 15), True])
    sheet['A1'].font = Font(bold=True, color='FFFFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='FF2563EB')
    sheet['A4'] = '結合セル（編集禁止）'
    sheet.merge_cells('A4:C4')
    sheet.column_dimensions['A'].width = 28
    book.create_sheet('補足')['A1'] = '日本語と先頭ゼロ・型・数式・書式の検証用'
    for name in ('basic.xlsx', 'validation_readonly.xlsx'):
        if name.startswith('validation'):
            rule = DataValidation(type='list', formula1='"はい,いいえ"')
            sheet.add_data_validation(rule)
            rule.add('F2')
        data = BytesIO()
        book.save(data)
        target = folder / name
        with target.open('xb') as handle:
            handle.write(data.getvalue())
        print(target)


def create_display_sample(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = '表示書式'
    sheet.append(['確認項目', '値（表示書式あり）', '表示形式'])
    cases = [
        ('桁区切り', 12345, '#,##0'),
        ('小数と桁区切り', 1234.5, '#,##0.00'),
        ('負の数', -1234.567, '#,##0.00'),
        ('パーセント', 0.125, '0.00%'),
        ('日付', datetime(2026, 10, 15), 'yyyy/mm/dd'),
        ('日時', datetime(2026, 10, 15, 9, 5, 3), 'yyyy-mm-dd hh:mm:ss'),
        ('時刻', time(9, 5, 3), 'hh:mm:ss'),
        ('先頭ゼロの文字列', '00123', 'General'),
        ('未対応書式（元の値を表示）', 1234, '0.00E+00'),
        ('明示的な左寄せ', 1234, '#,##0'),
        ('明示的な中央寄せ', '中央', 'General'),
        ('明示的な右寄せ', '右側', 'General'),
        ('数式（文字列・編集禁止）', '=B2*2', '#,##0'),
    ]
    for label, value, pattern in cases:
        sheet.append([label, value, pattern])
        sheet.cell(sheet.max_row, 2).number_format = pattern
    for row, horizontal in ((11, 'left'), (12, 'center'), (13, 'right')):
        sheet.cell(row, 2).alignment = Alignment(horizontal=horizontal)
    sheet.column_dimensions['A'].width = 36
    sheet.column_dimensions['B'].width = 30
    sheet.column_dimensions['C'].width = 32
    book.create_sheet('書式なし').append(['比較用', 1234.5, 0.125])
    target = folder / 'display_formats.xlsx'
    with target.open('xb') as handle:
        book.save(handle)
    print(target)


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument('--sheet-operations', action='store_true')
    modes.add_argument('--display-formats', action='store_true')
    args = parser.parse_args()
    folder = Path(__file__).parent / 'xlsx_samples'
    if args.display_formats:
        create_display_sample(folder)
    elif args.sheet_operations:
        folder.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = '一覧'
        sheet.append(['商品コード', '商品名'])
        sheet.append(['00123', '日本語テスト'])
        sheet['A1'].font = Font(bold=True, color='FF2563EB')
        sheet['A4'] = '結合テスト'
        sheet.merge_cells('A4:B4')
        sheet.column_dimensions['A'].width = 24
        book.create_sheet('補足')['A1'] = '数式を含まないシート操作用サンプル'
        target = folder / 'sheet_operations.xlsx'
        with target.open('xb') as stream:
            book.save(stream)
        print(target)
    else:
        create_samples(folder)
