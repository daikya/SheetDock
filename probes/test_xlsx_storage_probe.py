from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

import probes.xlsx_storage_probe as storage
from probes.xlsx_gui_probe import XlsxSession, XlsxGuiProbe
from probes.xlsx_probe import snapshot
from sheetdock.storage import ExternalChangeError, fingerprint, save_lock


@pytest.fixture
def path(tmp_path):
    book = Workbook()
    book.active['A1'] = 'before'
    book.active['B1'] = '=1+1'
    path = tmp_path / 'source.xlsx'
    book.save(path)
    return path


def external_edit(path, value='external'):
    book = load_workbook(path)
    book.active['A1'] = value
    book.save(path)


def assert_no_temps(path):
    assert not list(path.parent.glob(f'.{path.name}.*'))


def test_overwrite_backup_and_history(path):
    before = path.read_bytes()
    session = XlsxSession(path)
    session.edit(0, 0, 'after')
    session.save()
    assert load_workbook(path).active['A1'].value == 'after'
    assert Path(str(path) + '.bak').read_bytes() == before
    assert session.stamp == fingerprint(path)
    assert not session.dirty
    assert session.history()
    assert session.dirty
    assert_no_temps(path)


def test_backup_disabled(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'after')
    session.save(backup=False)
    assert not Path(str(path) + '.bak').exists()


def test_second_save_replaces_backup_with_previous_version(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'first')
    session.save()
    first = path.read_bytes()
    session.edit(0, 0, 'second')
    session.save()
    assert Path(str(path) + '.bak').read_bytes() == first


def test_external_change_retains_edits_and_file(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    stamp = session.stamp
    external_edit(path)
    before = path.read_bytes()
    with pytest.raises(ExternalChangeError):
        session.save()
    assert path.read_bytes() == before
    assert session.probe.book.active['A1'].value == 'local'
    assert session.stamp == stamp
    assert session.dirty
    assert not Path(str(path) + '.bak').exists()
    assert_no_temps(path)


def test_explicit_overwrite_backs_up_latest_external_file(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    external_edit(path)
    external = path.read_bytes()
    session.save(expected=fingerprint(path))
    assert Path(str(path) + '.bak').read_bytes() == external
    assert load_workbook(path).active['A1'].value == 'local'


def test_change_after_consent_is_not_overwritten(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    external_edit(path, 'first')
    consent_stamp = fingerprint(path)
    external_edit(path, 'second')
    with pytest.raises(ExternalChangeError):
        session.save(expected=consent_stamp)
    assert load_workbook(path).active['A1'].value == 'second'


def test_missing_file_requires_explicit_recreate(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    path.unlink()
    with pytest.raises(ExternalChangeError):
        session.save()
    session.save(expected=None)
    assert load_workbook(path).active['A1'].value == 'local'
    assert not session.dirty


@pytest.mark.parametrize('failure', ['write', 'verification', 'backup', 'replace'])
def test_failure_keeps_source_and_dirty(path, monkeypatch, failure):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    before = path.read_bytes()
    def fail(*args, **kwargs):
        raise OSError('simulated failure')
    if failure == 'write':
        monkeypatch.setattr(session.probe.book, 'save', fail)
    elif failure == 'verification':
        monkeypatch.setattr(storage, 'load_workbook', fail)
    elif failure == 'backup':
        monkeypatch.setattr(storage.shutil, 'copy2', fail)
    else:
        original = storage.os.replace
        def fail_target(source, destination):
            if Path(destination) == path:
                fail()
            return original(source, destination)
        monkeypatch.setattr(storage.os, 'replace', fail_target)
    with pytest.raises(OSError):
        session.save()
    assert path.read_bytes() == before
    assert session.probe.book.active['A1'].value == 'local'
    assert session.dirty
    assert_no_temps(path)


def test_backup_failure_keeps_previous_backup(path, monkeypatch):
    backup = Path(str(path) + '.bak')
    backup.write_bytes(b'previous backup')
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    def partial(source, destination):
        Path(destination).write_bytes(b'partial')
        raise OSError('failure')
    monkeypatch.setattr(storage.shutil, 'copy2', partial)
    with pytest.raises(OSError):
        session.save()
    assert backup.read_bytes() == b'previous backup'
    assert_no_temps(path)


def test_readback_mismatch_blocks_replace(path, monkeypatch):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    before = path.read_bytes()
    original = storage.snapshot
    monkeypatch.setattr(storage, 'snapshot', lambda b: ('different',) if b is session.probe.book else original(b))
    with pytest.raises(ValueError):
        session.save()
    assert path.read_bytes() == before
    assert session.dirty
    assert_no_temps(path)


def test_second_writer_lock(path):
    session = XlsxSession(path)
    with save_lock(path):
        with pytest.raises(RuntimeError, match='保存ロック'):
            session.save()
    assert_no_temps(path)


def test_change_during_backup_blocks_replace(path, monkeypatch):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    original = storage.shutil.copy2
    def change(source, destination):
        original(source, destination)
        external_edit(path, 'during backup')
    monkeypatch.setattr(storage.shutil, 'copy2', change)
    with pytest.raises(ExternalChangeError):
        session.save()
    assert load_workbook(path).active['A1'].value == 'during backup'
    assert session.dirty
    assert_no_temps(path)


def test_new_file_publication_race_does_not_overwrite(path, tmp_path, monkeypatch):
    session = XlsxSession(path)
    destination = tmp_path / 'new.xlsx'
    original = storage.os.link
    def race(source, target):
        Path(target).write_bytes(b'other writer')
        return original(source, target)
    monkeypatch.setattr(storage.os, 'link', race)
    with pytest.raises(FileExistsError):
        session.save_as(destination)
    assert destination.read_bytes() == b'other writer'
    assert session.probe.path == path.resolve()
    assert_no_temps(destination)


def test_load_race_rejected(path, monkeypatch):
    original = storage.XlsxProbe
    def changing(source):
        probe = original(source)
        external_edit(source)
        return probe
    monkeypatch.setattr(storage, 'XlsxProbe', changing)
    with pytest.raises(ExternalChangeError):
        storage.load_checked(path)


def test_readonly_blocks_checked_save(path):
    session = XlsxSession(path)
    session.probe.reasons.append('unsupported')
    with pytest.raises(PermissionError):
        session.save()


def fake_gui(session, answer, monkeypatch):
    gui = XlsxGuiProbe.__new__(XlsxGuiProbe)
    gui.session, gui.root = session, None
    gui.backup = SimpleNamespace(get=lambda: True)
    gui.names = SimpleNamespace(configure=lambda **kw: None)
    gui.render = lambda: None
    monkeypatch.setattr('probes.xlsx_gui_probe.messagebox.askyesnocancel', lambda *a, **kw: answer)
    monkeypatch.setattr('probes.xlsx_gui_probe.messagebox.askyesno', lambda *a, **kw: True)
    return gui


def test_gui_conflict_cancel_keeps_edits(path, monkeypatch):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    external_edit(path)
    gui = fake_gui(session, None, monkeypatch)
    gui.resolve_external_change()
    assert gui.session is session
    assert session.dirty
    assert load_workbook(path).active['A1'].value == 'external'


def test_gui_conflict_reload_updates_session(path, monkeypatch):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    external_edit(path)
    gui = fake_gui(session, False, monkeypatch)
    gui.resolve_external_change()
    assert gui.session.probe.book.active['A1'].value == 'external'
    assert not gui.session.dirty


def test_gui_conflict_overwrite(path, monkeypatch):
    session = XlsxSession(path)
    session.edit(0, 0, 'local')
    external_edit(path)
    gui = fake_gui(session, True, monkeypatch)
    gui.resolve_external_change()
    assert load_workbook(path).active['A1'].value == 'local'
    assert not session.dirty
