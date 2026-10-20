"""Standalone XLSX projection/edit GUI; not imported by the main app."""
from __future__ import annotations

import csv
from io import StringIO
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk

from openpyxl.utils import get_column_letter
from tksheet import Sheet

from sheetdock.csv_tools_gui import history_shortcut, INPUT_CLASSES
from sheetdock.xlsx_axis import HISTORY_LIMIT
from sheetdock.xlsx_display import format_value
from sheetdock.xlsx_extent import extent, extents
from sheetdock.xlsx_gui import cell_state, display_value, parse_input, protection_reason, restore_cell
from sheetdock.xlsx_io import snapshot

if __package__:
    from .xlsx_sort_filter_probe import SortFilterSession
else:
    from xlsx_sort_filter_probe import SortFilterSession


class XlsxProjection:
    def __init__(self, path, *, header=True):
        if type(header) is not bool:
            raise TypeError('見出し指定はboolです')
        self.session = SortFilterSession(path)
        self.header, self.mapping, self.view_token = header, [], None
        self.generation = 0

    def signature(self):
        s = self.session
        return (id(s.probe.book), s.sheet_name, self.header, tuple(sorted(s.filters.items())),
                snapshot(s.probe.book), extents(s.probe.book))

    def view(self):
        s = self.session
        s.check_column(1)
        self.mapping = s.visible_rows(header=self.header)
        _, columns = extent(s.sheet)
        def text(row, column, formatted=True):
            cell = s.sheet._cells.get((row, column))
            return '' if cell is None else (format_value(cell.value, cell.number_format)[0]
                                           if formatted else display_value(cell.value))
        headers = [f'{get_column_letter(c)}: {text(1, c, False)}' if self.header
                   else get_column_letter(c) for c in range(1, columns + 1)]
        body = [[text(row, c) for c in range(1, columns + 1)] for row in self.mapping]
        self.generation += 1
        self.view_token = (self.generation, self.signature())
        return headers, body, self.mapping[:], self.view_token

    def data_row(self, row, token):
        if token != self.view_token or token is None or token[1] != self.signature():
            raise ValueError('表示・ブック・シートが変わりました。セルを選び直してください')
        if type(row) is not int or not 0 <= row < len(self.mapping):
            raise IndexError('表示中の本文セルを選択してください')
        return self.mapping[row]

    def paste(self, row, column, values, token):
        self.data_row(row, token)
        s = self.session
        if s.probe.read_only:
            raise PermissionError('読み取り専用のブックは編集できません')
        _, columns = extent(s.sheet)
        if type(column) is not int or not 0 <= column < columns:
            raise IndexError('表示中の列を選択してください')
        if (not isinstance(values, list) or not values or
                any(not isinstance(cells, list) or not cells for cells in values) or
                len({len(cells) for cells in values}) != 1 or
                any(not isinstance(value, str) for cells in values for value in cells)):
            raise ValueError('空でない長方形の文字列範囲を指定してください')
        if row + len(values) > len(self.mapping) or column + len(values[0]) > columns:
            raise IndexError('貼り付けが表示範囲を超えます。全体を中止しました')
        pending = []
        for offset, cells in enumerate(values):
            data_row = self.mapping[row + offset]
            for c, text in enumerate(cells, column + 1):
                cell = s.sheet.cell(data_row, c)
                reason = protection_reason(s.probe, s.sheet, cell)
                if reason:
                    raise PermissionError(f'{cell.coordinate}: {reason}')
                pending.append((cell, cell_state(cell), parse_input(text, cell.value)))
        try:
            for cell, before, value in pending:
                s.probe.set_value(s.sheet_name, cell.coordinate, value)
        except Exception:
            for cell, before, value in pending:
                restore_cell(cell, before)
            raise
        changes = [(s.sheet_name, cell.coordinate, before, cell_state(cell))
                   for cell, before, value in pending if before != cell_state(cell)]
        if not changes:
            return False
        s.undo_stack.append((s.sheet_name, changes))
        del s.undo_stack[:-HISTORY_LIMIT]
        s.redo_stack.clear()
        return True

    def edit(self, row, column, value, token):
        return self.paste(row, column, [[value]], token)


class XlsxProjectionGui:
    def __init__(self, root):
        self.root, self.model = root, None
        self.mapping, self.view_token = [], None
        root.title('SheetDock — XLSX並べ替え・フィルター 独立GUI検証')
        root.geometry('1080x700')
        root.columnconfigure(0, weight=1)
        root.rowconfigure(4, weight=1)
        self.header = tk.BooleanVar(value=True)
        self.sheet_name = tk.StringVar()
        self.column = tk.StringVar(value='1')
        self.descending = tk.BooleanVar(value=False)
        self.operator = tk.StringVar(value='部分一致')
        self.query = tk.StringVar()
        self.status = tk.StringVar(value='検証用XLSXを開いてください。保存は全ブックの新規名への別名保存です。')
        self.detail = tk.StringVar()
        self.notice = tk.StringVar()
        opening = ttk.Frame(root)
        opening.grid(row=0, column=0, sticky='ew', padx=8, pady=5)
        ttk.Button(opening, text='開く', command=self.open).pack(side='left')
        ttk.Checkbutton(opening, text='次に開くブックの先頭行を見出しにする', variable=self.header).pack(side='left', padx=6)
        self.sheets = ttk.Combobox(opening, textvariable=self.sheet_name, state='readonly', width=16)
        self.sheets.pack(side='left', padx=6)
        self.sheets.bind('<<ComboboxSelected>>', self.switch_sheet)
        actions = ttk.Frame(root)
        actions.grid(row=1, column=0, sticky='ew', padx=8, pady=3)
        self.buttons = {}
        for name, label, callback in [('save', '全ブックを別名保存', self.save), ('edit', 'セル編集', self.edit),
                ('paste', '範囲貼り付け', self.paste), ('undo', 'Undo', self.history),
                ('redo', 'Redo', lambda: self.history(True))]:
            self.buttons[name] = ttk.Button(actions, text=label, command=callback, state='disabled')
            self.buttons[name].pack(side='left', padx=3)
        ttk.Label(actions, textvariable=self.notice).pack(side='left', padx=8)
        controls = ttk.Frame(root)
        controls.grid(row=2, column=0, sticky='ew', padx=8, pady=3)
        ttk.Label(controls, text='対象列番号').pack(side='left')
        ttk.Spinbox(controls, from_=1, to=50, textvariable=self.column, width=5).pack(side='left', padx=4)
        ttk.Checkbutton(controls, text='降順', variable=self.descending).pack(side='left')
        self.sort_button = ttk.Button(controls, text='型を保って並べ替え', command=self.sort)
        self.sort_button.pack(side='left', padx=4)
        ttk.Combobox(controls, textvariable=self.operator, values=('部分一致', '完全一致', '空', '空以外'), state='readonly', width=9).pack(side='left')
        ttk.Entry(controls, textvariable=self.query, width=17).pack(side='left', padx=4)
        ttk.Button(controls, text='フィルター適用', command=self.filter).pack(side='left')
        ttk.Button(controls, text='解除', command=self.clear).pack(side='left', padx=4)
        ttk.Label(root, text='行番号は元のシート行。セル選択→「セル編集」で元の値を編集。表内直接入力・行列操作・寸法変更は今回対象外。\n'
                  'フィルター中も保存は非表示行・他シートを含む全ブック。並べ替え前にフィルターを解除してください。',
                  wraplength=1040).grid(row=3, column=0, sticky='ew', padx=8, pady=4)
        self.sheet = Sheet(root, data=[], headers=[], cell_auto_resize_enabled=False)
        self.sheet.grid(row=4, column=0, sticky='nsew', padx=8, pady=4)
        self.sheet.enable_bindings('single_select', 'drag_select', 'row_select', 'column_select',
                                   'select_all', 'arrowkeys', 'copy')
        self.sheet.bind('<<SheetSelect>>', self.selection_detail)
        ttk.Label(root, textvariable=self.status, wraplength=1040).grid(row=5, column=0, sticky='ew', padx=8, pady=4)
        ttk.Label(root, textvariable=self.detail, wraplength=1040).grid(row=6, column=0, sticky='ew', padx=8, pady=4)
        self.key_tag = f'XlsxProjectionHistory_{id(self)}'
        root.bind_class(self.key_tag, '<KeyPress>', self.key)
        self.install_keys(root)
        root.bind('<Control-s>', lambda event: (self.save(), 'break')[1])
        root.protocol('WM_DELETE_WINDOW', self.close)

    def install_keys(self, widget):
        widget.bindtags((self.key_tag,) + tuple(tag for tag in widget.bindtags() if tag != self.key_tag))
        for child in widget.winfo_children():
            self.install_keys(child)

    def key(self, event):
        if event.widget.winfo_class() in INPUT_CLASSES:
            return None
        system = self.root.tk.call('tk', 'windowingsystem')
        redo = history_shortcut(event, system)
        if redo is not None:
            self.history(redo)
            return 'break'
        if event.state & 4 and (event.keysym.lower() == 'v' or (system == 'win32' and event.keycode == 86)):
            self.paste()
            return 'break'
        return None

    def run(self, action):
        if not self.model:
            return
        try:
            action()
            self.render()
            self.sheet.MT.focus_set()
        except Exception as error:
            messagebox.showerror('操作中止', str(error), parent=self.root)

    def may_discard(self):
        return not (self.model and self.model.session.dirty) or messagebox.askyesno(
            '未保存の編集', '未保存の編集を破棄しますか？', parent=self.root)

    def open(self):
        path = filedialog.askopenfilename(parent=self.root, filetypes=[('XLSX', '*.xlsx')])
        if path:
            self.open_path(path)

    def open_path(self, path):
        candidate = None
        try:
            candidate = XlsxProjection(path, header=self.header.get())
            candidate.view()
            if not self.may_discard():
                return
            old = self.model
            self.model, candidate = candidate, None
            if old:
                old.session.probe.book.close()
            self.sync_conditions()
            self.render()
            self.sheet.MT.focus_set()
        except Exception as error:
            messagebox.showerror('読込中止', str(error), parent=self.root)
        finally:
            if candidate:
                candidate.session.probe.book.close()

    def sync_conditions(self):
        s = self.model.session
        self.sheets.configure(values=s.probe.book.sheetnames)
        self.sheet_name.set(s.sheet_name)
        column, operator, query, _ = s.filters.get(s.sheet_name, (1, 'contains', '', self.model.header))
        self.column.set(str(column))
        self.operator.set({'contains': '部分一致', 'equals': '完全一致', 'empty': '空', 'nonempty': '空以外'}[operator])
        self.query.set(query)

    def switch_sheet(self, event=None):
        if not self.model:
            return
        s, previous = self.model.session, self.model.session.sheet_name
        try:
            name = self.sheet_name.get()
            if name not in s.probe.book.sheetnames:
                raise ValueError('指定シートが存在しません')
            s.sheet_name = name
            self.model.view()
            self.sync_conditions()
            self.render()
        except (ValueError, TypeError) as error:
            s.sheet_name = previous
            self.sheet_name.set(previous)
            messagebox.showerror('シート切替中止', str(error), parent=self.root)

    def render(self):
        headers, body, self.mapping, self.view_token = self.model.view()
        s = self.model.session
        self.sheet.set_sheet_data(body, delete_options=True)
        self.sheet.headers(headers)
        self.sheet.row_index([str(row) for row in self.mapping])
        self.sheet.set_column_widths(s.column_widths())
        heights = s.row_heights()
        self.sheet.set_row_heights([max(self.sheet.MT.min_row_height, heights[row - 1]) for row in self.mapping])
        self.sheet.deselect('all')
        self.sheet.reset_undos()
        self.sheet_name.set(s.sheet_name)
        self.sort_button.configure(state='disabled' if s.probe.read_only or s.sheet_name in s.filters else 'normal')
        self.buttons['save'].configure(state='disabled' if s.probe.read_only else 'normal')
        self.buttons['undo'].configure(state='normal' if s.undo_stack and not s.probe.read_only else 'disabled')
        self.buttons['redo'].configure(state='normal' if s.redo_stack and not s.probe.read_only else 'disabled')
        self.status.set(f'{s.probe.path.name} / {s.sheet_name} / {"読み取り専用" if s.probe.read_only else "編集可"} / '
                        f'表示 {len(self.mapping)}行 / 未保存: {"あり" if s.dirty else "なし"} / '
                        f'Undo {len(s.undo_stack)}件 / Redo {len(s.redo_stack)}件 / 保存は全ブック')
        self.selection_detail()

    def selected(self):
        cell = self.sheet.get_currently_selected()
        if not cell:
            raise ValueError('表示中の本文セルを選択してください')
        self.model.data_row(cell.row, self.view_token)
        return cell.row, cell.column, self.view_token

    def selection_detail(self, event=None):
        editable = False
        try:
            if not self.model:
                raise ValueError('XLSXを開いてください')
            row, column, token = self.selected()
            target = self.model.data_row(row, token)
            s = self.model.session
            cell = s.sheet.cell(target, column + 1)
            reason = protection_reason(s.probe, s.sheet, cell)
            editable = not reason
            self.detail.set(f'表示本文{row + 1}行 → {cell.coordinate} / 型={cell.data_type} / {reason or "編集可"}')
        except (ValueError, IndexError):
            self.detail.set('表示中の本文セルを選択してください')
        for name in ('edit', 'paste'):
            self.buttons[name].configure(state='normal' if editable else 'disabled')

    def edit(self):
        def action():
            model = self.model
            row, column, token = self.selected()
            target = model.data_row(row, token)
            cell = model.session.sheet.cell(target, column + 1)
            reason = protection_reason(model.session.probe, model.session.sheet, cell)
            if reason:
                raise PermissionError(reason)
            value = simpledialog.askstring('セル編集', f'{cell.coordinate}（書式を外した元の値）',
                                          initialvalue=display_value(cell.value), parent=self.root)
            if value is not None:
                if model is not self.model:
                    raise ValueError('編集中にブックが変わりました')
                model.edit(row, column, value, token)
        self.run(action)

    def paste(self):
        def action():
            row, column, token = self.selected()
            values = list(csv.reader(StringIO(self.root.clipboard_get(), newline=''), delimiter='\t', strict=True))
            self.model.paste(row, column, values, token)
        self.run(action)

    def sort(self):
        column, descending = self.column.get(), self.descending.get()
        self.run(lambda: self.model.session.sort_rows(int(column), header=self.model.header, descending=descending))

    def filter(self):
        column, operator, query = self.column.get(), self.operator.get(), self.query.get()
        self.run(lambda: self.model.session.set_filter(int(column),
                 operator={'部分一致': 'contains', '完全一致': 'equals', '空': 'empty', '空以外': 'nonempty'}[operator],
                 query=query, header=self.model.header))

    def clear(self):
        self.run(lambda: self.model.session.clear_filter())

    def history(self, redo=False):
        def action():
            changed = self.model.session.history(redo)
            self.notice.set(('Redo' if redo else 'Undo') + (': 1操作を反映しました' if changed else ': 対象なし'))
        self.run(action)

    def save(self):
        if not self.model or self.model.session.probe.read_only:
            return
        s = self.model.session
        path = filedialog.asksaveasfilename(parent=self.root, initialdir=str(s.probe.path.parent),
                    initialfile=s.probe.path.stem + '_gui_copy.xlsx', defaultextension='.xlsx', filetypes=[('XLSX', '*.xlsx')])
        if path:
            def action():
                self.model.session.save_as(path)
                self.notice.set('非表示行・他シートを含む全ブックを保存しました')
            self.run(action)

    def close(self):
        if self.may_discard():
            if self.model:
                self.model.session.probe.book.close()
            self.root.destroy()


def main():
    root = tk.Tk()
    XlsxProjectionGui(root)
    root.mainloop()


if __name__ == '__main__':
    main()
