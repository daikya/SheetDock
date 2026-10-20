from datetime import datetime
from unittest.mock import Mock

import pytest
from types import SimpleNamespace

from sheetdock.xlsx_inline import InlineBridge
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sheetdock.xlsx_gui import XlsxSession, XlsxPane
from sheetdock.xlsx_io import snapshot


@pytest.fixture
def bridge(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Data'
    sheet.append(['文字列', '数値', '真偽値', '日時', '数式', '結合', 'エラー'])
    sheet.append(['00123', 1234.5, True, datetime(2026, 10, 20, 8, 15), '=B2*2', '結合', '#DIV/0!'])
    sheet['A3'] = 'セル内\n改行'
    sheet['B3'] = 42
    sheet['B2'].number_format = '#,##0.00'
    sheet['D2'].number_format = 'yyyy/mm/dd hh:mm'
    sheet.merge_cells('F2:F3')
    sheet.row_dimensions[5].height = 15
    sheet.column_dimensions['H'].width = 12
    book.create_sheet('Other')['A1'] = '他シート'
    book.save(tmp_path / 'inline_basic.xlsx')
    rule = DataValidation(type='list', formula1='"はい,いいえ"')
    sheet.add_data_validation(rule)
    rule.add('A2')
    book.save(tmp_path / 'inline_readonly.xlsx')
    book.close()
    session = XlsxSession(tmp_path / 'inline_basic.xlsx')
    return InlineBridge(session, Mock(), Mock())


def begin(bridge, row, column, **kwargs):
    return bridge.begin({'row': row, 'column': column, 'key': 'F2', 'value': 'display', **kwargs})


def commit(bridge, row, column, text):
    assert begin(bridge, row, column) is not None
    return bridge.validate({'row': row, 'column': column, 'value': text})


@pytest.mark.parametrize('row,column,raw', [(1, 0, '00123'), (1, 1, '1234.5'),
    (1, 2, 'TRUE'), (1, 3, '2026-10-20T08:15:00'), (4, 7, '')])
def test_begin_uses_raw_values_without_changing_model(bridge, row, column, raw):
    before = snapshot(bridge.session.probe.book)
    assert begin(bridge, row, column) == raw
    assert snapshot(bridge.session.probe.book) == before and not bridge.session.dirty


@pytest.mark.parametrize('key,value', [('x', 'x'), ('BackSpace', '')])
def test_replace_input_keeps_typed_character_or_backspace(bridge, key, value):
    assert begin(bridge, 1, 0, key=key, value=value) == value
    assert not bridge.session.dirty


@pytest.mark.parametrize('row,column,text,value,display', [
    (1, 0, '00045', '00045', '00045'), (1, 1, '2500.25', 2500.25, '2,500.25'),
    (1, 2, 'false', False, 'FALSE'),
    (1, 3, '2026-10-21T09:30:00', datetime(2026, 10, 21, 9, 30), '2026/10/21 09:30'),
    (4, 7, '00100', '00100', '00100'), (2, 0, '日本語\n改行', '日本語\n改行', '日本語\n改行')])
def test_valid_commit_preserves_type_format_history_and_roundtrip(bridge, tmp_path, row, column, text, value, display):
    session = bridge.session
    before = snapshot(session.probe.book)
    assert commit(bridge, row, column, text) == display
    assert session.sheet.cell(row + 1, column + 1).value == value
    assert len(session.undo_stack) == 1 and session.dirty
    bridge.refresh.assert_called_once()
    after = snapshot(session.probe.book)
    session.history()
    assert snapshot(session.probe.book) == before and not session.dirty
    session.history(True)
    target = tmp_path / 'copy.xlsx'
    session.save_as(target)
    assert snapshot(XlsxSession(target).probe.book) == after


@pytest.mark.parametrize('row,column', [(1, 4), (1, 5), (2, 5), (1, 6)])
def test_formula_merge_and_error_cells_cannot_start_edit(bridge, row, column):
    before = snapshot(bridge.session.probe.book)
    assert begin(bridge, row, column) is None
    bridge.notice.assert_called_once()
    assert snapshot(bridge.session.probe.book) == before and not bridge.session.undo_stack


def test_readonly_blocks_even_plain_text_cell(bridge):
    bridge.session = XlsxSession(bridge.session.probe.path.with_name('inline_readonly.xlsx'))
    assert begin(bridge, 1, 0) is None
    assert not bridge.session.undo_stack and not bridge.session.dirty


@pytest.mark.parametrize('row,column,text', [(1, 1, 'not a number'), (1, 2, 'maybe'),
    (1, 3, 'not a date'), (1, 0, '=1+1'), (1, 0, 'bad\x00text')])
def test_invalid_input_is_rejected_by_both_model_and_widget(bridge, row, column, text):
    before = snapshot(bridge.session.probe.book)
    assert commit(bridge, row, column, text) is None
    assert bridge.last_error and not bridge.session.undo_stack
    assert snapshot(bridge.session.probe.book) == before and not bridge.session.dirty


def test_same_raw_value_and_escape_do_not_create_history(bridge):
    assert commit(bridge, 1, 1, '1234.5') == '1,234.50'
    assert not bridge.session.undo_stack and not bridge.session.dirty
    begin(bridge, 1, 0)
    bridge.cancel()
    assert bridge.context is None and not bridge.session.undo_stack


@pytest.mark.parametrize('change', ['sheet', 'book', 'session', 'cell', 'readonly', 'coordinates'])
def test_context_or_protection_changed_during_edit_is_rejected(bridge, tmp_path, change):
    session = bridge.session
    begin(bridge, 1, 0)
    row, column = 1, 0
    if change == 'sheet':
        session.sheet_name = 'Other'
    elif change == 'book':
        replacement = XlsxSession(session.probe.path)
        session.probe.book = replacement.probe.book
    elif change == 'session':
        bridge.session = XlsxSession(session.probe.path)
    elif change == 'cell':
        session.sheet['A2'] = 'changed elsewhere'
    elif change == 'readonly':
        session.probe.reasons.append('unsupported')
    else:
        row, column = 2, 0
    before = snapshot(bridge.session.probe.book)
    assert bridge.validate({'row': row, 'column': column, 'value': 'unsafe'}) is None
    assert bridge.last_error and snapshot(bridge.session.probe.book) == before
    assert not bridge.session.undo_stack


def test_validation_without_begin_and_outside_view_are_rejected(bridge):
    assert bridge.validate({'row': 0, 'column': 0, 'value': 'unsafe'}) is None
    assert begin(bridge, 20, 20) is None
    assert not bridge.session.undo_stack


def test_error_is_not_overwritten_by_next_begin_before_idle_notification(bridge):
    assert commit(bridge, 1, 1, 'bad') is None
    message = bridge.pending_error
    assert begin(bridge, 1, 0) is None and bridge.pending_error == message
    bridge.pending_error = None  # GUI acknowledged the queued error.
    assert commit(bridge, 1, 0, 'recovered') == 'recovered'
    assert bridge.last_error is None and len(bridge.session.undo_stack) == 1


@pytest.fixture
def pane(bridge, monkeypatch):
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane._rendering = bridge.session, False
    pane.root = SimpleNamespace(after_idle=Mock(return_value='idle'), after_cancel=Mock())
    pane.detail = SimpleNamespace(set=Mock())
    editor = SimpleNamespace(open=False, coords=(1, 0))
    pane.sheet_widget = SimpleNamespace(MT=SimpleNamespace(text_editor=editor,
        hide_text_editor_and_dropdown=Mock(side_effect=lambda: setattr(editor, 'open', False))),
        extra_bindings=Mock(), edit_validation=Mock(), bind_key_text_editor=Mock(),
        get_text_editor_value=Mock(return_value='typed'), set_cell_data=Mock(),
        close_text_editor=Mock(side_effect=lambda **kw: setattr(editor, 'open', False)))
    pane.selection_changed, pane._on_change, pane.render = Mock(), Mock(), Mock()
    pane.setup_inline()
    monkeypatch.setattr('sheetdock.xlsx_inline.dialogs.showerror', Mock())
    return pane


def start_pane(pane, row=1, column=0):
    assert pane.begin_inline_edit({'row': row, 'column': column, 'key': 'F2', 'value': ''}) is not None
    pane.sheet_widget.MT.text_editor.open = True
    pane.sheet_widget.MT.text_editor.coords = (row, column)


def test_pane_binds_validation_and_keeps_native_mutation_controlled(pane):
    validate = pane.sheet_widget.edit_validation.call_args.args[0]
    assert validate.__self__ is pane._inline
    assert pane.sheet_widget.extra_bindings.call_args.args == ('begin_edit_cell', pane.begin_inline_edit)
    start_pane(pane)
    assert pane.finish_inline_edit()
    assert pane.session.sheet['A2'].value == 'typed' and len(pane.session.undo_stack) == 1
    pane.sheet_widget.set_cell_data.assert_called_once_with(1, 0, value='typed', redraw=True)
    pane._on_change.assert_called_once()
    pane.render.assert_not_called()


def test_pending_invalid_editor_blocks_action_and_keeps_values(pane):
    start_pane(pane, 1, 1)
    pane.sheet_widget.get_text_editor_value.return_value = 'invalid number'
    before = snapshot(pane.session.probe.book)
    assert not pane.finish_inline_edit()
    assert snapshot(pane.session.probe.book) == before and not pane.session.undo_stack
    pane.sheet_widget.set_cell_data.assert_not_called()
    assert not pane.sheet_widget.MT.text_editor.open


def test_already_closed_error_still_blocks_immediate_action(pane):
    start_pane(pane, 1, 1)
    pane._inline.validate({'row': 1, 'column': 1, 'value': 'invalid'})
    pane.sheet_widget.MT.text_editor.open = False
    assert not pane.finish_inline_edit()
    assert pane.finish_inline_edit()  # A later explicit action can use the retained old value.


def test_pane_cancel_never_adds_history(pane):
    start_pane(pane)
    pane.cancel_inline_edit()
    assert pane._inline.context is None and not pane.session.undo_stack and not pane.session.dirty
    pane.sheet_widget.close_text_editor.assert_called_once_with(set_data=False)


def test_lifecycle_replacement_is_rejected_before_old_model_edit(pane):
    start_pane(pane)
    old = pane.session
    pane.session = XlsxSession(old.probe.path)
    assert not pane.finish_inline_edit()
    assert old.sheet['A2'].value == '00123' and pane.session.sheet['A2'].value == '00123'


def test_decorated_history_and_sheet_switch_stop_on_failed_input(pane):
    pane.finish_inline_edit = Mock(return_value=False)
    pane.session.history = Mock()
    pane.history()
    pane.session.history.assert_not_called()
    pane.sheet_tabs = SimpleNamespace(sync=Mock())
    pane.switch_sheet('Other')
    assert pane.session.sheet_name == 'Data'
    pane.sheet_tabs.sync.assert_called_once_with(['Data', 'Other'], 'Data')


def test_failed_input_restores_native_column_resize_without_model_change(pane):
    pane.finish_inline_edit = Mock(return_value=False)
    pane.session.resize_columns = Mock()
    pane.apply_column_widths({0: 200})
    pane.session.resize_columns.assert_not_called()
    pane.render.assert_called_once()


def test_failed_input_clears_pending_row_drag_and_restores_view(pane):
    pane.finish_inline_edit = Mock(return_value=False)
    pane._pending_row_sizes = {0: (20, 80)}
    pane.session.resize_rows = Mock()
    pane.finish_row_resize()
    assert not pane._pending_row_sizes
    pane.session.resize_rows.assert_not_called()
    pane.render.assert_called_once()

