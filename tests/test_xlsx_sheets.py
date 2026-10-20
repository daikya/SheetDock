import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from sheetdock.xlsx_gui import XlsxSession
from sheetdock.xlsx_io import ProtectedCellError


@pytest.fixture
def path(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Data'
    sheet['A1'] = '00123'
    sheet['A1'].font = Font(bold=True)
    sheet['A3'] = 'merged'
    sheet.merge_cells('A3:B3')
    sheet.column_dimensions['A'].width = 25
    book.create_sheet('Notes')['A1'] = '日本語'
    path = tmp_path / 'sheets.xlsx'
    book.save(path)
    return path


def test_add_undo_redo(path):
    session = XlsxSession(path)
    session.sheet_operation('add', '追加')
    assert session.probe.book.sheetnames == ['Data', 'Notes', '追加']
    assert session.sheet_name == '追加' and session.dirty
    session.history()
    assert session.probe.book.sheetnames == ['Data', 'Notes']
    assert not session.dirty
    session.history(True)
    assert session.sheet_name == '追加'


def test_rename_undo_redo(path):
    session = XlsxSession(path)
    session.sheet_operation('rename', '集計')
    assert session.probe.book.sheetnames == ['集計', 'Notes']
    assert session.sheet['A1'].value == '00123'
    session.history()
    assert session.sheet_name == 'Data' and not session.dirty
    session.history(True)
    assert session.sheet_name == '集計'


def test_case_only_rename(path):
    session = XlsxSession(path)
    session.sheet_operation('rename', 'data')
    assert session.probe.book.sheetnames == ['data', 'Notes']


def test_delete_restores_data_styles_merges_widths(path):
    session = XlsxSession(path)
    session.sheet_operation('delete')
    assert session.probe.book.sheetnames == ['Notes']
    session.history()
    assert session.sheet_name == 'Data' and not session.dirty
    assert session.sheet['A1'].value == '00123'
    assert session.sheet['A1'].font.bold
    assert str(next(iter(session.sheet.merged_cells.ranges))) == 'A3:B3'
    assert session.sheet.column_dimensions['A'].width == 25
    session.history(True)
    assert session.probe.book.sheetnames == ['Notes']


def test_cell_edits_and_structure_interleave(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.sheet_operation('rename', 'New')
    session.sheet_operation('add', 'Extra')
    session.edit(0, 0, 'extra value')
    for _ in range(4):
        assert session.history()
    assert session.sheet_name == 'Data'
    assert not session.dirty
    for _ in range(4):
        assert session.history(True)
    assert session.probe.book['Extra']['A1'].value == 'extra value'
    assert session.probe.book['New']['A1'].value == 'changed'


@pytest.mark.parametrize('name', ['', '  ', 'x'*32, 'a/b', 'a\\b', 'a?b',
                                'a*b', '[a]', 'a:b', "'a", "a'", 'History', 'a\x00b'])
def test_bad_names_leave_model_and_history_unchanged(path, name):
    session = XlsxSession(path)
    with pytest.raises(ValueError):
        session.sheet_operation('add', name)
    assert not session.dirty and not session.undo_stack


@pytest.mark.parametrize('operation,name', [('add', 'data'), ('rename', 'notes')])
def test_duplicate_names_rejected(path, operation, name):
    session = XlsxSession(path)
    with pytest.raises(ValueError):
        session.sheet_operation(operation, name)
    assert not session.dirty


@pytest.mark.parametrize('operation', ['rename', 'delete'])
def test_formula_in_any_sheet_blocks_destructive_structure(path, operation):
    book = load_workbook(path)
    book['Notes']['B1'] = '=Data!A1'
    book.save(path)
    session = XlsxSession(path)
    with pytest.raises(ProtectedCellError):
        session.sheet_operation(operation, 'Renamed')
    assert session.probe.book['Notes']['B1'].value == '=Data!A1'
    assert not session.dirty


def test_formula_book_allows_append_and_undo(path):
    book = load_workbook(path)
    book['Notes']['B1'] = '=Data!A1'
    book.save(path)
    session = XlsxSession(path)
    session.sheet_operation('add', 'Extra')
    assert session.probe.book['Notes']['B1'].value == '=Data!A1'
    session.history()
    assert not session.dirty


def test_last_visible_sheet_is_protected(path):
    book = load_workbook(path)
    book['Notes'].sheet_state = 'hidden'
    book.save(path)
    session = XlsxSession(path)
    with pytest.raises(ProtectedCellError, match='最後の表示シート'):
        session.sheet_operation('delete')
    session.sheet_name = 'Notes'
    session.sheet_operation('delete')
    assert session.probe.book.sheetnames == ['Data']


@pytest.mark.parametrize('operation', ['add', 'rename', 'delete'])
def test_readonly_blocks_all_structure(path, operation):
    session = XlsxSession(path)
    session.probe.reasons.append('unsupported')
    with pytest.raises(ProtectedCellError):
        session.sheet_operation(operation, 'New')


def test_structure_save_and_reopen(path, tmp_path):
    session = XlsxSession(path)
    session.sheet_operation('rename', '名前変更')
    session.sheet_operation('add', '追加')
    destination = tmp_path / 'saved.xlsx'
    session.save_as(destination)
    restored = XlsxSession(destination)
    assert restored.probe.book.sheetnames == ['名前変更', 'Notes', '追加']
    assert not restored.dirty
    session.history()
    assert session.dirty
    session.history(True)
    assert not session.dirty


def test_failed_structure_rolls_back_and_keeps_redo(path, monkeypatch):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.history()
    original = session.pack_book
    calls = []
    def fail_after_mutation():
        calls.append(True)
        if len(calls) == 2:
            raise OSError('simulated snapshot failure')
        return original()
    monkeypatch.setattr(session, 'pack_book', fail_after_mutation)
    with pytest.raises(OSError):
        session.sheet_operation('add', 'New')
    assert session.probe.book.sheetnames == ['Data', 'Notes']
    assert session.redo_stack
    assert not session.dirty


def test_delete_and_recreate_same_name_does_not_hide_lost_content(path):
    session = XlsxSession(path)
    session.sheet_operation('delete')
    session.sheet_operation('add', 'Data')
    assert session.dirty


def test_noop_rename_keeps_redo(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.history()
    assert not session.sheet_operation('rename', 'Data')
    assert session.redo_stack and not session.dirty
