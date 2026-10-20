from io import BytesIO
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.dimensions import ColumnDimension, RowDimension

from sheetdock.xlsx_gui import XlsxSession as AppendSession
from sheetdock.xlsx_extent import extent, extents
from sheetdock.xlsx_io import ProtectedCellError, snapshot
from sheetdock.xlsx_widths import stored_width
from sheetdock.xlsx_extent import data_extent
from sheetdock.xlsx_gui import XlsxPane
from types import SimpleNamespace
from unittest.mock import Mock
import csv


@pytest.fixture
def path(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Data'
    sheet.append(['00123', 3, True])
    sheet.append(['日本語\n改行', datetime(2026, 10, 20), None])
    sheet['A2'].font = Font(bold=True, color='FF2563EB')
    sheet['A2'].fill = PatternFill('solid', fgColor='FFFDE68A')
    sheet.row_dimensions[2].height = 32
    sheet.column_dimensions['A'].width = 24
    book.create_sheet('Other')['A1'] = 'unchanged'
    path = tmp_path / 'data.xlsx'
    book.save(path)
    return path


def test_plain_blank_cell_does_not_survive_save():
    book = Workbook()
    book.active['A1'] = 'data'
    book.active.cell(6, 4)
    assert (book.active.max_row, book.active.max_column) == (6, 4)
    stream = BytesIO()
    book.save(stream)
    checked = load_workbook(BytesIO(stream.getvalue()))
    try:
        assert (checked.active.max_row, checked.active.max_column) == (1, 1)
    finally:
        checked.close()


@pytest.mark.parametrize('axis,expected', [('rows', (4, 3)), ('columns', (2, 5))])
def test_append_history_and_repeated_save_keep_blank_extent(path, tmp_path, axis, expected):
    session = AppendSession(path)
    original = path.read_bytes()
    before, sizes = snapshot(session.probe.book), extents(session.probe.book)
    session.append_axis(axis, 2)
    after = snapshot(session.probe.book)
    assert extent(session.sheet) == expected and session.dirty
    # Values, types, cell formatting and the other sheet are all unchanged.
    assert after[0][5] == before[0][5] and after[1] == before[1]
    assert len(session.undo_stack) == 1
    session.history()
    assert snapshot(session.probe.book) == before and extents(session.probe.book) == sizes
    assert not session.dirty
    session.history(True)
    output = tmp_path / 'copy.xlsx'
    session.save_as(output)
    assert not session.dirty and path.read_bytes() == original
    reopened = AppendSession(output)
    assert extent(reopened.sheet) == expected
    assert snapshot(reopened.probe.book) == after
    # A second ordinary openpyxl save must retain dimension-based extents too.
    stream = BytesIO()
    reopened.probe.book.save(stream)
    checked = load_workbook(BytesIO(stream.getvalue()))
    assert extent(checked['Data']) == expected
    assert (checked['Data'].max_row, checked['Data'].max_column) == (2, 3)
    checked.close()
    session.history()
    assert session.dirty
    session.history(True)
    assert not session.dirty


def test_repeated_append_uses_logical_tail_and_mixed_history(path, tmp_path):
    session = AppendSession(path)
    baseline = snapshot(session.probe.book)
    session.append_axis('rows', 2)
    session.append_axis('columns', 2)
    session.append_axis('rows', 3)
    assert extent(session.sheet) == (7, 5)
    for expected in [(4, 5), (4, 3), (2, 3)]:
        session.history()
        assert extent(session.sheet) == expected
    assert snapshot(session.probe.book) == baseline and not session.dirty
    for _ in range(3):
        session.history(True)
    session.save_as(tmp_path / 'first.xlsx')
    reopened = AppendSession(tmp_path / 'first.xlsx')
    reopened.append_axis('columns', 2)
    assert extent(reopened.sheet) == (7, 7)
    # Blank cells can later receive data with an ordinary workbook editor.
    reopened.probe.set_value('Data', 'G7', 'new value')
    reopened.save_as(tmp_path / 'edited.xlsx')
    checked = AppendSession(tmp_path / 'edited.xlsx')
    assert checked.sheet['G7'].value == 'new value' and extent(checked.sheet) == (7, 7)


@pytest.mark.parametrize('axis,expected', [('rows', (3, 1)), ('columns', (1, 3))])
def test_empty_sheet_baseline_is_one_cell(tmp_path, axis, expected):
    path = tmp_path / 'empty.xlsx'
    Workbook().save(path)
    session = AppendSession(path)
    session.append_axis(axis, 2)
    session.save_as(tmp_path / 'empty_copy.xlsx')
    assert extent(AppendSession(tmp_path / 'empty_copy.xlsx').sheet) == expected


@pytest.mark.parametrize('axis,default', [('rows', 22.5), ('columns', 18.25)])
def test_new_dimensions_use_sheet_defaults_not_last_cell_style(path, axis, default):
    session = AppendSession(path)
    if axis == 'rows':
        session.sheet.sheet_format.defaultRowHeight = default
    else:
        session.sheet.sheet_format.defaultColWidth = default
    session.append_axis(axis, 2)
    if axis == 'rows':
        assert [session.sheet.row_dimensions[i].height for i in (3, 4)] == [default] * 2
    else:
        assert [session.sheet.column_dimensions[c].width for c in ('D', 'E')] == [default] * 2
    assert all(not cell.has_style for row in session.sheet for cell in row if cell.value is None)


def test_existing_blank_dimension_extents_are_respected(path):
    session = AppendSession(path)
    session.sheet.row_dimensions[6] = RowDimension(session.sheet, index=6, height=25)
    session.sheet.column_dimensions['G'] = ColumnDimension(session.sheet, index='G', min=7, max=7, width=20)
    session.append_axis('rows', 2)
    session.append_axis('columns', 1)
    assert extent(session.sheet) == (8, 8)
    assert session.sheet.row_dimensions[6].height == 25
    assert session.sheet.column_dimensions['G'].width == 20


def test_default_column_width_matches_existing_display_conversion(path):
    session = AppendSession(path)
    expected = stored_width(session.sheet, 4)
    session.append_axis('columns', 2)
    assert session.sheet.column_dimensions['D'].width == expected
    assert session.sheet.column_dimensions['E'].width == expected


@pytest.mark.parametrize('count', [True, 0, -1, 101, '2', 1.5])
def test_invalid_count_is_atomic(path, count):
    session = AppendSession(path)
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError):
        session.append_axis('rows', count)
    assert snapshot(session.probe.book) == before and not session.undo_stack and not session.dirty


@pytest.mark.parametrize('restriction', ['formula', 'merged', 'readonly', 'hidden', 'grouped', 'default_hidden'])
def test_existing_structure_restrictions_are_preserved(path, restriction):
    session = AppendSession(path)
    if restriction == 'formula':
        session.probe.book['Other']['A2'] = '=Data!B1'
    elif restriction == 'merged':
        session.sheet.merge_cells('A1:B1')
    elif restriction == 'readonly':
        session.probe.reasons.append('unsupported')
    elif restriction == 'hidden':
        session.sheet.row_dimensions[2].hidden = True
    elif restriction == 'grouped':
        session.sheet.row_dimensions[2].outlineLevel = 1
    else:
        session.sheet.sheet_format.zeroHeight = True
    before, sizes = snapshot(session.probe.book), extents(session.probe.book)
    with pytest.raises(ProtectedCellError):
        session.append_axis('rows', 2)
    assert snapshot(session.probe.book) == before and extents(session.probe.book) == sizes
    assert not session.undo_stack


@pytest.mark.parametrize('axis', ['rows', 'columns'])
def test_extent_limits_include_blank_dimensions(path, axis):
    session = AppendSession(path)
    if axis == 'rows':
        session.sheet.row_dimensions[1000] = RowDimension(session.sheet, index=1000, height=15)
    else:
        session.sheet.column_dimensions['AX'] = ColumnDimension(session.sheet, index='AX', min=50, max=50, width=13)
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError, match='1000行'):
        session.append_axis(axis)
    assert snapshot(session.probe.book) == before and not session.undo_stack


def test_book_area_counts_other_sheet_dimensions(path):
    session = AppendSession(path)
    other = session.probe.book['Other']
    other.row_dimensions[1000] = RowDimension(other, index=1000, height=15)
    other.column_dimensions['CV'] = ColumnDimension(other, index='CV', min=100, max=100, width=13)
    with pytest.raises(ValueError, match='10万'):
        session.append_axis('rows')
    assert not session.undo_stack


def test_verification_failure_does_not_publish_candidate(path, monkeypatch):
    session = AppendSession(path)
    before = snapshot(session.probe.book)
    monkeypatch.setattr('sheetdock.xlsx_axis.inspect_package', lambda data: ('unsupported',))
    with pytest.raises(ValueError, match='未対応'):
        session.append_axis('columns', 2)
    assert snapshot(session.probe.book) == before and not session.dirty and not session.undo_stack


def test_save_failure_keeps_original_edit_and_history(path, tmp_path, monkeypatch):
    session = AppendSession(path)
    original = path.read_bytes()
    session.append_axis('rows', 2)
    def fail(*args, **kwargs):
        raise OSError('simulated failure')
    monkeypatch.setattr('sheetdock.xlsx_gui.save_checked', fail)
    with pytest.raises(OSError):
        session.save_as(tmp_path / 'copy.xlsx')
    assert path.read_bytes() == original and session.dirty and len(session.undo_stack) == 1
    assert extent(session.sheet) == (4, 3)
    session.history()
    assert not session.dirty


def test_new_append_clears_redo(path):
    session = AppendSession(path)
    session.append_axis('rows', 2)
    session.history()
    session.append_axis('columns', 2)
    assert not session.history(True) and extent(session.sheet) == (2, 5)


def test_reading_blank_cells_does_not_expand_edit_bounds_export_or_dirty(path, tmp_path):
    session = AppendSession(path)
    before = snapshot(session.probe.book)
    session.sheet.cell(900, 49)
    assert extent(session.sheet) == (2, 3) and data_extent(session.sheet) == (2, 3)
    assert len(session.row_heights()) == 2 and len(session.column_widths()) == 3
    assert not session.dirty and snapshot(session.probe.book) == before
    with pytest.raises(ValueError):
        session.edit(4, 4, 'outside')
    session.export_csv(tmp_path / 'data.csv')
    with (tmp_path / 'data.csv').open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.reader(stream))
    assert len(rows) == 2 and all(len(row) == 3 for row in rows)


def test_new_tail_edit_paste_and_csv_export_survive_reopen(path, tmp_path):
    session = AppendSession(path)
    session.append_axis('rows', 2)
    session.append_axis('columns', 2)
    # Simulate the GUI allocating all empty cells of its displayed rectangle.
    for row in session.sheet.iter_rows(min_row=1, max_row=4, min_col=1, max_col=5):
        pass
    session.export_csv(tmp_path / 'blank_tail.csv')
    with (tmp_path / 'blank_tail.csv').open(encoding='utf-8-sig', newline='') as stream:
        assert len(list(csv.reader(stream))) == 2
    session.edit(3, 4, '00123')
    session.paste(2, 3, 'left\tright\n下段\t"cell\nline"')
    session.save_as(tmp_path / 'tail.xlsx')
    checked = AppendSession(tmp_path / 'tail.xlsx')
    assert extent(checked.sheet) == (4, 5)
    assert checked.sheet['E3'].value == 'right' and checked.sheet['E3'].data_type == 's'
    assert checked.sheet['E4'].value == 'cell\nline'
    checked.export_csv(tmp_path / 'edited_tail.csv')
    with (tmp_path / 'edited_tail.csv').open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.reader(stream))
    assert len(rows) == 4 and all(len(row) == 5 for row in rows)
    assert rows[3][4] == 'cell\nline'


def test_new_blank_cell_preserves_zeroes_and_atomic_paste_boundary(path, tmp_path):
    session = AppendSession(path)
    session.append_axis('rows', 1)
    session.edit(2, 0, '00045')
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError):
        session.paste(2, 2, 'inside\toutside')
    assert snapshot(session.probe.book) == before
    session.save_as(tmp_path / 'zeros.xlsx')
    checked = AppendSession(tmp_path / 'zeros.xlsx')
    assert checked.sheet['A3'].value == '00045' and checked.sheet['A3'].data_type == 's'


@pytest.mark.parametrize('action,expected', [
    ('insert-rows', (5, 5)), ('delete-rows', (3, 5)),
    ('insert-columns', (4, 6)), ('delete-columns', (4, 4)),
])
def test_structure_changes_preserve_tail_extents_and_history(path, tmp_path, action, expected):
    session = AppendSession(path)
    session.append_axis('rows', 2)
    session.append_axis('columns', 2)
    before = snapshot(session.probe.book)
    # Materialize blanks exactly as rendering does, without making them data.
    list(session.sheet.iter_rows(min_row=1, max_row=4, min_col=1, max_col=5))
    assert session.axis_operation(action, 3, 1)
    assert extent(session.sheet) == expected
    session.history()
    assert extent(session.sheet) == (4, 5) and snapshot(session.probe.book) == before
    session.history(True)
    session.save_as(tmp_path / 'structure.xlsx')
    assert extent(AppendSession(tmp_path / 'structure.xlsx').sheet) == expected


def test_size_edit_search_and_mixed_history_on_tail(path, tmp_path):
    session = AppendSession(path)
    before = snapshot(session.probe.book)
    session.append_axis('rows', 2)
    session.append_axis('columns', 2)
    session.resize_rows({3: 80})
    session.resize_columns({4: 200})
    session.edit(3, 4, 'tail word')
    assert session.find('tail', all_sheets=True) == [('Data', 3, 4, 'tail word')]
    after = snapshot(session.probe.book)
    for _ in range(5):
        assert session.history()
    assert snapshot(session.probe.book) == before and extent(session.sheet) == (2, 3) and not session.dirty
    for _ in range(5):
        assert session.history(True)
    assert snapshot(session.probe.book) == after and extent(session.sheet) == (4, 5)
    session.save_as(tmp_path / 'mixed.xlsx')
    checked = AppendSession(tmp_path / 'mixed.xlsx')
    assert checked.row_heights()[3] == 80 and checked.column_widths()[4] == 200


@pytest.fixture
def pane(path, monkeypatch):
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane.root, pane._digit_width = AppendSession(path), None, 9
    pane.sheet_widget = SimpleNamespace(select_cell=Mock(), see=Mock())
    pane.render, pane.selection_changed = Mock(), Mock()
    pane.ask, pane.info, pane.error = Mock(return_value=2), Mock(), Mock()
    for name, mock in [('askinteger', pane.ask), ('showinfo', pane.info), ('showerror', pane.error)]:
        monkeypatch.setattr('sheetdock.dialogs.' + name, mock)
    return pane


@pytest.mark.parametrize('axis,cell,expected', [('rows', (2, 0), (4, 3)), ('columns', (0, 3), (2, 5))])
def test_gui_append_selects_first_added_cell_and_records_once(pane, axis, cell, expected):
    pane.append_axis(axis)
    pane.error.assert_not_called()
    pane.render.assert_called_once()
    pane.sheet_widget.select_cell.assert_called_once_with(*cell)
    assert extent(pane.session.sheet) == expected and len(pane.session.undo_stack) == 1
    if axis == 'columns':
        assert pane.session.sheet.column_dimensions['D'].width == stored_width(Workbook().active, 4, 9)


def test_gui_append_cancel_is_unchanged(pane):
    pane.ask.return_value = None
    pane.append_axis('rows')
    pane.render.assert_not_called()
    assert not pane.session.dirty and not pane.session.undo_stack


def test_gui_append_rechecks_protection_and_bounds(pane):
    pane.session.probe.reasons.append('unsupported')
    pane.append_axis('rows')
    pane.ask.assert_not_called()
    assert not pane.session.undo_stack
    pane.session.probe.reasons.clear()
    pane.session.sheet.row_dimensions[1000] = RowDimension(pane.session.sheet, index=1000, height=15)
    pane.append_axis('rows')
    pane.ask.assert_not_called()
    assert not pane.session.undo_stack


def test_append_buttons_work_without_cell_selection_and_respect_readonly(pane):
    pane.selected_cell = lambda: None
    pane._axis_block_reason = ''
    pane.append_buttons = {'rows': Mock(), 'columns': Mock()}
    pane.update_axis_commands()
    for button in pane.append_buttons.values():
        button.configure.assert_called_with(state='normal')
    pane.session.probe.reasons.append('unsupported')
    pane.update_axis_commands()
    for button in pane.append_buttons.values():
        button.configure.assert_called_with(state='disabled')

