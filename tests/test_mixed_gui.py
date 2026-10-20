from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from openpyxl import Workbook, load_workbook

from sheetdock.csv_io import CsvDocument, CsvFormat, save_csv
from sheetdock.gui import Session
from sheetdock.mixed_gui import SheetDockApp, is_xlsx
from sheetdock.tabs import Documents
from sheetdock.xlsx_gui import XlsxSession, XlsxPane


@pytest.fixture
def app(tmp_path, monkeypatch):
    book = Workbook()
    book.active['A1'] = 'original'
    path = tmp_path / 'workbook.xlsx'
    book.save(path)
    csv_path = tmp_path / 'data.csv'
    document = CsvDocument([['00123']], CsvFormat('utf-8', False, '\n', ','))
    save_csv(document, csv_path)
    app = SheetDockApp.__new__(SheetDockApp)
    app.root = None
    app.documents = Documents()
    csv_tab = app.documents.add(csv_path, Session(document, False))
    xlsx_tab = app.documents.add(path, XlsxSession(path))
    app.active = csv_tab
    app.session = csv_tab.session
    app.source = csv_path
    app.pages = {}
    app.backup_enabled = SimpleNamespace(get=lambda: True)
    app.status = SimpleNamespace(set=lambda text: None)
    app.labels = lambda: None
    app.remember_file = lambda path: None
    app.pages[xlsx_tab] = (None, SimpleNamespace(session=xlsx_tab.session,
        render=lambda: app.sync_xlsx(xlsx_tab)))
    app.snapshot = lambda: csv_tab.session.snapshot()
    monkeypatch.setattr('sheetdock.mixed_gui.messagebox.showerror', lambda *a, **kw: None)
    return app, csv_tab, xlsx_tab


def test_mixed_titles_and_duplicate_detection(app):
    app, csv_tab, xlsx_tab = app
    csv_tab.session.editor.set_cell(0, 0, '00999')
    xlsx_tab.session.edit(0, 0, 'changed')
    assert csv_tab.title.endswith(' *') and xlsx_tab.title.endswith(' *')
    assert app.documents.find(xlsx_tab.source) is xlsx_tab
    assert not is_xlsx(csv_tab) and is_xlsx(xlsx_tab)


def test_inactive_xlsx_save_does_not_change_active_csv(app):
    app, csv_tab, xlsx_tab = app
    csv_tab.session.editor.set_cell(0, 0, '00999')
    xlsx_tab.session.edit(0, 0, 'changed')
    assert app.save_tab(xlsx_tab)
    assert load_workbook(xlsx_tab.source).active['A1'].value == 'changed'
    assert not xlsx_tab.session.dirty
    assert csv_tab.session.editor.dirty
    assert app.source == csv_tab.source
    assert app.session is csv_tab.session


def test_xlsx_save_as_updates_one_tab_only(app, tmp_path, monkeypatch):
    app, csv_tab, xlsx_tab = app
    xlsx_tab.session.edit(0, 0, 'changed')
    destination = tmp_path / 'new.xlsx'
    monkeypatch.setattr('sheetdock.mixed_gui.filedialog.asksaveasfilename', lambda **kw: str(destination))
    assert app.save_as_tab(xlsx_tab)
    assert xlsx_tab.source == destination
    assert not xlsx_tab.session.dirty
    assert csv_tab.source.name == 'data.csv'


def test_save_as_cannot_target_another_open_tab(app, monkeypatch):
    app, csv_tab, xlsx_tab = app
    before = csv_tab.source.read_bytes()
    original_path = xlsx_tab.source
    monkeypatch.setattr('sheetdock.mixed_gui.filedialog.asksaveasfilename', lambda **kw: str(csv_tab.source))
    assert not app.save_as_tab(xlsx_tab)
    assert csv_tab.source.read_bytes() == before
    assert xlsx_tab.source == original_path


def test_csv_cannot_save_text_with_xlsx_extension(app, tmp_path, monkeypatch):
    app, csv_tab, _ = app
    target = tmp_path / 'wrong.xlsx'
    monkeypatch.setattr('sheetdock.mixed_gui.filedialog.asksaveasfilename', lambda **kw: str(target))
    assert not app.save_as_tab(csv_tab)
    assert not target.exists()


def test_export_target_cannot_be_open_deleted_csv(app):
    app, csv_tab, _ = app
    csv_tab.source.unlink()
    with pytest.raises(ValueError, match='別のタブ'):
        app.check_export_target(csv_tab.source)


def test_open_backup_target_blocks_xlsx_save(app):
    app, csv_tab, xlsx_tab = app
    backup = Path(str(xlsx_tab.source) + '.bak')
    backup.write_bytes(b'opened backup')
    app.documents.add(backup, csv_tab.session)
    before = xlsx_tab.source.read_bytes()
    xlsx_tab.session.edit(0, 0, 'changed')
    assert not app.save_tab(xlsx_tab)
    assert xlsx_tab.source.read_bytes() == before
    assert backup.read_bytes() == b'opened backup'
    assert xlsx_tab.session.dirty


def test_external_cancel_keeps_both_sessions(app, monkeypatch):
    app, csv_tab, xlsx_tab = app
    xlsx_tab.session.edit(0, 0, 'local')
    book = load_workbook(xlsx_tab.source)
    book.active['A1'] = 'external'
    book.save(xlsx_tab.source)
    monkeypatch.setattr('sheetdock.mixed_gui.messagebox.askyesnocancel', lambda *a, **kw: None)
    assert not app.save_tab(xlsx_tab)
    assert xlsx_tab.session.probe.book.active['A1'].value == 'local'
    assert xlsx_tab.session.dirty
    assert app.active is csv_tab


def test_inactive_reload_replaces_only_xlsx_session(app, monkeypatch):
    app, csv_tab, xlsx_tab = app
    original = xlsx_tab.session
    xlsx_tab.session.edit(0, 0, 'local')
    monkeypatch.setattr('sheetdock.mixed_gui.messagebox.askyesno', lambda *a, **kw: True)
    assert not app.reload_tab(xlsx_tab)
    assert xlsx_tab.session is not original
    assert not xlsx_tab.session.dirty
    assert app.session is csv_tab.session
    assert xlsx_tab.session.stamp == xlsx_tab.stamp


def test_readonly_save_never_opens_dialog_or_changes_file(app):
    app, _, xlsx_tab = app
    xlsx_tab.session.probe.reasons.append('unsupported')
    before = xlsx_tab.source.read_bytes()
    assert not app.save_tab(xlsx_tab)
    assert not app.save_as_tab(xlsx_tab)
    assert xlsx_tab.source.read_bytes() == before


def test_mixed_close_cancel_preserves_all_tabs(app):
    app, csv_tab, xlsx_tab = app
    csv_tab.session.editor.set_cell(0, 0, 'changed')
    xlsx_tab.session.edit(0, 0, 'changed')
    destroyed = []
    app.root = SimpleNamespace(destroy=lambda: destroyed.append(True))
    app.decide_close = lambda tab: 'discard' if tab is csv_tab else 'cancel'
    app.close()
    assert not destroyed
    assert app.documents.tabs == [csv_tab, xlsx_tab]


def test_xlsx_may_close_saves_before_closing(app):
    app, _, xlsx_tab = app
    xlsx_tab.session.edit(0, 0, 'saved on close')
    assert app.documents.may_close(xlsx_tab, lambda tab: 'save', app.save_tab)
    assert not xlsx_tab.session.dirty
    assert load_workbook(xlsx_tab.source).active['A1'].value == 'saved on close'


def test_duplicate_open_activates_existing_tab(app):
    app, _, xlsx_tab = app
    activated = []
    app.activate = lambda tab: activated.append(tab)
    app.notebook = SimpleNamespace(select=lambda page: None)
    app.open_path(xlsx_tab.source)
    app.open_path(xlsx_tab.source)
    assert activated == [xlsx_tab, xlsx_tab]
    assert len(app.documents.tabs) == 2


def test_xlsm_is_rejected_before_csv_parsing(app, tmp_path):
    app, _, _ = app
    app.open_path(tmp_path / 'macro.xlsm')
    assert len(app.documents.tabs) == 2


def test_pending_xlsx_error_prevents_save_as_and_reload(app, monkeypatch):
    app, _, tab = app
    pane = app.pane(tab)
    pane.finish_inline_edit = lambda: False
    original = tab.source.read_bytes()
    ask = Mock()
    monkeypatch.setattr('sheetdock.mixed_gui.filedialog.asksaveasfilename', ask)
    assert not app.save_tab(tab) and not app.save_as_tab(tab) and not app.reload_tab(tab)
    ask.assert_not_called()
    assert tab.source.read_bytes() == original


def test_pending_error_restores_original_file_tab(app):
    app, csv_tab, tab = app
    app.active = tab
    app.pane(tab).finish_inline_edit = lambda: False
    app.notebook = SimpleNamespace(select=Mock())
    app.switching = False
    app.activate(csv_tab)
    assert app.active is tab and not app.switching
    app.notebook.select.assert_called_once_with(app.pages[tab][0])


def test_pending_error_prevents_open_find_and_close(app, monkeypatch):
    app, _, tab = app
    app.active = tab
    app.pane(tab).finish_inline_edit = lambda: False
    ask = Mock()
    monkeypatch.setattr('sheetdock.mixed_gui.filedialog.askopenfilename', ask)
    app.open()
    app.open_path(tab.source)
    app.find()
    app.close_tab()
    app.close()
    ask.assert_not_called()
    assert app.active is tab and len(app.documents.tabs) == 2


def test_pane_save_controls_delegate_to_host():
    pane = XlsxPane.__new__(XlsxPane)
    calls = []
    pane._save_handler = lambda: calls.append('save')
    pane._save_as_handler = lambda: calls.append('save-as')
    assert pane.save() == 'break'
    pane.save_as()
    assert calls == ['save', 'save-as']


def test_xlsx_history_does_not_touch_csv_history(app):
    app, csv_tab, xlsx_tab = app
    csv_tab.session.editor.set_cell(0, 0, 'changed')
    xlsx_tab.session.edit(0, 0, 'changed')
    app.active = xlsx_tab
    app.pages[xlsx_tab][1].history = lambda redo: xlsx_tab.session.history(redo)
    app.history(False)
    assert not xlsx_tab.session.dirty
    assert csv_tab.session.editor.dirty


def test_runtime_modules_do_not_import_probes():
    import ast
    root = Path(__file__).resolve().parents[1] / 'sheetdock'
    for path in root.glob('*.py'):
        for node in ast.walk(ast.parse(path.read_text(encoding='utf-8'))):
            if isinstance(node, ast.ImportFrom):
                assert not (node.module or '').startswith('probes')


def test_runtime_licence_spelling_is_collected():
    import build_release
    for name in ('openpyxl', 'et-xmlfile'):
        assert build_release.package_license_paths(name)


class Control:
    def __init__(self):
        self.state = None
        self.visible = False

    def configure(self, **kwargs):
        self.state = kwargs.get('state', self.state)

    def grid(self):
        self.visible = True

    def grid_remove(self):
        self.visible = False


@pytest.mark.parametrize('mode', ['csv', 'xlsx', 'readonly', 'protected', 'no-selection', 'empty'])
def test_shared_controls_follow_active_file_and_selection(app, mode):
    app, csv_tab, xlsx_tab = app
    app.layout_ready = True
    app.command_buttons = {name: Control() for name in ['undo', 'redo', 'find', 'edit', 'paste', 'export']}
    for name in ['save_current_button', 'save_button', 'close_button', 'csv_options',
                 'csv_encoding_actions', 'row_column_actions', 'guidance_label']:
        setattr(app, name, Control())
    if mode == 'csv':
        app.active = csv_tab
    elif mode == 'empty':
        app.active = None
    else:
        app.active = xlsx_tab
        pane = app.pane()
        cell = xlsx_tab.session.sheet['A1']
        if mode == 'readonly':
            xlsx_tab.session.probe.reasons.append('unsupported')
        if mode == 'protected':
            cell.value = '=1+1'
        pane.selected_cell = lambda: None if mode == 'no-selection' else cell
    app.update_commands()
    assert app.command_buttons['export'].state == ('normal' if mode not in {'csv', 'empty'} else 'disabled')
    assert app.command_buttons['edit'].state == ('normal' if mode == 'xlsx' else 'disabled')
    assert app.command_buttons['paste'].state == app.command_buttons['edit'].state
    assert app.save_button.state == ('disabled' if mode in {'readonly', 'empty'} else 'normal')
    assert app.row_column_actions.visible == (mode == 'csv')
    assert app.csv_options.visible == (mode in {'csv', 'empty'})


@pytest.mark.parametrize('dirty,label', [(False, 'なし'), (True, 'あり')])
def test_xlsx_status_uses_same_dirty_labels_as_csv(app, dirty, label):
    app, csv_tab, xlsx_tab = app
    app.active = xlsx_tab
    if dirty:
        xlsx_tab.session.edit(0, 0, 'changed')
    texts = []
    app.status = SimpleNamespace(set=texts.append)
    app.format_status()
    assert texts[-1].endswith(f'未保存={label}')
    assert 'True' not in texts[-1] and 'False' not in texts[-1]
