from dataclasses import replace
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from sheetdock.xlsx_gui import XlsxSession, XlsxPane
from sheetdock.xlsx_io import ProtectedCellError, snapshot
from sheetdock.xlsx_search import Match, XlsxSearch


@pytest.fixture
def session(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'One'
    sheet['A1'] = 'cat cat'
    sheet['B1'] = 123
    sheet['C1'] = 'cat merged'
    sheet.merge_cells('C1:D1')
    sheet['A2'] = 'cat'
    sheet['B2'] = '=1+1'
    sheet['A4'] = 'cat'
    book.create_sheet('Two')['A1'] = 'cat'
    path = tmp_path / 'search.xlsx'
    book.save(path)
    return XlsxSession(path)


def test_forward_backward_and_wrap_count_occurrences(session):
    search = XlsxSearch()
    first = search.move(session, 'cat')
    assert (first.sheet, first.row, first.start) == ('One', 0, 0)
    assert len(search.matches) == 6
    second = search.move(session, 'cat')
    assert second.start == 4 and second.column == 0
    assert search.move(session, 'cat', backward=True) == first
    assert search.move(session, 'cat', backward=True).sheet == 'Two'
    assert search.move(session, 'cat') == first
    assert not session.dirty and not session.undo_stack


def test_initial_backward_starts_at_last_match(session):
    assert XlsxSearch().move(session, 'cat', backward=True).sheet == 'Two'


def test_search_current_sheet_and_changed_scope(session):
    search = XlsxSearch()
    search.move(session, 'cat', all_sheets=False)
    assert len(search.matches) == 5
    session.sheet_name = 'Two'
    assert search.move(session, 'cat', all_sheets=False).sheet == 'Two'
    assert len(search.matches) == 1


def test_formula_search_is_explicit_and_not_editable(session):
    search = XlsxSearch()
    assert search.move(session, '=1+1') is None
    match = search.move(session, '=1+1', include_formulas=True)
    with pytest.raises(ProtectedCellError):
        session.replace_match(match, '=1+1', 'changed')
    assert session.sheet['B2'].value == '=1+1'
    assert not session.dirty


def test_replace_exact_occurrence_and_history_origin(session):
    search = XlsxSearch()
    search.move(session, 'cat')
    match = search.move(session, 'cat')
    session.sheet_name = 'Two'
    before = snapshot(session.probe.book)
    assert session.replace_match(match, 'cat', 'dog')
    assert session.sheet_name == 'One'
    assert session.sheet['A1'].value == 'cat dog'
    assert len(session.undo_stack) == 1
    session.sheet_name = 'Two'
    session.history()
    assert session.sheet_name == 'One'
    assert snapshot(session.probe.book) == before and not session.dirty
    session.history(True)
    assert session.sheet['A1'].value == 'cat dog'


def test_resume_does_not_immediately_search_inserted_text(session):
    search = XlsxSearch()
    match = search.move(session, 'cat')
    session.replace_match(match, 'cat', 'catcat')
    next_match = search.after_replace(session, 'cat', match, 'catcat')
    assert next_match.value == 'catcat cat'
    assert next_match.start == 7
    assert search.current(session, 'cat') == next_match


def test_deletion_resumes_at_shifted_next_occurrence(session):
    search = XlsxSearch()
    match = search.move(session, 'cat')
    session.replace_match(match, 'cat', '')
    next_match = search.after_replace(session, 'cat', match, '')
    assert next_match.start == 1 and next_match.value == ' cat'


def test_replace_only_match_to_blank_roundtrips_and_keeps_style(session, tmp_path):
    session.sheet['A2'].number_format = '@'
    match = Match('One', 1, 0, 0, 3, 'cat')
    session.replace_match(match, 'cat', '')
    assert session.sheet['A2'].value is None
    assert session.sheet['A2'].number_format == '@'
    path = tmp_path / 'blank.xlsx'
    session.save_as(path)
    reopened = load_workbook(path)
    assert reopened['One']['A2'].value is None
    assert reopened['One']['A2'].number_format == '@'


@pytest.mark.parametrize('change', ['value', 'query', 'scope', 'formulas', 'book'])
def test_old_search_result_requires_new_search(session, change):
    search = XlsxSearch()
    search.move(session, 'cat')
    query, options = 'cat', {}
    if change == 'value':
        session.edit(0, 0, 'changed cat')
    elif change == 'query':
        query = 'ca'
    elif change == 'scope':
        options['all_sheets'] = False
    elif change == 'formulas':
        options['include_formulas'] = True
    else:
        session.sheet_operation('add', 'New')
        session.history()  # Structural Undo replaces the workbook object.
    with pytest.raises(ValueError, match='もう一度検索'):
        search.current(session, query, **options)


def test_replace_without_search_is_rejected(session):
    with pytest.raises(ValueError):
        XlsxSearch().current(session, 'cat')


@pytest.mark.parametrize('match,query', [
    (Match('One', 0, 2, 0, 3, 'cat merged'), 'cat'),
    (Match('One', 0, 1, 0, 1, '123'), '1'),
])
def test_merged_and_numeric_matches_are_protected(session, match, query):
    with pytest.raises(ProtectedCellError):
        session.replace_match(match, query, 'x')
    assert not session.dirty and not session.undo_stack


def test_readonly_rejects_single_replace_but_allows_search(session):
    session.probe.reasons.append('unsupported feature')
    match = XlsxSearch().move(session, 'cat')
    with pytest.raises(ProtectedCellError):
        session.replace_match(match, 'cat', 'x')
    assert not session.dirty


@pytest.mark.parametrize('replacement', ['=1+1', '\x00', 'a' * 32768],
                         ids=['formula', 'control-character', 'too-long'])
def test_invalid_replacement_is_atomic(session, replacement):
    before = snapshot(session.probe.book)
    match = XlsxSearch().move(session, 'cat')
    with pytest.raises((ValueError, ProtectedCellError)):
        session.replace_match(match, 'cat', replacement)
    assert snapshot(session.probe.book) == before
    assert not session.undo_stack and not session.dirty


def test_noop_keeps_redo_and_does_not_add_history(session):
    session.edit(1, 0, 'dog')
    session.history()
    match = XlsxSearch().move(session, 'cat')
    assert not session.replace_match(match, 'cat', 'cat')
    assert session.redo_stack and not session.undo_stack


@pytest.mark.parametrize('field,value', [('start', -1), ('end', 99), ('value', 'stale'), ('sheet', 'missing'), ('row', 999)])
def test_invalid_or_stale_match_does_not_edit(session, field, value):
    match = XlsxSearch().move(session, 'cat')
    with pytest.raises(ValueError):
        session.replace_match(replace(match, **{field: value}), 'cat', 'dog')
    assert not session.dirty and not session.undo_stack


def test_last_match_replaced_and_wraps_to_first(session):
    search = XlsxSearch()
    match = search.move(session, 'cat', backward=True)
    session.replace_match(match, 'cat', 'dog')
    assert search.after_replace(session, 'cat', match, 'dog').sheet == 'One'


def test_zero_remaining_matches_has_no_current_result(session):
    search = XlsxSearch()
    match = search.move(session, 'merged')
    # Use an editable unique string for the end-of-results case.
    session.edit(0, 0, 'unique')
    match = search.move(session, 'unique')
    session.replace_match(match, 'unique', 'done')
    assert search.after_replace(session, 'unique', match, 'done') is None
    with pytest.raises(ValueError):
        search.current(session, 'unique')


def test_gui_single_replace_updates_display_and_moves_to_next(session):
    pane = XlsxPane.__new__(XlsxPane)
    pane.session = session
    pane.query = SimpleNamespace(get=lambda: 'cat')
    pane.replacement = SimpleNamespace(get=lambda: 'dog')
    pane.all_sheets = SimpleNamespace(get=lambda: True)
    pane.formula_search = SimpleNamespace(get=lambda: False)
    pane.search_tracker = XlsxSearch()
    pane.search_tracker.move(session, 'cat')
    pane.selected_cell = lambda: session.sheet['A1']
    shown = []
    pane.render = lambda: shown.append('render')
    pane.show_match = lambda match: shown.append(match)
    pane.replace_one()
    assert session.sheet['A1'].value == 'dog cat'
    assert shown[-1].start == 4


def test_occurrence_limit_prevents_unbounded_expansion(session):
    session.sheet['A1'] = 'x' * 10_000
    search = XlsxSearch()
    assert search.move(session, 'x') is not None
    assert len(search.matches) == 10_000
    session.sheet['A2'] = 'x'
    with pytest.raises(ValueError, match='1万件'):
        search.move(session, 'x')


def test_gui_rejects_selection_change_since_search(session, monkeypatch):
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane.root = session, None
    pane.query = SimpleNamespace(get=lambda: 'cat')
    pane.replacement = SimpleNamespace(get=lambda: 'dog')
    pane.all_sheets = SimpleNamespace(get=lambda: True)
    pane.formula_search = SimpleNamespace(get=lambda: False)
    pane.search_tracker = XlsxSearch()
    pane.search_tracker.move(session, 'cat')
    pane.selected_cell = lambda: session.sheet['A2']
    errors = []
    monkeypatch.setattr('sheetdock.xlsx_gui.messagebox.showerror', lambda *args, **kwargs: errors.append(args))
    pane.replace_one()
    assert errors and not session.dirty


def test_post_replace_search_limit_keeps_edit_visible_and_undoable(session, monkeypatch):
    session.sheet['A1'] = 'unique'
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane.root = session, None
    pane.query = SimpleNamespace(get=lambda: 'u')
    pane.replacement = SimpleNamespace(get=lambda: 'u' * 6)
    pane.all_sheets = SimpleNamespace(get=lambda: True)
    pane.formula_search = SimpleNamespace(get=lambda: False)
    pane.search_tracker = XlsxSearch()
    pane.search_tracker.move(session, 'u')
    pane.selected_cell = lambda: session.sheet['A1']
    rendered, warnings = [], []
    pane.render = lambda: rendered.append(session.sheet['A1'].value)
    monkeypatch.setattr('sheetdock.xlsx_search.MAX_MATCHES', 5)
    monkeypatch.setattr('sheetdock.xlsx_gui.messagebox.showwarning', lambda *args, **kwargs: warnings.append(args))
    pane.replace_one()
    assert rendered == ['uuuuuunique']
    assert warnings and '次の検索を中止' in warnings[0][1]
    assert len(session.undo_stack) == 1 and pane.search_tracker.signature is None
    session.history()
    assert session.sheet['A1'].value == 'unique'
