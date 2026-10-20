from datetime import date, datetime, time, timedelta
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from sheetdock.xlsx_gui import (XlsxSession, XlsxPane, VIEW_BINDINGS,
                                  clipboard_rows, display_value, parse_input, protection_reason)
from sheetdock.xlsx_io import ProtectedCellError


@pytest.fixture
def path(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet['A1'], sheet['B1'], sheet['C1'] = '00123', 12, '=B1*2'
    sheet['A2'] = 'merged'
    sheet.merge_cells('A2:B2')
    book.create_sheet('Second')['A1'] = True
    path = tmp_path / 'test.xlsx'
    book.save(path)
    return path


@pytest.mark.parametrize('text,old,expected', [
    ('00045', '00123', '00045'), ('001', None, '001'), ('42', 12, 42),
    ('1.5', 12, 1.5), ('FALSE', True, False), (' true ', False, True),
    ('2026-10-15', date(2025, 1, 1), date(2026, 10, 15)),
    ('2026-10-15T12:00:00', datetime(2025, 1, 1), datetime(2026, 10, 15, 12)),
    ('12:30:00', time(0), time(12, 30)),
])
def test_parse_preserves_type(text, old, expected):
    value = parse_input(text, old)
    assert value == expected
    assert type(value) is type(expected)


@pytest.mark.parametrize('text,old', [('no', True), ('abc', 3), ('wrong', date.today()),
                                    ('1', timedelta(days=1))])
def test_invalid_conversion(text, old):
    with pytest.raises(ValueError):
        parse_input(text, old)


def test_display():
    assert display_value(None) == ''
    assert display_value(True) == 'TRUE'
    assert display_value(False) == 'FALSE'
    assert display_value('00123') == '00123'
    assert display_value('=SUM(A1:A3)') == '=SUM(A1:A3)'
    assert display_value(date(2026, 10, 15)) == '2026-10-15'


def test_sheet_switch_and_dirty(path):
    session = XlsxSession(path)
    assert not session.dirty
    session.edit(0, 0, '00045')
    session.sheet_name = 'Second'
    session.edit(0, 0, 'FALSE')
    session.sheet_name = 'Sheet'
    assert session.sheet['A1'].value == '00045'
    assert session.probe.book['Second']['A1'].value is False
    assert session.dirty


@pytest.mark.parametrize('row,column', [(0, 2), (1, 0), (1, 1)])
def test_formula_and_merge_blocked(path, row, column):
    session = XlsxSession(path)
    with pytest.raises(ProtectedCellError):
        session.edit(row, column, 'changed')
    assert not session.dirty


def test_input_error_leaves_model_unchanged(path):
    session = XlsxSession(path)
    with pytest.raises(ValueError):
        session.edit(0, 1, 'abc')
    assert session.sheet['B1'].value == 12
    assert not session.dirty


def test_revert_is_clean(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'abc')
    session.edit(0, 0, '00123')
    assert not session.dirty


def test_save_as_resets_dirty_without_changing_original(path, tmp_path):
    before = path.read_bytes()
    session = XlsxSession(path)
    session.edit(0, 0, '00045')
    destination = tmp_path / 'copy.xlsx'
    session.save_as(destination)
    assert not session.dirty
    assert session.probe.path == destination.resolve()
    assert path.read_bytes() == before
    assert load_workbook(destination)['Sheet']['A1'].value == '00045'


def test_failed_save_retains_dirty_and_path(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'change')
    with pytest.raises(FileExistsError):
        session.save_as(path)
    assert session.dirty
    assert session.probe.path == path.resolve()


def test_readonly_blocks_edits_and_saves(path, tmp_path):
    book = load_workbook(path)
    book.active.add_data_validation(DataValidation(type='list', formula1='"A,B"'))
    book.save(path)
    session = XlsxSession(path)
    assert protection_reason(session.probe, session.sheet, session.sheet['A1'])
    with pytest.raises(ProtectedCellError):
        session.edit(0, 0, 'abc')
    with pytest.raises(PermissionError):
        session.save_as(tmp_path / 'forbidden.xlsx')
    assert not session.dirty


def test_only_validated_inline_editing_is_added_to_bindings():
    assert 'edit_cell' in VIEW_BINDINGS
    assert set(VIEW_BINDINGS).isdisjoint({'all', 'paste', 'cut', 'delete',
                                        'undo', 'redo', 'row_drag_and_drop',
                                        'column_drag_and_drop', 'replace'})


def test_open_cancel_does_not_discard_session(path, monkeypatch):
    session = XlsxSession(path)
    gui = XlsxPane.__new__(XlsxPane)
    gui.session, gui.root = session, None
    monkeypatch.setattr('sheetdock.xlsx_gui.filedialog.askopenfilename', lambda **kw: '')
    gui.open_file()
    assert gui.session is session


def test_close_cancel_preserves_window(path, monkeypatch):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    gui = XlsxPane.__new__(XlsxPane)
    destroyed = []
    gui.session = session
    gui.root = SimpleNamespace(destroy=lambda: destroyed.append(True))
    monkeypatch.setattr('sheetdock.xlsx_gui.messagebox.askyesno', lambda *a, **kw: False)
    gui.close()
    assert not destroyed


def test_sparse_oversized_sheet_rejected(tmp_path):
    book = Workbook()
    book.active['A100001'] = 'far away'
    path = tmp_path / 'large.xlsx'
    book.save(path)
    with pytest.raises(ValueError, match='10万セル'):
        XlsxSession(path)


def test_edit_undo_redo(path):
    session = XlsxSession(path)
    session.edit(0, 0, '00045')
    assert session.history()
    assert session.sheet['A1'].value == '00123'
    assert not session.dirty
    assert session.history(redo=True)
    assert session.sheet['A1'].value == '00045'
    assert session.dirty


def test_range_paste_is_one_history_action(path):
    session = XlsxSession(path)
    session.paste(0, 0, '00045\t42\r\n')
    assert session.sheet['A1'].value == '00045'
    assert session.sheet['B1'].value == 42
    assert len(session.undo_stack) == 1
    session.history()
    assert not session.dirty
    session.history(True)
    assert session.sheet['B1'].value == 42


@pytest.mark.parametrize('text', ['changed\tabc', 'changed\t=1+1', 'changed\tNaN'])
def test_invalid_range_is_atomic(path, text):
    session = XlsxSession(path)
    with pytest.raises((ValueError, ProtectedCellError)):
        session.paste(0, 0, text)
    assert session.sheet['A1'].value == '00123'
    assert not session.dirty
    assert not session.undo_stack


def test_backend_validation_failure_rolls_back_prior_cells(path):
    session = XlsxSession(path)
    # The first assignment succeeds before the backend rejects the next value.
    with pytest.raises(ValueError):
        session.paste(0, 0, 'changed\tinf')
    assert session.sheet['A1'].value == '00123'
    assert session.sheet['B1'].value == 12
    assert not session.undo_stack


@pytest.mark.parametrize('row,column,text', [(0, 0, 'x\t3\tnew'),
                                           (1, 0, 'new\tnew'), (0, 2, 'new')])
def test_paste_formula_or_merge_is_atomic(path, row, column, text):
    session = XlsxSession(path)
    with pytest.raises(ProtectedCellError):
        session.paste(row, column, text)
    assert not session.dirty
    assert not session.undo_stack


@pytest.mark.parametrize('row,column,text', [(0, 3, 'x'), (2, 0, 'x'),
                                           (0, 1, '2\t3\t4'), (-1, 0, 'x')])
def test_paste_does_not_expand_sheet(path, row, column, text):
    session = XlsxSession(path)
    before = (session.sheet.max_row, session.sheet.max_column)
    with pytest.raises(ValueError):
        session.paste(row, column, text)
    assert (session.sheet.max_row, session.sheet.max_column) == before
    assert not session.dirty


def test_history_switches_back_to_edited_sheet(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.sheet_name = 'Second'
    session.edit(0, 0, 'FALSE')
    session.sheet_name = 'Sheet'
    session.history()
    assert session.sheet_name == 'Second'
    assert session.sheet['A1'].value is True
    session.history()
    assert session.sheet_name == 'Sheet'
    assert not session.dirty


def test_branch_clears_redo_but_noop_does_not(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.history()
    assert not session.edit(0, 0, '00123')
    assert session.redo_stack
    session.edit(0, 1, '30')
    assert not session.redo_stack


def test_failed_paste_keeps_redo(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.history()
    with pytest.raises(ValueError):
        session.paste(0, 0, 'x\tinf')
    assert session.redo_stack
    assert session.history(True)


def test_save_history_dirty_follows_saved_content(path, tmp_path):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.save_as(tmp_path / 'saved.xlsx')
    assert not session.dirty
    session.history()
    assert session.dirty
    session.history(True)
    assert not session.dirty


def test_undo_restores_style_and_type(path):
    session = XlsxSession(path)
    before = session.sheet['B1']._style
    session.edit(0, 1, '2.5')
    session.history()
    assert session.sheet['B1'].value == 12
    assert session.sheet['B1'].data_type == 'n'
    assert session.sheet['B1']._style == before


def test_quoted_clipboard_and_empty():
    assert clipboard_rows('"first\nsecond"\t"a\tb"\r\n') == [['first\nsecond', 'a\tb']]
    assert clipboard_rows('00123\t日本語\r\n00045\t末尾\r\n') == [['00123', '日本語'], ['00045', '末尾']]
    assert clipboard_rows('') == []


@pytest.mark.parametrize('text', ['a\tb\nc', '"unclosed'])
def test_bad_clipboard_rejected(text):
    import csv
    with pytest.raises((ValueError, csv.Error)):
        clipboard_rows(text)


def test_empty_clipboard_no_history(path):
    session = XlsxSession(path)
    assert not session.paste(0, 0, '')
    assert not session.undo_stack


def test_history_limit(path):
    session = XlsxSession(path)
    for i in range(105):
        session.edit(0, 0, str(i))
    assert len(session.undo_stack) == 100


def test_readonly_blocks_paste_and_history(path):
    session = XlsxSession(path)
    session.probe.reasons.append('unsupported')
    with pytest.raises(ProtectedCellError):
        session.paste(0, 0, 'changed')
    with pytest.raises(ProtectedCellError):
        session.history()
