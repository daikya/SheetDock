from copy import copy
from datetime import datetime, time, timedelta
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

import sheetdock.xlsx_tools as probe
from sheetdock.xlsx_tools import typed_key
from sheetdock.xlsx_gui import XlsxSession as SortFilterSession
from xlsx_sort_fixtures import make_samples
from sheetdock.xlsx_extent import extents
from sheetdock.xlsx_io import ProtectedCellError, snapshot


@pytest.fixture
def samples(tmp_path):
    make_samples(tmp_path)
    return tmp_path


@pytest.fixture
def session(samples):
    s = SortFilterSession(samples/'sort_basic.xlsx')
    try:
        yield s
    finally:
        s.probe.book.close()


def row_records(sheet):
    return {sheet.cell(row, 1).value: (
        tuple((cell.value, cell.data_type, copy(cell.font), copy(cell.fill),
               copy(cell.border), copy(cell.alignment), cell.number_format, copy(cell.protection))
              for cell in sheet[row]), sheet.row_dimensions[row].height)
        for row in range(2, 7)}


@pytest.mark.parametrize('column,descending,ids', [
    (2, False, ['00005', '00002', '00003', '00123', '00004']),
    (2, True, ['00123', '00002', '00003', '00005', '00004']),
    (1, False, ['00002', '00003', '00004', '00005', '00123']),
    (1, True, ['00123', '00005', '00004', '00003', '00002']),
])
def test_stable_sort_preserves_whole_typed_styled_rows(session, column, descending, ids):
    records = row_records(session.sheet)
    before_other = snapshot(session.probe.book)[1]
    old_columns = tuple((key, dict(value)) for key, value in session.sheet.column_dimensions.items())
    assert session.sort_rows(column, descending=descending)
    assert [session.sheet.cell(r, 1).value for r in range(2, 7)] == ids
    assert row_records(session.sheet) == records
    assert snapshot(session.probe.book)[1] == before_other
    assert tuple((key, dict(value)) for key, value in session.sheet.column_dimensions.items()) == old_columns
    assert session.sheet['A1'].value == 'コード'


def test_mixed_history_roundtrip_and_source_protection(session, samples):
    original = (samples/'sort_basic.xlsx').read_bytes()
    before, bounds = snapshot(session.probe.book), extents(session.probe.book)
    session.edit(1, 0, '00099')
    widths = session.column_widths()
    widths[0] += 30
    session.resize_columns({0: widths[0]})
    session.sort_rows(2)
    after = snapshot(session.probe.book)
    for _ in range(3):
        assert session.history()
    assert snapshot(session.probe.book) == before and not session.dirty
    for _ in range(3):
        assert session.history(True)
    assert snapshot(session.probe.book) == after and session.dirty
    target = samples/'copy.xlsx'
    session.save_as(target)
    reopened = SortFilterSession(target)
    try:
        assert snapshot(reopened.probe.book) == after and extents(reopened.probe.book) == bounds
    finally:
        reopened.probe.book.close()
    assert not session.dirty
    assert session.history() and session.dirty
    assert session.history(True) and not session.dirty
    assert (samples/'sort_basic.xlsx').read_bytes() == original


@pytest.mark.parametrize('name,exception', [('formula', ProtectedCellError), ('merged', ProtectedCellError),
    ('readonly', ProtectedCellError), ('hidden', ProtectedCellError), ('error', ProtectedCellError), ('mixed', TypeError)])
def test_unsupported_sort_is_atomic(samples, name, exception):
    path = samples/f'sort_{name}.xlsx'
    raw = path.read_bytes()
    s = SortFilterSession(path)
    try:
        before, bounds = snapshot(s.probe.book), extents(s.probe.book)
        with pytest.raises(exception):
            s.sort_rows(2)
        assert snapshot(s.probe.book) == before and extents(s.probe.book) == bounds
        assert not s.dirty and not s.undo_stack and not s.redo_stack
        assert path.read_bytes() == raw
    finally:
        s.probe.book.close()


@pytest.mark.parametrize('operator,query,rows', [('contains', '東京', [2, 4, 6]),
    ('equals', '東京', [2, 4, 6]), ('equals', 'none', []), ('nonempty', '', [2, 3, 4, 5, 6])])
def test_filter_is_only_projection_and_keeps_history(session, operator, query, rows):
    session.edit(1, 0, '00099')
    session.history()
    before = snapshot(session.probe.book)
    undo, redo = list(session.undo_stack), list(session.redo_stack)
    cells = set(session.sheet._cells)
    session.set_filter(5, operator=operator, query=query)
    assert session.visible_rows() == rows
    assert set(session.sheet._cells) == cells and snapshot(session.probe.book) == before
    assert session.undo_stack == undo and session.redo_stack == redo and not session.dirty
    session.clear_filter()
    assert session.visible_rows() == [2, 3, 4, 5, 6]
    assert session.history(True) and session.sheet['A2'].value == '00099'


def test_filter_raw_number_boolean_datetime_and_blanks(session):
    for column, query, expected in [(2, '2', [3, 4]), (4, 'TRUE', [2, 4, 6]),
                                    (3, '2026-10-04', [4])]:
        session.set_filter(column, query=query)
        assert session.visible_rows() == expected
    session.set_filter(2, operator='empty')
    assert session.visible_rows() == [5]
    session.clear_filter()
    session.append_axis('rows', 2)
    allocated = set(session.sheet._cells)
    session.set_filter(2, operator='empty')
    assert session.visible_rows() == [5, 7, 8]
    assert set(session.sheet._cells) == allocated


@pytest.mark.parametrize('name', ['readonly', 'formula', 'merged', 'error'])
def test_view_filters_remain_available_in_protected_books(samples, name):
    s = SortFilterSession(samples/f'sort_{name}.xlsx')
    try:
        before = snapshot(s.probe.book)
        s.set_filter(5, query='東京')
        assert s.visible_rows() == [2, 4, 6]
        assert snapshot(s.probe.book) == before and not s.dirty
        assert not s.undo_stack and not s.redo_stack
    finally:
        s.probe.book.close()


def test_existing_hidden_flags_are_not_changed_or_redisplayed(samples):
    s = SortFilterSession(samples/'sort_hidden.xlsx')
    try:
        s.set_filter(5, query='大阪')
        assert s.visible_rows() == [] and s.sheet.row_dimensions[3].hidden
        s.clear_filter()
        assert s.visible_rows() == [2, 4, 5, 6]
    finally:
        s.probe.book.close()


def test_filter_conditions_are_per_sheet_and_sort_requires_clear(session):
    session.set_filter(5, query='東京')
    with pytest.raises(ValueError, match='解除'):
        session.sort_rows(2)
    session.sheet_name = 'Other'
    session.set_filter(1, query='他', header=False)
    assert session.visible_rows() == [1]
    session.sheet_name = 'Data'
    assert session.visible_rows() == [2, 4, 6]
    session.clear_filter()
    assert session.sort_rows(2)
    session.sheet_name = 'Other'
    assert session.visible_rows() == [1]


@pytest.mark.parametrize('value,family', [(10, 'number'), (2.5, 'number'), ('00123', 'text'),
    (True, 'boolean'), (datetime(2026, 10, 20), 'datetime'), (time(9, 30), 'time'), (timedelta(days=1), 'timedelta')])
def test_typed_comparison_families_do_not_coerce_original_values(value, family):
    cell = SimpleNamespace(value=value, data_type='s' if isinstance(value, str) else 'n')
    assert typed_key(cell)[0] == family and cell.value == value


@pytest.mark.parametrize('column', [0, -1, True, 7, '2'])
def test_invalid_columns_do_not_change_book_or_filters(session, column):
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError):
        session.sort_rows(column)
    with pytest.raises(ValueError):
        session.set_filter(column)
    assert snapshot(session.probe.book) == before and not session.filters


def test_noop_sort_preserves_redo_and_no_header_sorts_first_row(samples):
    path = samples/'simple.xlsx'
    book = Workbook()
    book.active.append([2])
    book.active.append([1])
    book.save(path)
    book.close()
    s = SortFilterSession(path)
    try:
        s.edit(0, 0, '3')
        s.history()
        assert not s.sort_rows(1) and not s.undo_stack and s.redo_stack
        assert s.sort_rows(1, header=False)
        assert [s.sheet.cell(r, 1).value for r in (1, 2)] == [1, 2]
        assert not s.redo_stack
    finally:
        s.probe.book.close()


@pytest.mark.parametrize('failure', ['unsupported', 'reload_mismatch'])
def test_candidate_validation_failure_keeps_live_book_and_history(session, monkeypatch, failure):
    session.edit(1, 0, '00099')
    session.history()
    book, before = session.probe.book, snapshot(session.probe.book)
    undo, redo = list(session.undo_stack), list(session.redo_stack)
    if failure == 'unsupported':
        monkeypatch.setattr(probe, 'inspect_package', lambda data: ('unsupported fixture',))
    else:
        original_load = probe.load_workbook
        calls = []
        def load(*args, **kwargs):
            result = original_load(*args, **kwargs)
            calls.append(result)
            if len(calls) == 2:
                result['Data']['A2'] = 'corrupt fixture'
            return result
        monkeypatch.setattr(probe, 'load_workbook', load)
    with pytest.raises(ValueError):
        session.sort_rows(2)
    assert session.probe.book is book and snapshot(book) == before
    assert session.undo_stack == undo and session.redo_stack == redo


@pytest.mark.parametrize('values,expected', [
    ([datetime(2026, 10, 22), datetime(2026, 10, 20), datetime(2026, 10, 21)], [2, 3, 1]),
    ([time(12), time(8), time(10)], [2, 3, 1]),
    ([timedelta(hours=12), timedelta(hours=8), timedelta(hours=10)], [2, 3, 1]),
    ([True, False, True], [2, 1, 3]),
    (['10', '2', '001'], [3, 1, 2]),
])
def test_actual_typed_key_sort_roundtrip(tmp_path, values, expected):
    path = tmp_path/'typed.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.append(['ID', 'key'])
    for i, value in enumerate(values, 1):
        sheet.append([i, value])
    book.save(path)
    book.close()
    s = SortFilterSession(path)
    try:
        assert s.sort_rows(2)
        assert [s.sheet.cell(r, 1).value for r in range(2, 5)] == expected
        assert s.history() and not s.dirty
    finally:
        s.probe.book.close()


def test_row_dimension_style_moves_with_its_row(session):
    session.sheet.row_dimensions[2].font = Font(bold=True, color='FFAA0033')
    before = copy(session.sheet.row_dimensions[2].font)
    session.sort_rows(2)
    assert session.sheet['A5'].value == '00123'
    assert session.sheet.row_dimensions[5].font == before
    assert session.history() and session.sheet.row_dimensions[2].font == before


def test_default_hidden_rows_reject_before_condition_or_book_changes(session):
    for name, value in [('zeroHeight', True), ('defaultRowHeight', 0)]:
        original = getattr(session.sheet.sheet_format, name)
        setattr(session.sheet.sheet_format, name, value)
        before = snapshot(session.probe.book)
        with pytest.raises(ProtectedCellError):
            session.sort_rows(2)
        with pytest.raises(ValueError):
            session.set_filter(5, query='東京')
        assert not session.filters and snapshot(session.probe.book) == before and not session.undo_stack
        setattr(session.sheet.sheet_format, name, original)


def test_oversized_sheet_view_is_rejected_without_mutating_book(session):
    session.sheet.row_dimensions[1001].height = 20
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError, match='1000'):
        session.sort_rows(2)
    with pytest.raises(ValueError, match='1000'):
        session.set_filter(5, query='東京')
    assert snapshot(session.probe.book) == before and not session.filters
