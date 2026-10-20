from copy import copy

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.dimensions import ColumnDimension

from sheetdock.xlsx_gui import XlsxSession, XlsxPane
from sheetdock.xlsx_io import ProtectedCellError, snapshot
from sheetdock.xlsx_widths import display_widths, stored_width, pixels_to_width, width_to_pixels


@pytest.fixture
def path(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'One'
    sheet.append(['00123', '=1+1', 'merged'])
    sheet.merge_cells('C1:D1')
    sheet.column_dimensions['A'].width = 24.123046875
    sheet.column_dimensions['A'].font = Font(bold=True)
    book.create_sheet('Two').append(['日本語', 123])
    book['Two'].column_dimensions['A'].width = 12
    path = tmp_path / 'widths.xlsx'
    book.save(path)
    return path


def test_spec_pixel_example():
    assert width_to_pixels(8.7109375, 7) == 61


@pytest.mark.parametrize('metric', [5, 7, 9, 12])
def test_integer_pixel_roundtrip(metric):
    for pixels in range(20, width_to_pixels(255, metric) + 1):
        assert width_to_pixels(pixels_to_width(pixels, metric), metric) == pixels


@pytest.mark.parametrize('pixels', [0, 19, 5000, float('nan'), float('inf'), True])
def test_bad_pixel_width_rejected(pixels):
    with pytest.raises(ValueError):
        pixels_to_width(pixels)


def test_read_widths_does_not_create_dimensions_or_dirty(path):
    session = XlsxSession(path)
    before = snapshot(session.probe.book)
    assert len(session.column_widths()) == 4
    assert snapshot(session.probe.book) == before
    assert not session.dirty


def test_resize_dirty_undo_redo_preserves_original_precision(path):
    session = XlsxSession(path)
    original = session.sheet.column_dimensions['A'].width
    session.resize_columns({0: 250})
    assert session.dirty and session.column_widths()[0] == 250
    assert session.sheet.column_dimensions['A'].font.bold
    session.history()
    assert session.sheet.column_dimensions['A'].width == original
    assert not session.dirty
    session.history(True)
    assert session.column_widths()[0] == 250


def test_default_dimension_removed_on_undo(path):
    session = XlsxSession(path)
    assert 'B' not in session.sheet.column_dimensions
    session.resize_columns({1: 180})
    assert 'B' in session.sheet.column_dimensions
    session.history()
    assert 'B' not in session.sheet.column_dimensions
    assert not session.dirty


def test_sheets_keep_distinct_widths(path):
    session = XlsxSession(path)
    session.resize_columns({0: 210})
    session.sheet_name = 'Two'
    assert session.column_widths()[0] == width_to_pixels(12)
    session.resize_columns({0: 140})
    session.sheet_name = 'One'
    assert session.column_widths()[0] == 210
    session.history()
    assert session.sheet_name == 'Two'
    assert session.column_widths()[0] == width_to_pixels(12)


def test_width_saved_inside_xlsx_and_reopened(path, tmp_path):
    session = XlsxSession(path)
    session.resize_columns({0: 220})
    target = tmp_path / 'saved.xlsx'
    session.save_as(target)
    assert not session.dirty
    assert XlsxSession(target).column_widths()[0] == 220
    assert load_workbook(target)['One'].column_dimensions['A'].width == session.sheet.column_dimensions['A'].width
    session.history()
    assert session.dirty
    session.history(True)
    assert not session.dirty


def test_unchanged_pixels_do_not_quantize_existing_width(path):
    session = XlsxSession(path)
    original = session.sheet.column_dimensions['A'].width
    assert not session.resize_columns({0: session.column_widths()[0]})
    assert session.sheet.column_dimensions['A'].width == original
    assert not session.dirty and not session.undo_stack


def test_multi_column_resize_one_history(path):
    session = XlsxSession(path)
    before = session.column_widths()
    session.resize_columns({0: 170, 1: 180})
    assert len(session.undo_stack) == 1
    session.history()
    assert session.column_widths() == before


def test_invalid_batch_is_atomic(path):
    session = XlsxSession(path)
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError):
        session.resize_columns({0: 250, 99: 200})
    assert snapshot(session.probe.book) == before
    assert not session.undo_stack


def test_grouped_columns_split_and_restore(path):
    book = load_workbook(path)
    sheet = book['One']
    sheet.column_dimensions['B'] = ColumnDimension(sheet, index='B', min=2, max=4, width=15, outlineLevel=1)
    book.save(path)
    session = XlsxSession(path)
    session.resize_columns({2: 220})
    assert stored_width(session.sheet, 2) == 15
    assert stored_width(session.sheet, 4) == 15
    assert session.column_widths()[2] == 220
    assert session.sheet.column_dimensions['C'].outlineLevel == 1
    session.pack_book()  # The split ranges must pass the real save/reload verifier.
    session.history()
    assert 'C' not in session.sheet.column_dimensions
    assert session.sheet.column_dimensions['B'].max == 4
    assert not session.dirty


def test_hidden_column_rejected_and_previous_changes_rolled_back(path):
    book = load_workbook(path)
    book['One'].column_dimensions['B'].hidden = True
    book.save(path)
    session = XlsxSession(path)
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError, match='非表示'):
        session.resize_columns({0: 250, 1: 170})
    assert snapshot(session.probe.book) == before
    assert not session.dirty


def test_readonly_rejects_resize(path):
    session = XlsxSession(path)
    session.probe.reasons.append('unsupported')
    with pytest.raises(ProtectedCellError):
        session.resize_columns({0: 180})


def test_cell_history_and_width_history_interleave(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.resize_columns({0: 230})
    session.sheet_operation('add', 'Extra')
    for _ in range(3):
        session.history()
    assert not session.dirty
    for _ in range(3):
        session.history(True)
    assert session.probe.book['One']['A1'].value == 'changed'
    assert session.probe.book.sheetnames[-1] == 'Extra'


def test_width_does_not_change_formula_or_merge(path):
    session = XlsxSession(path)
    session.resize_columns({1: 180, 2: 200})
    assert session.sheet['B1'].value == '=1+1'
    assert str(next(iter(session.sheet.merged_cells.ranges))) == 'C1:D1'


def test_pane_resize_uses_only_changed_columns_and_not_json(path):
    pane = XlsxPane.__new__(XlsxPane)
    pane.session = XlsxSession(path)
    pane._rendering, pane._digit_width = False, 7
    calls = []
    pane._on_change = lambda: calls.append(True)
    pane.column_resized({'resized': {'columns': {0: {'old_size': 170, 'new_size': 240}}}})
    assert pane.session.column_widths()[0] == 240
    assert calls == [True]


def test_rendering_resize_event_is_ignored(path):
    pane = XlsxPane.__new__(XlsxPane)
    pane.session = XlsxSession(path)
    pane._rendering = True
    pane.column_resized({'resized': {'columns': {0: {'new_size': 240}}}})
    assert not pane.session.dirty


def test_failed_save_keeps_width_edit_and_history(path, monkeypatch):
    session = XlsxSession(path)
    session.resize_columns({0: 245})
    before = path.read_bytes()
    def fail(*args, **kwargs):
        raise OSError('simulated failure')
    monkeypatch.setattr('sheetdock.xlsx_gui.save_checked', fail)
    with pytest.raises(OSError):
        session.save()
    assert session.dirty and session.undo_stack
    assert session.column_widths()[0] == 245
    assert path.read_bytes() == before


def test_normal_style_digit_measurement_uses_font_object(path, monkeypatch):
    from types import SimpleNamespace
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane.root = XlsxSession(path), None
    observed = []
    def fake_font(**kwargs):
        observed.append(kwargs)
        return SimpleNamespace(measure=lambda text: 9)
    monkeypatch.setattr('sheetdock.xlsx_gui.tkfont.Font', fake_font)
    assert pane.measure_digit_width() == 9
    assert observed[0]['family'] == 'Calibri'
    assert observed[0]['size'] == 11


def drag_pane(path):
    from types import SimpleNamespace
    pane = XlsxPane.__new__(XlsxPane)
    pane.session = XlsxSession(path)
    pane._rendering, pane._digit_width = False, 7
    pane._pending_column_widths = {}
    pane.sheet_widget = SimpleNamespace(CH=SimpleNamespace(currently_resizing_width=True))
    pane._on_change = lambda: None
    return pane


def drag_to(pane, pixels):
    pane.column_resized({'resized': {'columns': {0: {'new_size': pixels}}}})


def release_drag(pane):
    pane.sheet_widget.CH.currently_resizing_width = False
    pane.finish_column_resize()


def test_drag_motion_is_one_undo_step(path):
    pane = drag_pane(path)
    original = pane.session.column_widths()
    for pixels in range(200, 260):
        drag_to(pane, pixels)
    assert not pane.session.undo_stack
    release_drag(pane)
    assert len(pane.session.undo_stack) == 1
    assert pane.session.column_widths()[0] == 259
    pane.session.history()
    assert pane.session.column_widths() == original
    assert not pane.session.dirty
    pane.session.history(True)
    assert pane.session.column_widths()[0] == 259


def test_separate_drags_remain_separate_history_steps(path):
    pane = drag_pane(path)
    drag_to(pane, 240)
    release_drag(pane)
    pane.sheet_widget.CH.currently_resizing_width = True
    drag_to(pane, 260)
    drag_to(pane, 280)
    release_drag(pane)
    assert len(pane.session.undo_stack) == 2
    pane.session.history()
    assert pane.session.column_widths()[0] == 240
    pane.finish_column_resize()  # Release without another drag must do nothing.
    assert len(pane.session.undo_stack) == 1


def test_drag_back_to_original_keeps_exact_width_and_redo(path):
    pane = drag_pane(path)
    original_width = pane.session.sheet.column_dimensions['A'].width
    pane.session.edit(0, 0, 'changed')
    pane.session.history()
    original_pixels = pane.session.column_widths()[0]
    drag_to(pane, original_pixels + 40)
    drag_to(pane, original_pixels)
    release_drag(pane)
    assert not pane.session.dirty
    assert not pane.session.undo_stack
    assert pane.session.redo_stack
    assert pane.session.sheet.column_dimensions['A'].width == original_width
