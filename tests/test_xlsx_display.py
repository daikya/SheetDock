from datetime import date, datetime, time, timedelta
import csv

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from sheetdock.xlsx_display import format_value, cell_alignment
from sheetdock.xlsx_gui import XlsxSession, XlsxPane
from sheetdock.xlsx_io import snapshot


@pytest.mark.parametrize('value,pattern,expected', [
    (12345, '#,##0', '12,345'), (1234.5, '#,##0.00', '1,234.50'),
    (-1234.567, '#,##0.00', '-1,234.57'), (2.675, '0.00', '2.68'),
    (-2.675, '0.00', '-2.68'), (0.125, '0%', '13%'),
    (0.125, '0.00%', '12.50%'), (-0.0001, '0.00', '0.00'),
    (1234.5, '0.000000', '1234.500000'), (12, 'General', '12'),
    (1234, '@', '1234'), ('00123', '#,##0', '00123'),
    ('=1+1', '0.00', '=1+1'), (True, '0.00', 'TRUE'),
    (False, '0%', 'FALSE'), (None, '0.00', ''),
    (datetime(2026, 10, 15, 9, 5, 3), 'yyyy/mm/dd hh:mm:ss', '2026/10/15 09:05:03'),
    (date(2026, 1, 5), 'YYYY-MM-DD', '2026-01-05'),
    (date(2026, 1, 5), 'yyyy/m/d', '2026/1/5'),
    (date(2026, 1, 5), 'mm-dd-yy', '01-05-26'),
    (time(9, 5, 3), 'hh:mm:ss', '09:05:03'),
    (time(9, 5), 'h:mm', '9:05'),
])
def test_supported_display(value, pattern, expected):
    assert format_value(value, pattern) == (expected, True)


@pytest.mark.parametrize('value,pattern,expected', [
    (1234, '$#,##0', '1234'), (1, '[Red]0', '1'),
    (1, '0;[Red]-0', '1'), (1.2, '0.00E+00', '1.2'),
    (date(2026, 10, 15), 'yyyy"年"m"月"d"日"', '2026-10-15'),
    (time(1, 2), 'yyyy-mm-dd', '01:02:00'),
    (timedelta(days=1), '[h]:mm:ss', '1 day, 0:00:00'),
    (float('inf'), '0.00', 'inf'),
])
def test_unknown_format_uses_raw_value(value, pattern, expected):
    assert format_value(value, pattern) == (expected, False)


@pytest.mark.parametrize('value,explicit,expected', [
    ('text', None, 'left'), (123, None, 'right'),
    (date(2026, 10, 15), None, 'right'), (True, None, 'center'),
    ('=1+1', None, 'left'), (123, 'left', 'left'),
    ('text', 'center', 'center'), ('text', 'right', 'right'),
])
def test_cell_alignment(value, explicit, expected):
    cell = Workbook().active['A1']
    cell.value = value
    cell.alignment = Alignment(horizontal=explicit)
    assert cell_alignment(cell) == expected


def test_display_does_not_modify_values_styles_or_dirty_state(tmp_path):
    book = Workbook()
    book.active.append([1234.5, 0.125, datetime(2026, 10, 15)])
    for cell, pattern in zip(book.active[1], ('#,##0.00', '0.0%', 'yyyy/mm/dd')):
        cell.number_format = pattern
    path = tmp_path / 'formats.xlsx'
    book.save(path)
    session = XlsxSession(path)
    before = snapshot(session.probe.book)
    assert [format_value(c.value, c.number_format)[0] for c in session.sheet[1]] == [
        '1,234.50', '12.5%', '2026/10/15']
    assert not session.dirty and snapshot(session.probe.book) == before
    session.edit(0, 0, '2345.5')
    assert format_value(session.sheet['A1'].value, session.sheet['A1'].number_format)[0] == '2,345.50'
    session.history()
    assert snapshot(session.probe.book) == before and not session.dirty
    session.history(True)
    saved = tmp_path / 'saved.xlsx'
    session.save_as(saved)
    reopened = load_workbook(saved)
    assert reopened.active['A1'].value == 2345.5
    assert reopened.active['A1'].number_format == '#,##0.00'
    export = tmp_path / 'raw.csv'
    session.export_csv(export)
    with export.open(encoding='utf-8-sig', newline='') as handle:
        assert list(csv.reader(handle)) == [['2345.5', '0.125', '2026-10-15T00:00:00']]


def test_edit_dialog_uses_raw_value_instead_of_formatted_display(tmp_path, monkeypatch):
    book = Workbook()
    book.active['A1'] = 0.125
    book.active['A1'].number_format = '0.00%'
    path = tmp_path / 'percent.xlsx'
    book.save(path)
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane.root = XlsxSession(path), None
    pane.selected_cell = lambda: pane.session.sheet['A1']
    observed = []
    def ask(*args, **kwargs):
        observed.append(kwargs['initialvalue'])
        return '0.25'
    monkeypatch.setattr('sheetdock.xlsx_gui.simpledialog.askstring', ask)
    pane.render = lambda: None
    pane.edit_cell()
    assert observed == ['0.125']
    assert pane.session.sheet['A1'].value == 0.25
    assert pane.session.sheet['A1'].number_format == '0.00%'
