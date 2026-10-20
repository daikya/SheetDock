from copy import copy
import json
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from sheetdock.editor import Editor
from sheetdock.gui import MultiFileApp, Session
from sheetdock.csv_io import CsvDocument, CsvFormat
from sheetdock.heights import fit_heights
from sheetdock.settings import SettingsStore
from sheetdock.tabs import Documents, path_key
from sheetdock.xlsx_gui import XlsxSession, XlsxPane
from sheetdock.xlsx_io import snapshot, ProtectedCellError
from sheetdock.xlsx_heights import height_to_pixels, pixels_to_height


def test_csv_height_history_is_clean_and_redoable():
    editor = Editor([['a'], ['b']])
    editor.resize_rows({0: 60})
    assert editor.row_heights == [60, None] and not editor.dirty
    editor.undo()
    assert editor.row_heights == [None, None] and not editor.dirty
    editor.redo()
    assert editor.row_heights == [60, None] and not editor.dirty


def test_csv_height_follows_insert_delete_and_mixed_history():
    editor = Editor([['a'], ['b']])
    editor.resize_rows({1: 60})
    editor.resize_columns([220])
    editor.set_cell(0, 0, 'edited')
    editor.insert_rows(0, 2)
    assert editor.row_heights == [None, None, None, 60]
    editor.delete_rows([1, 2])
    assert editor.row_heights == [None, 60]
    for _ in range(5):
        assert editor.undo()
    assert editor.rows == [['a'], ['b']]
    assert editor.row_heights == [None, None] and editor.column_widths == [120]
    for _ in range(5):
        assert editor.redo()
    assert editor.row_heights == [None, 60]


def test_csv_header_toggle_and_column_delete_keep_heights():
    session = Session(CsvDocument([['head'], ['body']], CsvFormat('utf-8', False, '\n', ',')), True)
    session.editor.resize_rows({1: 60})
    session.set_header_view(False)
    assert session.editor.row_heights == [None, 60]
    session.editor.delete_columns([0])
    assert session.editor.row_heights == [None, 60]
    session.editor.undo()
    session.editor.undo()
    assert session.editor.row_heights == [None, None]


def test_csv_invalid_height_is_atomic_and_new_change_discards_redo():
    editor = Editor([['a'], ['b']])
    with pytest.raises(ValueError):
        editor.resize_rows({0: 60, 9: 40})
    assert editor.row_heights == [None, None] and not editor._undo
    editor.resize_rows({0: 60})
    editor.undo()
    assert not editor.resize_rows({}) and editor._redo
    editor.resize_rows({1: 70})
    assert not editor.redo()


def test_settings_legacy_migration_and_other_updates_keep_heights(tmp_path):
    store = SettingsStore(tmp_path / 'settings.json')
    store.path.write_text(json.dumps({'version': 1, 'file_widths': {'old': [150]}}))
    assert store.load().file_heights == {}
    path = tmp_path / 'data.csv'
    store.update(heights_file=path, row_heights=[None, 70, False, -1, 'bad'])
    store.update(backup_enabled=False, widths_file=path, column_widths=[180])
    loaded = store.load()
    assert loaded.file_heights[path_key(path)] == [None, 70, None, None, None]
    assert loaded.file_widths['old'] == [150]


@pytest.mark.parametrize('value', [None, True, -1, 0, float('nan'), float('inf'), 2001, '60'])
def test_bad_height_settings_fall_back_to_default(value):
    assert fit_heights([value], 2) == [None, None]


def test_csv_drag_coalesces_and_accounts_for_header_offset(tmp_path):
    app = MultiFileApp.__new__(MultiFileApp)
    tab = Documents().add(tmp_path / 'data.csv', Session(
        CsvDocument([['head'], ['body']], CsvFormat('utf-8', False, '\n', ',')), True))
    widget = SimpleNamespace(RI=SimpleNamespace(currently_resizing_height=True), _sheetdock_height_drag={})
    app.pages, app.rendering, app.applying_widths = {tab: (None, widget)}, False, False
    app.settings_store = SettingsStore(tmp_path / 'settings.json')
    for old, new in ((20, 40), (40, 60), (60, 80)):
        app.row_resized(tab, {'resized': {'rows': {0: {'old_size': old, 'new_size': new}}}})
    assert not tab.session.editor._undo
    widget.RI.currently_resizing_height = False
    app.finish_row_resize(tab)
    assert tab.session.editor.row_heights == [None, 80]
    assert len(tab.session.editor._undo) == 1 and not tab.session.editor.dirty
    assert app.settings_store.load().file_heights[path_key(tab.source)] == [None, 80]
    tab.session.editor.undo()
    app.persist_widths(tab)
    assert app.settings_store.load().file_heights[path_key(tab.source)] == [None, None]


def test_csv_drag_returning_to_start_does_not_store_default(tmp_path):
    app = MultiFileApp.__new__(MultiFileApp)
    tab = Documents().add(tmp_path / 'data.csv', Session(
        CsvDocument([['body']], CsvFormat('utf-8', False, '\n', ',')), False))
    widget = SimpleNamespace(RI=SimpleNamespace(currently_resizing_height=True), _sheetdock_height_drag={})
    app.pages, app.rendering, app.applying_widths = {tab: (None, widget)}, False, False
    app.row_resized(tab, {'resized': {'rows': {0: {'old_size': 20, 'new_size': 60}}}})
    app.row_resized(tab, {'resized': {'rows': {0: {'old_size': 60, 'new_size': 20}}}})
    widget.RI.currently_resizing_height = False
    app.finish_row_resize(tab)
    assert tab.session.editor.row_heights == [None] and not tab.session.editor._undo


@pytest.fixture
def path(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'One'
    sheet.append(['a', 1]); sheet.append(['b', 2]); sheet.append(['c', 3])
    sheet.row_dimensions[2].height = 32.123456
    sheet.row_dimensions[2].font = Font(bold=True)
    book.create_sheet('Two').append(['other'])
    path = tmp_path / 'height.xlsx'
    book.save(path)
    return path


@pytest.mark.parametrize('scale', [1, 4/3, 1.5, 2])
def test_xlsx_point_pixel_roundtrip(scale):
    for pixels in range(5, height_to_pixels(409.5, scale) + 1):
        assert height_to_pixels(pixels_to_height(pixels, scale), scale) == pixels


def test_xlsx_read_is_nonmutating_and_noop_keeps_precision(path):
    session = XlsxSession(path)
    before = snapshot(session.probe.book)
    heights = session.row_heights()
    assert len(heights) == 3 and heights[1] > heights[0]
    assert not session.resize_rows({1: heights[1]})
    assert session.sheet.row_dimensions[2].height == 32.123456
    assert snapshot(session.probe.book) == before and not session.dirty


def test_xlsx_resize_undo_keeps_style_and_original_precision(path):
    session = XlsxSession(path)
    before = snapshot(session.probe.book)
    session.resize_rows({1: 90})
    assert session.dirty and session.row_heights()[1] == 90
    assert session.sheet.row_dimensions[2].font.bold
    session.history()
    assert snapshot(session.probe.book) == before and not session.dirty
    session.history(True)
    assert session.row_heights()[1] == 90


def test_default_row_definition_removed_on_undo(path):
    session = XlsxSession(path)
    assert 1 not in session.sheet.row_dimensions
    session.resize_rows({0: 60})
    session.history()
    assert 1 not in session.sheet.row_dimensions and not session.dirty


@pytest.mark.parametrize('case', ['hidden', 'zero', 'zero-default', 'default-zero-height'])
def test_xlsx_rejects_hidden_rows_atomically(path, case):
    session = XlsxSession(path)
    if case == 'hidden':
        session.sheet.row_dimensions[2].hidden = True
    elif case == 'zero':
        session.sheet.row_dimensions[2].height = 0
    elif case == 'zero-default':
        session.sheet.sheet_format.zeroHeight = True
    else:
        session.sheet.sheet_format.defaultRowHeight = 0
    before = snapshot(session.probe.book)
    changes = {0: 80, 1: 90} if case != 'zero-default' else {0: 80}
    with pytest.raises(ValueError):
        session.resize_rows(changes)
    assert snapshot(session.probe.book) == before and not session.undo_stack


@pytest.mark.parametrize('pixels', [True, 0, -1, 5000, float('nan'), float('inf'), None])
def test_xlsx_invalid_height_rejected(path, pixels):
    session = XlsxSession(path)
    with pytest.raises(ValueError):
        session.resize_rows({0: pixels})
    assert not session.dirty and not session.undo_stack


def test_xlsx_height_save_reopen_and_sheet_isolation(path, tmp_path):
    session = XlsxSession(path)
    session.resize_rows({0: 60})
    session.sheet_name = 'Two'
    session.resize_rows({0: 90})
    target = tmp_path / 'saved.xlsx'
    session.save_as(target)
    reopened = XlsxSession(target)
    assert reopened.row_heights()[0] == 60
    reopened.sheet_name = 'Two'
    assert reopened.row_heights()[0] == 90
    assert not session.dirty
    session.history()
    assert session.dirty
    session.history(True)
    assert not session.dirty


def test_xlsx_height_cell_width_axis_histories_work_together(path):
    session = XlsxSession(path)
    before = snapshot(session.probe.book)
    session.resize_rows({1: 90})
    session.resize_columns({0: 180})
    session.edit(0, 0, 'changed')
    session.axis_operation('insert-rows', 2, 1)
    assert session.sheet.row_dimensions[3].height == 67.5
    after = snapshot(session.probe.book)
    for _ in range(4):
        session.history()
    assert snapshot(session.probe.book) == before and not session.dirty
    for _ in range(4):
        session.history(True)
    assert snapshot(session.probe.book) == after


def test_xlsx_readonly_and_save_failure_keep_original(path, monkeypatch):
    session = XlsxSession(path)
    session.probe.reasons.append('unsupported')
    with pytest.raises(ProtectedCellError):
        session.resize_rows({0: 60})
    session.probe.reasons.clear()
    session.resize_rows({0: 60})
    original = path.read_bytes()
    def fail(*args, **kwargs):
        raise OSError('simulated save failure')
    monkeypatch.setattr('sheetdock.xlsx_gui.save_checked', fail)
    with pytest.raises(OSError):
        session.save()
    assert session.dirty and session.undo_stack and path.read_bytes() == original


def test_xlsx_drag_one_step_and_cancel_button(path, monkeypatch):
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane.root = XlsxSession(path), None
    pane._rendering, pane._point_scale, pane._row_minimum = False, 4/3, 5
    pane._pending_row_sizes = {}
    pane.sheet_widget = SimpleNamespace(RI=SimpleNamespace(currently_resizing_height=True))
    calls = []
    pane._on_change = lambda: calls.append(True)
    for old, new in ((20, 40), (40, 60), (60, 80)):
        pane.row_resized({'resized': {'rows': {0: {'old_size': old, 'new_size': new}}}})
    assert not pane.session.undo_stack
    pane.sheet_widget.RI.currently_resizing_height = False
    pane.finish_row_resize()
    assert len(pane.session.undo_stack) == 1 and calls == [True]
    pane.selected_cell = lambda: pane.session.sheet['A1']
    monkeypatch.setattr('sheetdock.xlsx_gui.simpledialog.askinteger', lambda *args, **kwargs: None)
    pane.set_row_height()
    assert len(pane.session.undo_stack) == 1
