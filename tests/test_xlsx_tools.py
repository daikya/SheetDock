from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from sheetdock.xlsx_gui import XlsxSession
from sheetdock.xlsx_io import ProtectedCellError
from sheetdock.csv_io import load_csv


@pytest.fixture
def path(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'One'
    sheet.append(['猫と猫', 123, '=123+1', True, datetime(2026, 10, 15)])
    sheet['A2'] = '猫'
    sheet.merge_cells('A2:B2')
    sheet['A1'].font = Font(bold=True)
    book.create_sheet('Two').append(['別の猫', '00123', '引用,"と\n改行'])
    path = tmp_path / 'tools.xlsx'
    book.save(path)
    return path


def test_find_all_sheets_and_order(path):
    session = XlsxSession(path)
    matches = session.find('猫')
    assert [(name, r, c) for name, r, c, v in matches] == [('One', 0, 0), ('One', 1, 0), ('Two', 0, 0)]
    assert len(session.find('猫', all_sheets=False)) == 2
    assert not session.dirty


def test_find_formula_opt_in(path):
    session = XlsxSession(path)
    assert len(session.find('123')) == 2
    assert len(session.find('123', include_formulas=True)) == 3


def test_find_numeric_boolean_datetime(path):
    session = XlsxSession(path)
    assert session.find('TRUE')[0][2] == 3
    assert session.find('2026-10-15')[0][2] == 4


def test_empty_query_rejected(path):
    session = XlsxSession(path)
    with pytest.raises(ValueError):
        session.find('')
    with pytest.raises(ValueError):
        session.replace_all('', 'x')


def test_replace_across_sheets_one_history_action(path):
    session = XlsxSession(path)
    assert session.replace_all('猫', '犬') == (2, 3, 1)
    assert session.probe.book['One']['A1'].value == '犬と犬'
    assert session.probe.book['One']['A2'].value == '猫'
    assert session.probe.book['Two']['A1'].value == '別の犬'
    assert len(session.undo_stack) == 1
    session.history()
    assert not session.dirty
    session.history(True)
    assert session.probe.book['Two']['A1'].value == '別の犬'
    assert session.probe.book['One']['A1'].font.bold


def test_replace_current_sheet_only(path):
    session = XlsxSession(path)
    assert session.replace_all('猫', '犬', all_sheets=False) == (1, 2, 1)
    assert session.probe.book['Two']['A1'].value == '別の猫'


def test_replace_skips_types_and_formulas(path):
    session = XlsxSession(path)
    assert session.replace_all('123', '456') == (1, 1, 2)
    assert session.probe.book['Two']['B1'].value == '00456'
    assert session.probe.book['One']['B1'].value == 123
    assert session.probe.book['One']['C1'].value == '=123+1'


def test_failed_replace_rolls_back_all_sheets(path):
    book = load_workbook(path)
    book['One']['A1'] = 'prefix猫'
    book['Two']['A1'] = '猫'
    book.save(path)
    session = XlsxSession(path)
    with pytest.raises(ProtectedCellError):
        session.replace_all('猫', '=SUM(A1)')
    assert session.probe.book['One']['A1'].value == 'prefix猫'
    assert not session.dirty
    assert not session.undo_stack


def test_backend_failure_restores_prior_replacements(path):
    session = XlsxSession(path)
    # One!A1 becomes longer than the backend limit, so the entire operation fails.
    with pytest.raises(ValueError):
        session.replace_all('猫', 'x' * 32768)
    assert not session.dirty
    assert not session.undo_stack


def test_noop_replace_keeps_redo(path):
    session = XlsxSession(path)
    session.edit(0, 0, 'changed')
    session.history()
    assert session.replace_all('猫', '猫') == (0, 0, 1)
    assert session.redo_stack


def test_delete_text_roundtrips_and_undoes(path, tmp_path):
    session = XlsxSession(path)
    session.sheet_name = 'Two'
    session.replace_all('別の猫', '', all_sheets=False)
    assert session.sheet['A1'].value is None
    destination = tmp_path / 'cleared.xlsx'
    session.save_as(destination)
    assert load_workbook(destination)['Two']['A1'].value is None
    session.history()
    assert session.sheet['A1'].value == '別の猫'
    assert session.dirty


def test_cleared_last_cell_can_be_saved(tmp_path):
    path = tmp_path / 'single.xlsx'
    book = Workbook()
    book.active['C5'] = 'text'
    book.save(path)
    session = XlsxSession(path)
    session.replace_all('text', '')
    session.save()
    assert load_workbook(path).active['C5'].value is None


def test_readonly_can_search_and_export_but_not_replace(path, tmp_path):
    session = XlsxSession(path)
    session.probe.reasons.append('unsupported')
    assert session.find('猫')
    with pytest.raises(ProtectedCellError):
        session.replace_all('猫', '犬')
    session.sheet_name = 'Two'
    session.export_csv(tmp_path / 'readonly.csv')


def test_formula_export_requires_opt_in(path, tmp_path):
    session = XlsxSession(path)
    target = tmp_path / 'formula.csv'
    with pytest.raises(ValueError, match='数式セル'):
        session.export_csv(target)
    assert not target.exists()
    session.export_csv(target, formulas='text')
    values = load_csv(target, encoding='utf-8').rows[0]
    assert values[1:] == ['123', '=123+1', 'TRUE', '2026-10-15T00:00:00']


def test_csv_content_and_format(path, tmp_path):
    session = XlsxSession(path)
    session.sheet_name = 'Two'
    target = tmp_path / 'export.csv'
    session.export_csv(target)
    doc = load_csv(target, encoding='utf-8')
    assert doc.rows == [['別の猫', '00123', '引用,"と\n改行']]
    assert doc.format.bom and doc.format.newline == '\r\n' and doc.format.delimiter == ','


def test_export_preserves_source_dirty_path_and_history(path, tmp_path):
    before = path.read_bytes()
    session = XlsxSession(path)
    session.sheet_name = 'Two'
    session.edit(0, 0, '未保存')
    session.export_csv(tmp_path / 'unsaved.csv')
    assert session.dirty and len(session.undo_stack) == 1
    assert session.probe.path == path.resolve()
    assert path.read_bytes() == before
    assert load_csv(tmp_path / 'unsaved.csv', encoding='utf-8').rows[0][0] == '未保存'


def test_export_never_overwrites_existing_file(path, tmp_path):
    session = XlsxSession(path)
    session.sheet_name = 'Two'
    target = tmp_path / 'existing.csv'
    target.write_bytes(b'original')
    with pytest.raises(FileExistsError):
        session.export_csv(target)
    assert target.read_bytes() == b'original'


def test_export_failure_cleanup(path, tmp_path, monkeypatch):
    session = XlsxSession(path)
    session.sheet_name = 'Two'
    def fail(*a, **kw):
        raise OSError('simulated failure')
    monkeypatch.setattr('sheetdock.xlsx_gui.os.link', fail)
    target = tmp_path / 'failed.csv'
    with pytest.raises(OSError):
        session.export_csv(target)
    assert not target.exists()
    assert not list(tmp_path.glob('.failed.csv.*'))


def test_export_uses_entire_sheet_not_only_viewport(tmp_path):
    path = tmp_path / 'beyond.xlsx'
    book = Workbook()
    book.active['A1001'] = 'outside'
    book.save(path)
    session = XlsxSession(path)
    assert session.find('outside')[0][1] == 1000
    target = tmp_path / 'entire.csv'
    session.export_csv(target)
    assert load_csv(target, encoding='utf-8').rows[1000][0] == 'outside'


def test_special_formula_export_is_rejected(path, tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    session = XlsxSession(path)
    session.sheet['C1'] = ArrayFormula('C1:C2', '=B1:B2+1')
    with pytest.raises(ValueError, match='特殊数式'):
        session.export_csv(tmp_path / 'special.csv', formulas='text')
