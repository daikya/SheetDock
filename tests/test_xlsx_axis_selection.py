from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from openpyxl import Workbook

from sheetdock.xlsx_gui import XlsxPane, XlsxSession, selected_axis_range
from sheetdock.xlsx_io import snapshot


@pytest.mark.parametrize('boxes,rows,expected', [
    (((1, 0, 3, 4),), True, (2, 2)),
    (((0, 1, 4, 3),), False, (2, 2)),
    (((2, 0, 4, 1), (1, 0, 3, 1)), True, (2, 3)),
    (((0, 2, 1, 3), (0, 1, 1, 2)), False, (2, 2)),
])
def test_axis_selection_projects_and_merges_intervals(boxes, rows, expected):
    assert selected_axis_range(boxes, rows, 4) == expected


@pytest.mark.parametrize('boxes,rows', [
    ((), True), (((0, 0, 1, 1), (2, 0, 3, 1)), True),
    (((0, 0, 1, 1), (0, 2, 1, 3)), False),
    (((-1, 0, 2, 1),), True), (((0, 0, 5, 1),), True),
])
def test_axis_selection_rejects_empty_gaps_or_outside(boxes, rows):
    with pytest.raises(ValueError):
        selected_axis_range(boxes, rows, 4)


@pytest.fixture
def pane(tmp_path, monkeypatch):
    book = Workbook()
    for i in range(4):
        book.active.append([f'{i},{j}' for j in range(4)])
    book.active.row_dimensions[2].height = 40
    book.active.column_dimensions['B'].width = 30
    book.create_sheet('Other')['A1'] = 'unchanged'
    path = tmp_path / 'selection.xlsx'
    book.save(path)
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane.root = XlsxSession(path), None
    # Current cell can be at the end of a reverse selection.
    pane.selected_cell = lambda: pane.session.sheet['C3']
    pane.sheet_widget = SimpleNamespace(get_all_selection_boxes=lambda: ((1, 1, 3, 3),),
        select_cell=Mock(), see=Mock())
    pane.render, pane.selection_changed = Mock(), Mock()
    pane.ask, pane.confirm, pane.error, pane.info = Mock(), Mock(return_value=True), Mock(), Mock()
    for name, mock in [('askinteger', pane.ask), ('askyesno', pane.confirm),
                       ('showerror', pane.error), ('showinfo', pane.info)]:
        monkeypatch.setattr('sheetdock.dialogs.' + name, mock)
    return pane


@pytest.mark.parametrize('action', ['insert-rows', 'delete-rows', 'insert-columns', 'delete-columns'])
def test_multiple_selection_routes_start_count_and_one_history(pane, action, tmp_path):
    before = snapshot(pane.session.probe.book)
    pane.axis_operation(action)
    pane.error.assert_not_called()
    pane.ask.assert_not_called()
    assert len(pane.session.undo_stack) == 1 and pane.session.dirty
    if action.endswith('rows'):
        assert pane.session.sheet.max_row == (6 if action.startswith('insert') else 2)
        pane.sheet_widget.select_cell.assert_called_once_with(1, 2)
    else:
        assert pane.session.sheet.max_column == (6 if action.startswith('insert') else 2)
    after = snapshot(pane.session.probe.book)
    assert after[1] == before[1]
    pane.session.history()
    assert snapshot(pane.session.probe.book) == before and not pane.session.dirty
    pane.session.history(True)
    assert snapshot(pane.session.probe.book) == after
    target = tmp_path / 'saved.xlsx'
    pane.session.save_as(target)
    assert snapshot(XlsxSession(target).probe.book) == after


def test_multiple_delete_confirmation_names_selected_start_and_count(pane):
    pane.confirm.return_value = False
    pane.axis_operation('delete-rows')
    assert '2行目から2行' in pane.confirm.call_args.args[1]
    assert not pane.session.dirty and not pane.session.undo_stack
    pane.render.assert_not_called()


def test_disjoint_selection_does_not_prompt_or_mutate(pane):
    pane.sheet_widget.get_all_selection_boxes = lambda: ((0, 0, 1, 1), (2, 0, 3, 1))
    pane.axis_operation('delete-rows')
    pane.error.assert_called_once()
    assert '連続' in pane.error.call_args.args[1]
    pane.ask.assert_not_called()
    pane.confirm.assert_not_called()
    assert not pane.session.dirty and not pane.session.undo_stack


def test_select_all_cannot_delete_all_rows(pane):
    pane.sheet_widget.get_all_selection_boxes = lambda: ((0, 0, 4, 4),)
    pane.axis_operation('delete-rows')
    pane.info.assert_called_once()
    pane.confirm.assert_not_called()
    assert not pane.session.dirty


def test_selected_range_over_insert_capacity_is_not_truncated(pane):
    pane.sheet_widget.get_all_selection_boxes = lambda: ((0, 0, 4, 4),)
    pane.session.sheet.cell(999, 1, 'limit')
    before = snapshot(pane.session.probe.book)
    pane.axis_operation('insert-rows')
    pane.info.assert_called_once()
    pane.ask.assert_not_called()
    assert snapshot(pane.session.probe.book) == before and not pane.session.undo_stack


def test_single_axis_selection_keeps_count_dialog(pane):
    pane.sheet_widget.get_all_selection_boxes = lambda: ((1, 1, 2, 3),)
    pane.ask.return_value = None
    pane.axis_operation('insert-rows')
    pane.ask.assert_called_once()
    assert '2行目' in pane.ask.call_args.args[1]
    assert not pane.session.dirty


def test_protection_precedes_multiple_selection(pane):
    pane.session.probe.book['Other']['A1'] = '=1+1'
    before = snapshot(pane.session.probe.book)
    pane.axis_operation('delete-columns')
    pane.info.assert_called_once()
    pane.confirm.assert_not_called()
    assert snapshot(pane.session.probe.book) == before and not pane.session.undo_stack
