from copy import copy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from datetime import datetime
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from sheetdock.xlsx_gui import XlsxSession as AxisSession, XlsxPane
from sheetdock.xlsx_io import ProtectedCellError, snapshot


def make_samples(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    targets = [folder / f'axis_{name}.xlsx' for name in ('basic', 'formula', 'merged', 'readonly')]
    if any(path.exists() for path in targets):
        raise FileExistsError('同名サンプルが存在します。既存ファイルは上書きしません')
    for name, target in zip(('basic', 'formula', 'merged', 'readonly'), targets):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Data'
        for row in (['コード', '数量', '日付', '確認'], ['00123', 3, datetime(2026, 10, 16), True],
                    ['00456', 8, datetime(2026, 10, 17), False], ['00789', 12, datetime(2026, 10, 18), True]):
            sheet.append(row)
        sheet['A2'].font = Font(bold=True, color='FF2563EB')
        sheet['A2'].fill = PatternFill('solid', fgColor='FFFDE68A')
        sheet.column_dimensions['A'].width = 24
        sheet.column_dimensions['B'].width = 14
        sheet.column_dimensions['C'].width = 22
        sheet.row_dimensions[2].height = 32
        notes = book.create_sheet('Notes')
        notes['A1'] = '他シートは変更しない'
        if name == 'formula':
            notes['A2'] = '=Data!B2'
        elif name == 'merged':
            sheet.merge_cells('A1:B1')
        elif name == 'readonly':
            rule = DataValidation(type='list', formula1='"はい,いいえ"')
            sheet.add_data_validation(rule)
            rule.add('D2')
        with target.open('xb') as handle:
            book.save(handle)
        book.close()



@pytest.fixture
def samples(tmp_path):
    make_samples(tmp_path)
    return tmp_path


@pytest.fixture
def session(samples):
    return AxisSession(samples / 'axis_basic.xlsx')


@pytest.mark.parametrize('action', ['insert-rows', 'delete-rows', 'insert-columns', 'delete-columns'])
def test_operation_undo_redo_and_roundtrip(session, tmp_path, action):
    original = session.probe.path.read_bytes()
    before = snapshot(session.probe.book)
    session.axis_operation(action, 2, 2)
    after = snapshot(session.probe.book)
    assert before != after and session.dirty
    session.history()
    assert snapshot(session.probe.book) == before and not session.dirty
    session.history(True)
    assert snapshot(session.probe.book) == after
    target = tmp_path / 'saved.xlsx'
    session.save_as(target)
    assert not session.dirty
    assert snapshot(load_workbook(target)) == after
    assert session.probe.path == target.resolve()
    assert original == (tmp_path / 'axis_basic.xlsx').read_bytes()


def test_row_insert_moves_values_styles_heights_and_keeps_other_sheet(session):
    style = copy(session.sheet['A2']._style)
    notes = snapshot(session.probe.book)[1]
    session.axis_operation('insert-rows', 2, 2)
    assert session.sheet['A4'].value == '00123'
    assert session.sheet['A4']._style == style
    assert session.sheet['C4'].value.year == 2026
    assert session.sheet['D4'].value is True
    assert session.sheet['A2'].value is None
    assert session.sheet.row_dimensions[4].height == 32
    assert 2 not in session.sheet.row_dimensions
    assert session.sheet.column_dimensions['A'].width == 24
    assert snapshot(session.probe.book)[1] == notes


def test_column_insert_moves_widths_styles_and_row_height(session):
    style = copy(session.sheet['A2']._style)
    session.axis_operation('insert-columns', 1, 2)
    assert session.sheet['C2'].value == '00123'
    assert session.sheet['C2']._style == style
    assert session.sheet['D2'].value == 3
    assert session.sheet.column_dimensions['C'].width == 24
    assert session.sheet.column_dimensions['D'].width == 14
    assert session.sheet.column_dimensions['E'].width == 22
    assert 'A' not in session.sheet.column_dimensions
    assert session.sheet.row_dimensions[2].height == 32


def test_delete_columns_removes_only_selected_widths_and_cells(session):
    session.axis_operation('delete-columns', 2, 2)
    assert session.sheet['A2'].value == '00123'
    assert session.sheet['B2'].value is True
    assert list(session.sheet.column_dimensions) == ['A']
    assert session.sheet.column_dimensions['A'].width == 24


def test_delete_rows_removes_only_selected_height_and_cells(session):
    session.axis_operation('delete-rows', 2)
    assert session.sheet['A2'].value == '00456'
    assert 2 not in session.sheet.row_dimensions


@pytest.mark.parametrize('name', ['formula', 'merged', 'readonly'])
def test_unsafe_workbook_is_rejected_without_changes(samples, name):
    session = AxisSession(samples / f'axis_{name}.xlsx')
    before = snapshot(session.probe.book)
    with pytest.raises(ProtectedCellError):
        session.axis_operation('insert-rows', 2)
    assert snapshot(session.probe.book) == before
    assert not session.dirty and not session.undo_stack


@pytest.mark.parametrize('action,index,count', [
    ('unknown', 2, 1), ('insert-rows', 0, 1), ('insert-rows', 5, 1),
    ('insert-columns', True, 1), ('insert-rows', 1, 0), ('insert-rows', 1, True),
    ('insert-rows', 1, 101), ('delete-rows', 1, 4), ('delete-columns', 4, 2),
])
def test_bad_requests_leave_editor_unchanged(session, action, index, count):
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError):
        session.axis_operation(action, index, count)
    assert snapshot(session.probe.book) == before and not session.undo_stack


@pytest.mark.parametrize('case', ['hidden-row', 'hidden-column', 'group', 'span', 'zero-width', 'zero-height'])
def test_unsupported_dimensions_rejected(session, case):
    if case == 'hidden-row':
        session.sheet.row_dimensions[2].hidden = True
    elif case == 'hidden-column':
        session.sheet.column_dimensions['A'].hidden = True
    elif case == 'group':
        session.sheet.row_dimensions[2].outlineLevel = 1
    elif case == 'span':
        session.sheet.column_dimensions['A'].max = 2
    elif case == 'zero-height':
        session.sheet.row_dimensions[2].height = 0
    else:
        session.sheet.column_dimensions['A'].width = 0
    before = snapshot(session.probe.book)
    with pytest.raises(ProtectedCellError):
        session.axis_operation('insert-rows', 2)
    assert snapshot(session.probe.book) == before


def test_failed_validation_does_not_publish_candidate(session, monkeypatch):
    original_book = session.probe.book
    original = snapshot(original_book)
    monkeypatch.setattr('sheetdock.xlsx_axis.inspect_package', lambda data: ['unsupported'])
    with pytest.raises(ValueError):
        session.axis_operation('insert-rows', 2)
    assert session.probe.book is original_book
    assert snapshot(original_book) == original
    assert not session.undo_stack and not session.dirty


def test_structural_history_interleaves_with_existing_cell_and_width_history(session):
    before = snapshot(session.probe.book)
    session.edit(1, 0, 'edited')
    session.resize_columns({0: 220})
    session.axis_operation('insert-rows', 2)
    session.axis_operation('delete-columns', 2)
    after = snapshot(session.probe.book)
    for _ in range(4):
        assert session.history()
    assert snapshot(session.probe.book) == before
    for _ in range(4):
        assert session.history(True)
    assert snapshot(session.probe.book) == after


def test_new_axis_operation_clears_redo(session):
    session.axis_operation('insert-rows', 2)
    session.history()
    assert session.redo_stack
    session.axis_operation('insert-columns', 2)
    assert not session.redo_stack


def test_display_limit_rejected_before_modification(session):
    session.sheet['A1000'] = 'tail'
    before = snapshot(session.probe.book)
    with pytest.raises(ValueError, match='1000行'):
        session.axis_operation('insert-rows', 2)
    assert snapshot(session.probe.book) == before


def test_outside_view_column_dimension_is_shifted_without_losing_width(session):
    session.sheet.column_dimensions['Z'].width = 20
    session.axis_operation('insert-columns', 2, 2)
    assert session.sheet.column_dimensions['AB'].width == 20
    assert 'Z' not in session.sheet.column_dimensions


def test_metadata_beyond_display_limit_is_rejected_before_shift(session):
    session.sheet.column_dimensions['XFD'].width = 20
    with pytest.raises(ValueError, match='1000行'):
        session.axis_operation('insert-columns', 2)
    assert 'XFD' in session.sheet.column_dimensions and not session.undo_stack








def test_merge_on_untouched_other_sheet_is_preserved(session):
    session.probe.book['Notes'].merge_cells('A1:B1')
    before = snapshot(session.probe.book)[1]
    session.axis_operation('insert-rows', 2)
    assert snapshot(session.probe.book)[1] == before


def test_operation_cannot_exceed_workbook_cell_limit(session):
    session.sheet['J1000'] = 'tail'
    session.probe.book['Notes']['KN300'] = 'other tail'
    with pytest.raises(ValueError, match='10万セル'):
        session.axis_operation('insert-columns', 2)
    assert session.sheet['A2'].value == '00123' and not session.undo_stack


@pytest.fixture
def pane(session, monkeypatch):
    from types import SimpleNamespace
    pane = XlsxPane.__new__(XlsxPane)
    pane.session, pane.root = session, None
    pane.selected_cell = lambda: session.sheet['A2']
    pane.render_calls, pane.selections, pane.errors = [], [], []
    pane.render = lambda: pane.render_calls.append(True)
    pane.selection_changed = lambda: None
    pane.sheet_widget = SimpleNamespace(select_cell=lambda *position: pane.selections.append(position),
        see=lambda *args: None, get_all_selection_boxes=lambda: ((1, 0, 2, 1),))
    monkeypatch.setattr('sheetdock.xlsx_gui.simpledialog.askinteger', lambda *args, **kwargs: 2)
    monkeypatch.setattr('sheetdock.xlsx_gui.messagebox.askyesno', lambda *args, **kwargs: True)
    monkeypatch.setattr('sheetdock.xlsx_gui.messagebox.showerror', lambda *args, **kwargs: pane.errors.append(args))
    monkeypatch.setattr('sheetdock.xlsx_gui.messagebox.showinfo', lambda *args, **kwargs: None)
    return pane


@pytest.mark.parametrize('action', ['insert-rows', 'delete-rows', 'insert-columns', 'delete-columns'])
def test_gui_routes_operation_count_and_preserves_history(pane, action):
    before = snapshot(pane.session.probe.book)
    pane.axis_operation(action)
    assert pane.render_calls and pane.selections and not pane.errors
    assert pane.session.dirty and len(pane.session.undo_stack) == 1
    pane.session.history()
    assert snapshot(pane.session.probe.book) == before


def test_gui_cancel_count_leaves_document_and_history_unchanged(pane, monkeypatch):
    monkeypatch.setattr('sheetdock.xlsx_gui.simpledialog.askinteger', lambda *args, **kwargs: None)
    pane.axis_operation('insert-rows')
    assert not pane.session.dirty and not pane.render_calls and not pane.session.undo_stack


def test_gui_delete_requires_confirmation_and_cancel_preserves_data(pane, monkeypatch):
    questions = []
    def reject(*args, **kwargs):
        questions.append(args)
        return False
    monkeypatch.setattr('sheetdock.xlsx_gui.messagebox.askyesno', reject)
    pane.axis_operation('delete-rows')
    assert questions and '2行目から2行' in questions[0][1]
    assert not pane.session.dirty and not pane.render_calls


def test_gui_missing_selection_does_not_prompt_for_count(pane, monkeypatch):
    pane.selected_cell = lambda: None
    def unexpected(*args, **kwargs):
        pytest.fail('No count dialog should appear')
    monkeypatch.setattr('sheetdock.xlsx_gui.simpledialog.askinteger', unexpected)
    pane.axis_operation('insert-rows')
    assert not pane.session.dirty and not pane.errors


def test_gui_rechecks_protection_at_execution(pane):
    pane.session.probe.book['Notes']['A2'] = '=Data!A2'
    before = snapshot(pane.session.probe.book)
    pane.axis_operation('insert-rows')
    assert snapshot(pane.session.probe.book) == before and not pane.render_calls


def test_gui_invalid_count_is_reported_without_changes(pane, monkeypatch):
    monkeypatch.setattr('sheetdock.xlsx_gui.simpledialog.askinteger', lambda *args, **kwargs: 100)
    pane.axis_operation('delete-rows')
    assert pane.errors and not pane.session.dirty and not pane.render_calls


def test_gui_buttons_follow_cached_reason_and_selection(pane):
    from types import SimpleNamespace
    states = []
    pane.axis_buttons = {'insert-rows': SimpleNamespace(configure=lambda **kwargs: states.append(kwargs['state']))}
    pane._axis_block_reason = ''
    pane.update_axis_commands()
    assert states[-1] == 'normal'
    pane._axis_block_reason = '数式あり'
    pane.update_axis_commands()
    assert states[-1] == 'disabled'
    pane._axis_block_reason = ''
    pane.selected_cell = lambda: None
    pane.update_axis_commands()
    assert states[-1] == 'disabled'
