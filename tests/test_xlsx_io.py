from datetime import datetime
from zipfile import ZipFile
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from sheetdock.xlsx_io import ProtectedCellError, XlsxProbe, snapshot


@pytest.fixture
def source(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = '売上'
    for address, value in {'A1': '00123', 'B1': 12.5, 'C1': True,
                           'D1': datetime(2026, 10, 15), 'E1': '=B1*2',
                           'A2': '日本語\n改行', 'B2': '結合'}.items():
        sheet[address] = value
    sheet['A1'].font = Font(name='Meiryo', bold=True, color='FF123456')
    sheet['A1'].fill = PatternFill('solid', fgColor='FF00FF00')
    sheet['A1'].alignment = Alignment(wrap_text=True)
    sheet['B1'].number_format = '0.00'
    sheet.column_dimensions['A'].width = 25
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells('B2:C2')
    book.create_sheet('補足')
    path = tmp_path / 'source.xlsx'
    book.save(path)
    return path


def test_roundtrip(source, tmp_path):
    before = source.read_bytes()
    probe = XlsxProbe(source)
    assert not probe.read_only, probe.reasons
    expected = snapshot(probe.book)
    destination = tmp_path / 'roundtrip.xlsx'
    probe.save_as(destination)
    reopened = load_workbook(destination)
    assert snapshot(reopened) == expected
    assert reopened['売上'].column_dimensions['A'].width == 25
    assert reopened['売上'].row_dimensions[1].height == 30
    assert reopened['売上']['E1'].value == '=B1*2'
    assert source.read_bytes() == before


@pytest.mark.parametrize('address,value', [('A1', '00045'), ('B1', 42),
                                         ('C1', False), ('D1', datetime(2027, 1, 1)),
                                         ('A2', '#N/A'), ('F1', '新規')])
def test_typed_edits(source, tmp_path, address, value):
    probe = XlsxProbe(source)
    probe.set_value('売上', address, value)
    target = tmp_path / 'edited.xlsx'
    probe.save_as(target)
    cell = load_workbook(target)['売上'][address]
    assert cell.value == value
    if isinstance(value, str):
        assert cell.data_type == 's'


@pytest.mark.parametrize('address,value', [('E1', 'delete'), ('B2', 'change'),
                                         ('C2', 'change'), ('A1', '=1+1')])
def test_protected_cells(source, address, value):
    probe = XlsxProbe(source)
    before = snapshot(probe.book)
    with pytest.raises(ProtectedCellError):
        probe.set_value('売上', address, value)
    assert snapshot(probe.book) == before


@pytest.mark.parametrize('address,value', [('A1', 123), ('B1', '12'), ('C1', 1),
                                         ('D1', '2026-10-15')])
def test_type_change_rejected(source, address, value):
    with pytest.raises(TypeError):
        XlsxProbe(source).set_value('売上', address, value)


@pytest.mark.parametrize(
    'value', ['x\x00y', 'x' * 32768, float('nan'), float('inf')],
    ids=['illegal-control-character', 'overlong-string', 'nan', 'infinity'],
)
def test_lossy_values_rejected(source, value):
    probe = XlsxProbe(source)
    with pytest.raises(ValueError):
        probe.set_value('売上', 'F1', value)
    assert probe.book['売上']['F1'].value is None


def test_existing_destination_preserved(source):
    before = source.read_bytes()
    with pytest.raises(FileExistsError):
        XlsxProbe(source).save_as(source)
    assert source.read_bytes() == before


def test_data_validation_readonly(source, tmp_path):
    book = load_workbook(source)
    rule = DataValidation(type='list', formula1='"A,B"')
    book['売上'].add_data_validation(rule)
    rule.add('A1')
    book.save(source)
    probe = XlsxProbe(source)
    assert probe.read_only
    with pytest.raises(PermissionError):
        probe.set_value('売上', 'A1', 'edit')
    with pytest.raises(PermissionError):
        probe.save_as(tmp_path / 'blocked.xlsx')
    assert not (tmp_path / 'blocked.xlsx').exists()


@pytest.mark.parametrize('part', ['xl/drawings/unknown.xml', 'customXml/item1.xml',
                                'xl/vbaProject.bin'])
def test_unknown_part_readonly(source, part):
    with ZipFile(source, 'a') as archive:
        archive.writestr(part, b'<unknown/>')
    assert XlsxProbe(source).read_only


def test_unsupported_extension(tmp_path):
    with pytest.raises(ValueError):
        XlsxProbe(tmp_path / 'macro.xlsm')


def test_invalid_xlsx(tmp_path):
    path = tmp_path / 'broken.xlsx'
    path.write_text('not a zip')
    with pytest.raises(Exception):
        XlsxProbe(path)


def test_save_failure_preserves_source(source, tmp_path, monkeypatch):
    probe = XlsxProbe(source)
    before = source.read_bytes()
    def fail(*args):
        raise OSError('simulated failure')
    monkeypatch.setattr(probe.book, 'save', fail)
    target = tmp_path / 'failure.xlsx'
    with pytest.raises(OSError):
        probe.save_as(target)
    assert not target.exists()
    assert source.read_bytes() == before
