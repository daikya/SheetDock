"""Validated row/column operations for the deliberately limited XLSX subset."""
from copy import copy
from io import BytesIO
import math

from openpyxl import load_workbook
from openpyxl.worksheet.dimensions import ColumnDimension, RowDimension
from openpyxl.utils import get_column_letter, column_index_from_string, coordinate_to_tuple
from .xlsx_io import ProtectedCellError, snapshot, inspect_package
from .xlsx_extent import extent, extents
from .xlsx_widths import stored_width

VIEW_ROWS, VIEW_COLUMNS = 1000, 50
HISTORY_LIMIT = 100
ACTIONS = ('insert-rows', 'delete-rows', 'insert-columns', 'delete-columns')


def shifted(index, at, count, deleting):
    if deleting and at <= index < at + count:
        return None
    return index - count if deleting and index >= at + count else (
        index + count if not deleting and index >= at else index)


def dimensions(sheet, axis, at, count, deleting):
    source = sheet.row_dimensions if axis == 'rows' else sheet.column_dimensions
    result = []
    for key, original in source.items():
        index = key if axis == 'rows' else column_index_from_string(key)
        target = shifted(index, at, count, deleting)
        if target is None:
            continue
        dimension = copy(original)
        dimension.parent = sheet
        if axis == 'rows':
            dimension.index = target
            result.append((target, dimension))
        else:
            dimension.index = get_column_letter(target)
            dimension.min = dimension.max = target
            result.append((dimension.index, dimension))
    source.clear()
    for key, dimension in result:
        source[key] = dimension


def shifted_cells(records, axis, at, count, deleting):
    result = []
    for record in records:
        row, column = coordinate_to_tuple(record[0])
        target = shifted(row if axis == 'rows' else column, at, count, deleting)
        if target is None:
            continue
        if axis == 'rows':
            row = target
        else:
            column = target
        result.append((row, column, (f'{get_column_letter(column)}{row}', *record[1:])))
    return tuple(item[2] for item in sorted(result, key=lambda item: item[:2]))


class AxisOperations:
    def axis_reason(self):
        if self.probe.read_only:
            return '読み取り専用のブックは行列操作できません'
        if any(cell.data_type == 'f' for sheet in self.probe.book for row in sheet for cell in row):
            return '別シートを含め数式があるブックは行列操作できません'
        if self.sheet.merged_cells.ranges:
            return '対象シートに結合セルがある場合は行列操作できません'
        if extent(self.sheet)[0] > VIEW_ROWS or extent(self.sheet)[1] > VIEW_COLUMNS:
            return '行列操作は対象シート全体が1000行×50列以内の場合だけです'
        for name, holder in (('rows', self.sheet.row_dimensions), ('columns', self.sheet.column_dimensions)):
            for key, dimension in holder.items():
                if dimension.hidden or dimension.outlineLevel or dimension.collapsed:
                    return '非表示・グループ化された行列は対象外です'
                if name == 'rows' and dimension.height == 0:
                    return '高さ0の行は対象外です'
                if name == 'columns':
                    position = column_index_from_string(key)
                    if ((dimension.min and dimension.min != position)
                            or (dimension.max and dimension.max != position) or dimension.width == 0):
                        return '複数列の定義・幅0・重複した列定義は対象外です'
        return ''

    def axis_operation(self, action, index, count=1):
        if action not in ACTIONS:
            raise ValueError('行列操作を指定してください')
        reason = self.axis_reason()
        if reason:
            raise ProtectedCellError(reason)
        axis = 'rows' if action.endswith('rows') else 'columns'
        deleting = action.startswith('delete')
        size = extent(self.sheet)[0 if axis == 'rows' else 1]
        if type(index) is not int or not 1 <= index <= size:
            raise ValueError('位置は既存範囲内の1始まりで指定してください（末尾への追加は今回対象外）')
        if type(count) is not int or not 1 <= count <= 100:
            raise ValueError('件数は1～100で指定してください')
        if deleting and (index + count - 1 > size or count >= size):
            raise ValueError('削除範囲を既存範囲内に収め、少なくとも1行・1列を残してください')
        if extent(self.sheet)[0] > VIEW_ROWS or extent(self.sheet)[1] > VIEW_COLUMNS:
            raise ValueError('今回の行列操作は対象シート全体が1000行×50列以内の場合だけです')
        new_rows = extent(self.sheet)[0] + (0 if axis != 'rows' else (-count if deleting else count))
        new_columns = extent(self.sheet)[1] + (0 if axis != 'columns' else (-count if deleting else count))
        if new_rows > VIEW_ROWS or new_columns > VIEW_COLUMNS:
            raise ValueError('操作後の範囲が1000行×50列を超えます')
        area = sum(extent(sheet)[0] * extent(sheet)[1] for sheet in self.probe.book if sheet is not self.sheet)
        if area + new_rows * new_columns > 100_000:
            raise ValueError('操作後のブックが10万セルを超えます')
        for name, holder, maximum in (
                ('rows', self.sheet.row_dimensions, 1_048_576),
                ('columns', self.sheet.column_dimensions, 16_384)):
            for key, dimension in holder.items():
                position = key if name == 'rows' else column_index_from_string(key)
                if dimension.hidden or dimension.outlineLevel or dimension.collapsed:
                    raise ProtectedCellError('非表示・グループ化された行列は今回対象外です')
                if name == 'rows' and dimension.height == 0:
                    raise ProtectedCellError('高さ0の行は今回対象外です')
                if name == 'columns' and (
                        (dimension.min and dimension.min != position)
                        or (dimension.max and dimension.max != position)
                        or dimension.width == 0):
                    raise ProtectedCellError('複数列の定義・幅0・重複した列定義は今回対象外です')
                target = shifted(position, index, count, deleting) if name == axis else position
                if target is not None and not 1 <= target <= maximum:
                    raise ValueError('移動先の行列定義がExcelの範囲を超えます')
        before = self.pack_book()
        baseline = snapshot(self.probe.book)
        candidate = load_workbook(BytesIO(before), data_only=False)
        try:
            sheet = candidate[self.sheet_name]
            getattr(sheet, action.replace('-', '_').replace('columns', 'cols'))(index, count)
            dimensions(sheet, axis, index, count, deleting)
            actual = snapshot(candidate)
            active = self.probe.book.sheetnames.index(self.sheet_name)
            expected_cells = shifted_cells(baseline[active][5], axis, index, count, deleting)
            if actual[active][5] != expected_cells:
                raise ValueError('値・型・セル書式の移動結果が期待と一致しません')
            if any(before_sheet != after_sheet for i, (before_sheet, after_sheet)
                   in enumerate(zip(baseline, actual)) if i != active):
                raise ValueError('他シートが変化したため操作を中止しました')
            if actual == baseline:
                return False
            if extent(sheet) != (new_rows, new_columns):
                raise ValueError('行列操作後の表示範囲が期待と一致しません')
            stream = BytesIO()
            candidate.save(stream)
            after = stream.getvalue()
            if inspect_package(after):
                raise ValueError('操作後に未対応の構造が含まれました')
            checked = load_workbook(BytesIO(after), data_only=False)
            try:
                if snapshot(checked) != actual or extents(checked) != extents(candidate):
                    raise ValueError('保存・再読込後の値・型・書式・行列定義が一致しません')
            finally:
                checked.close()
            self.probe.book.close()
            self.probe.book = candidate
            candidate = None
            self.undo_stack.append((self.sheet_name, {'before': before, 'after': after,
                'before_name': self.sheet_name, 'after_name': self.sheet_name}))
            del self.undo_stack[:-HISTORY_LIMIT]
            self.redo_stack.clear()
            return True
        finally:
            if candidate is not None:
                candidate.close()


    def append_axis(self, axis, count=1, digit_width=7):
        if axis not in {'rows', 'columns'}:
            raise ValueError('rowsまたはcolumnsを指定してください')
        if type(count) is not int or not 1 <= count <= 100:
            raise ValueError('追加件数は1～100で指定してください')
        reason = self.axis_reason()
        if reason:
            raise ProtectedCellError(reason)
        rows, columns = extent(self.sheet)
        if rows > VIEW_ROWS or columns > VIEW_COLUMNS:
            raise ValueError('表示範囲が1000行×50列を超えるシートは対象外です')
        if self.sheet.sheet_format.zeroHeight:
            raise ProtectedCellError('標準行が非表示のシートは対象外です')
        target_rows = rows + (count if axis == 'rows' else 0)
        target_columns = columns + (count if axis == 'columns' else 0)
        if target_rows > VIEW_ROWS or target_columns > VIEW_COLUMNS:
            raise ValueError('追加後の表示範囲が1000行×50列を超えます')
        area = sum(r * c for name, (r, c) in extents(self.probe.book) if name != self.sheet_name)
        if area + target_rows * target_columns > 100_000:
            raise ValueError('追加後の表示範囲を含めたブックが10万セルを超えます')
        default = (self.sheet.sheet_format.defaultRowHeight if axis == 'rows'
                   else self.sheet.sheet_format.defaultColWidth)
        if default is None:
            default = 15 if axis == 'rows' else stored_width(self.sheet, columns + 1, digit_width)
        if not math.isfinite(default) or not 0 < default <= (409.5 if axis == 'rows' else 255):
            raise ValueError('標準の行高または列幅が今回の検証範囲外です')

        before = self.pack_book()
        baseline = snapshot(self.probe.book)
        candidate = load_workbook(BytesIO(before), data_only=False)
        try:
            sheet = candidate[self.sheet_name]
            if axis == 'rows':
                for index in range(rows + 1, target_rows + 1):
                    sheet.row_dimensions[index] = RowDimension(sheet, index=index, height=default)
            else:
                for index in range(columns + 1, target_columns + 1):
                    name = get_column_letter(index)
                    sheet.column_dimensions[name] = ColumnDimension(sheet, index=name,
                        min=index, max=index, width=default)
            actual = snapshot(candidate)
            active = candidate.sheetnames.index(self.sheet_name)
            dimension_index = 3 if axis == 'rows' else 2
            for i, (old, new) in enumerate(zip(baseline, actual)):
                if i != active:
                    if old != new:
                        raise ValueError('他シートが変化したため中止しました')
                else:
                    if any(old[j] != new[j] for j in range(6) if j != dimension_index):
                        raise ValueError('既存の値・型・書式・結合等が変化したため中止しました')
                    old_dims, new_dims = dict(old[dimension_index]), dict(new[dimension_index])
                    if (len(new_dims) != len(old_dims) + count
                            or any(new_dims.get(key) != value for key, value in old_dims.items())):
                        raise ValueError('既存の行列定義が変化したため中止しました')
            if extent(sheet) != (target_rows, target_columns):
                raise ValueError('追加後の表示範囲が一致しません')
            stream = BytesIO()
            candidate.save(stream)
            after = stream.getvalue()
            if inspect_package(after):
                raise ValueError('追加後に未対応構造が含まれました')
            checked = load_workbook(BytesIO(after), data_only=False)
            try:
                if snapshot(checked) != actual or extents(checked) != extents(candidate):
                    raise ValueError('保存・再読込後の内容または表示範囲が一致しません')
            finally:
                checked.close()
            self.probe.book.close()
            self.probe.book, candidate = candidate, None
            self.undo_stack.append((self.sheet_name, {'before': before, 'after': after,
                'before_name': self.sheet_name, 'after_name': self.sheet_name}))
            del self.undo_stack[:-HISTORY_LIMIT]
            self.redo_stack.clear()
            return True
        finally:
            if candidate is not None:
                candidate.close()

