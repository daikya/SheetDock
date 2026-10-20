"""XLSX column dimensions and view pixels; no settings/JSON dependency."""
from copy import copy
import math

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

MIN_PIXELS = 20
MAX_EXCEL_WIDTH = 255


def width_to_pixels(width, digit_width=7):
    if not math.isfinite(width) or width < 0 or digit_width <= 0:
        raise ValueError('列幅が不正です')
    return math.floor((256 * width + math.floor(128 / digit_width)) * digit_width / 256)


def pixels_to_width(pixels, digit_width=7):
    if type(pixels) not in (int, float) or not math.isfinite(pixels) or digit_width <= 0:
        raise ValueError('列幅が不正です')
    pixels = round(pixels)
    if not MIN_PIXELS <= pixels <= width_to_pixels(MAX_EXCEL_WIDTH, digit_width):
        raise ValueError('列幅は20ピクセル以上、Excelの幅255以下で指定してください')
    units = round(256 * (pixels + 0.5) / digit_width - math.floor(128 / digit_width))
    return min(MAX_EXCEL_WIDTH, units / 256)


def covering_dimensions(sheet, column):
    result = []
    for key, dimension in sheet.column_dimensions.items():
        start = dimension.min or column_index_from_string(key)
        end = dimension.max or start
        if start <= column <= end:
            result.append((key, dimension, start, end))
    return result


def stored_width(sheet, column, digit_width=7):
    dimensions = covering_dimensions(sheet, column)
    if dimensions:
        return dimensions[-1][1].width
    if sheet.sheet_format.defaultColWidth is not None:
        return sheet.sheet_format.defaultColWidth
    # baseColWidth counts characters without the five-pixel margins/grid.
    return math.floor((sheet.sheet_format.baseColWidth * digit_width + 5) / digit_width * 256) / 256


def display_widths(sheet, count, digit_width=7):
    maximum = width_to_pixels(MAX_EXCEL_WIDTH, digit_width)
    return [max(MIN_PIXELS, min(maximum, width_to_pixels(stored_width(sheet, i + 1, digit_width), digit_width)))
            for i in range(count)]


def dimension_state(sheet):
    return [(key, copy(dimension)) for key, dimension in sheet.column_dimensions.items()]


def restore_dimensions(sheet, state):
    sheet.column_dimensions.clear()
    for key, dimension in state:
        restored = copy(dimension)
        restored.parent = sheet
        sheet.column_dimensions[key] = restored


def set_width(sheet, column, width):
    """Split an existing span without changing widths/flags/styles of neighbours."""
    dimensions = covering_dimensions(sheet, column)
    if len(dimensions) > 1:
        raise ValueError('列定義が重複しているため、この列幅は変更できません')
    if not dimensions:
        if stored_width(sheet, column) == 0:
            raise ValueError('幅0の列は変更できません')
        name = get_column_letter(column)
        sheet.column_dimensions[name] = ColumnDimension(sheet, index=name, min=column,
                                                        max=column, width=width)
        return
    key, original, start, end = dimensions[0]
    if original.hidden or original.width == 0:
        raise ValueError('非表示または幅0の列は変更できません')
    del sheet.column_dimensions[key]
    for low, high in [(start, column - 1), (column, column), (column + 1, end)]:
        if low > high:
            continue
        dimension = copy(original)
        dimension.index = get_column_letter(low)
        dimension.min, dimension.max = low, high
        if low == high == column:
            dimension.width, dimension.bestFit = width, False
        sheet.column_dimensions[dimension.index] = dimension
