"""XLSX session and embedded workbook editor for SheetDock file tabs."""
from __future__ import annotations

from datetime import date, datetime, time, timedelta
from copy import copy
import csv
from io import BytesIO, StringIO
import os
import tempfile
import re
from pathlib import Path
import tkinter as tk
from tkinter import font as tkfont
from tkinter import filedialog, ttk
from . import dialogs as messagebox
from . import dialogs as simpledialog

from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
from tksheet import Sheet
from .xlsx_io import ProtectedCellError, snapshot
from .storage import ExternalChangeError, fingerprint, save_lock
from .csv_io import CsvDocument, CsvFormat, save_csv
from .xlsx_storage import load_checked, save_checked
from .xlsx_axis import AxisOperations, VIEW_ROWS, VIEW_COLUMNS, HISTORY_LIMIT
from .xlsx_extent import extent, data_extent, check_area
from .xlsx_inline import InlineEditing, settled
from . import xlsx_heights
from .heights import accumulate_row_resize
from .xlsx_tabs import SheetTabs
from .xlsx_search import XlsxSearch
from .xlsx_display import raw_value as display_value, format_value, cell_alignment
from .xlsx_widths import (display_widths, pixels_to_width, width_to_pixels,
                          dimension_state, restore_dimensions, set_width, MIN_PIXELS, MAX_EXCEL_WIDTH)


_DEFAULT_EXPECTED = object()
# All model changes go through the controlled editor, never widget mutations.
VIEW_BINDINGS = ('single_select', 'drag_select', 'row_select', 'column_select',
                 'arrowkeys', 'copy', 'select_all', 'column_width_resize',
                 'row_height_resize', 'edit_cell')


def selected_axis_range(boxes, rows, size):
    """Project selection boxes onto the operated axis; require one interval."""
    intervals = []
    for r1, c1, r2, c2 in boxes:
        start, end = (r1, r2) if rows else (c1, c2)
        if not (0 <= start < end <= size):
            raise ValueError('選択範囲が現在の行列範囲外です。選択し直してください')
        intervals.append((start, end))
    if not intervals:
        raise ValueError('行列操作の対象範囲を選択してください')
    intervals.sort()
    first, last = intervals[0]
    for start, end in intervals[1:]:
        if start > last:
            raise ValueError('離れた行・列は一度に操作できません。連続した範囲を選択してください')
        last = max(last, end)
    return first + 1, last - first


def parse_input(text, previous):
    """Preserve existing type; new/blank cells accept literal strings only."""
    if previous is None or isinstance(previous, str):
        return text
    if isinstance(previous, bool):
        if text.strip().upper() not in {'TRUE', 'FALSE'}:
            raise ValueError('真偽値はTRUEまたはFALSEで入力してください')
        return text.strip().upper() == 'TRUE'
    if isinstance(previous, (int, float)):
        try:
            return int(text) if text.strip().lstrip('+-').isdigit() else float(text)
        except ValueError:
            raise ValueError('数値を入力してください') from None
    if isinstance(previous, (datetime, date, time)):
        return type(previous).fromisoformat(text)
    if isinstance(previous, timedelta):
        raise ValueError('時間間隔の編集は次の段階で対応します')
    raise ValueError('この型の編集には対応していません')


def protection_reason(probe, sheet, cell):
    if probe.read_only:
        return 'ブックは読み取り専用です'
    if cell.data_type == 'f':
        return '数式セルは編集できません（再計算しません）'
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            return f'結合範囲 {merged} は編集できません'
    if cell.data_type == 'e':
        return '既存のエラー値は編集できません'
    if isinstance(cell.value, timedelta):
        return '時間間隔の編集は未対応です'
    return ''


def clipboard_rows(text):
    if len(text) > 1_000_000:
        raise ValueError('今回の貼り付けは100万文字までです')
    rows = list(csv.reader(StringIO(text, newline=''), delimiter='\t', strict=True))
    if not rows:
        return []
    if not rows[0] or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError('貼り付け範囲の列数をそろえてください')
    if len(rows) * len(rows[0]) > 50_000:
        raise ValueError('今回の貼り付けは5万セルまでです')
    return rows


def cell_state(cell):
    return cell.value, cell.data_type, copy(cell._style)


def restore_cell(cell, state):
    cell._value, cell.data_type, cell._style = state[0], state[1], copy(state[2])


class XlsxSession(AxisOperations):
    def __init__(self, path):
        self.probe, self.stamp = load_checked(path)
        try:
            if sum(sheet.max_row * sheet.max_column for sheet in self.probe.book) > 100_000:
                raise ValueError('現在のXLSX編集はブック全体10万セルまでです')
            check_area(self.probe.book)
        except Exception:
            self.probe.book.close()
            raise
        self.sheet_name = self.probe.book.active.title
        self.saved_values = self.values()
        self.undo_stack, self.redo_stack = [], []

    @property
    def sheet(self):
        return self.probe.book[self.sheet_name]

    def values(self):
        # Include structure and formats so deleting/recreating a sheet is tracked.
        return snapshot(self.probe.book)

    @property
    def dirty(self):
        return self.values() != self.saved_values

    def edit(self, row, column, text):
        return self.apply_rows(row, column, [[text]])

    def paste(self, row, column, text):
        return self.apply_rows(row, column, clipboard_rows(text))

    def apply_rows(self, row, column, rows):
        if self.probe.read_only:
            raise ProtectedCellError('ブックは読み取り専用です')
        if not rows:
            return False
        sheet = self.sheet
        if (row < 0 or column < 0 or row + len(rows) > min(extent(sheet)[0], VIEW_ROWS)
                or column + len(rows[0]) > min(extent(sheet)[1], VIEW_COLUMNS)):
            raise ValueError('貼り付け先は現在表示中の行列範囲内を指定してください')
        pending = []
        for r, values in enumerate(rows, row + 1):
            for c, text in enumerate(values, column + 1):
                cell = sheet.cell(r, c)
                reason = protection_reason(self.probe, sheet, cell)
                if reason:
                    raise ProtectedCellError(f'{cell.coordinate}: {reason}')
                pending.append((cell, cell_state(cell), parse_input(text, cell.value)))
        try:
            for cell, before, value in pending:
                self.probe.set_value(self.sheet_name, cell.coordinate, value)
        except Exception:
            for cell, before, value in pending:
                restore_cell(cell, before)
            raise
        changes = [(self.sheet_name, cell.coordinate, before, cell_state(cell))
                   for cell, before, value in pending if before != cell_state(cell)]
        if not changes:
            return False
        self.undo_stack.append((self.sheet_name, changes))
        del self.undo_stack[:-HISTORY_LIMIT]
        self.redo_stack.clear()
        return True

    def history(self, redo=False):
        if self.probe.read_only:
            raise ProtectedCellError('ブックは読み取り専用です')
        source, target = (self.redo_stack, self.undo_stack) if redo else (self.undo_stack, self.redo_stack)
        if not source:
            return False
        name, changes = source[-1]
        if isinstance(changes, dict) and changes.get('kind') == 'widths':
            restore_dimensions(self.probe.book[name], changes['after' if redo else 'before'])
            self.sheet_name = name
        elif isinstance(changes, dict) and changes.get('kind') == 'heights':
            xlsx_heights.restore_dimensions(self.probe.book[name], changes['after' if redo else 'before'])
            self.sheet_name = name
        elif isinstance(changes, dict):
            book = load_workbook(BytesIO(changes['after' if redo else 'before']), data_only=False)
            self.probe.book.close()
            self.probe.book = book
            self.sheet_name = changes['after_name' if redo else 'before_name']
        else:
            for sheet_name, address, before, after in changes:
                restore_cell(self.probe.book[sheet_name][address], after if redo else before)
            self.sheet_name = name
        source.pop()
        target.append((name, changes))
        return True

    def column_widths(self, digit_width=7):
        return display_widths(self.sheet, min(extent(self.sheet)[1], VIEW_COLUMNS), digit_width)

    def resize_columns(self, changes, digit_width=7):
        if self.probe.read_only:
            raise ProtectedCellError('読み取り専用のブックは列幅を変更できません')
        current = self.column_widths(digit_width)
        pending = []
        for index, pixels in changes.items():
            if type(index) is not int or not 0 <= index < len(current):
                raise ValueError('表示中の列を指定してください')
            width = pixels_to_width(pixels, digit_width)
            if round(pixels) != current[index]:
                pending.append((index + 1, width))
        if not pending:
            return False
        before = dimension_state(self.sheet)
        try:
            for column, width in pending:
                set_width(self.sheet, column, width)
        except Exception:
            restore_dimensions(self.sheet, before)
            raise
        self.undo_stack.append((self.sheet_name, {'kind': 'widths', 'before': before,
                                                 'after': dimension_state(self.sheet)}))
        del self.undo_stack[:-HISTORY_LIMIT]
        self.redo_stack.clear()
        return True

    def row_heights(self, scale=4/3, minimum=5):
        return xlsx_heights.display_heights(self.sheet, min(extent(self.sheet)[0], VIEW_ROWS), scale, minimum)

    def resize_rows(self, changes, scale=4/3, minimum=5):
        if self.probe.read_only:
            raise ProtectedCellError('読み取り専用のブックは行の高さを変更できません')
        current = self.row_heights(scale, minimum)
        pending = []
        for row, pixels in changes.items():
            if type(row) is not int or not 0 <= row < len(current):
                raise ValueError('表示中の行を指定してください')
            points = xlsx_heights.pixels_to_height(pixels, scale)
            if round(pixels) != current[row]:
                pending.append((row + 1, points))
        if not pending:
            return False
        before = xlsx_heights.dimension_state(self.sheet)
        try:
            for row, points in pending:
                xlsx_heights.set_height(self.sheet, row, points)
        except Exception:
            xlsx_heights.restore_dimensions(self.sheet, before)
            raise
        self.undo_stack.append((self.sheet_name, {'kind': 'heights', 'before': before,
            'after': xlsx_heights.dimension_state(self.sheet)}))
        del self.undo_stack[:-HISTORY_LIMIT]
        self.redo_stack.clear()
        return True

    def structure_reason(self, operation):
        if self.probe.read_only:
            return 'ブックは読み取り専用です'
        if operation in {'rename', 'delete'} and any(c.data_type == 'f' for s in self.probe.book for row in s for c in row):
            return '数式を含むブックでは参照保護のため削除・名前変更はできません'
        if operation == 'delete' and self.sheet.sheet_state == 'visible' and sum(
                s.sheet_state == 'visible' for s in self.probe.book) <= 1:
            return '最後の表示シートは削除できません'
        return ''

    def validate_sheet_name(self, name, *, current=None):
        if (not isinstance(name, str) or not name.strip() or len(name) > 31
                or re.search(r'[\\/\?\*\[\]:\x00-\x1f]', name)
                or name.startswith("'") or name.endswith("'") or name.casefold() == 'history'):
            raise ValueError('シート名は1～31文字で、禁止文字・先頭末尾の引用符・Historyは使用できません')
        if any(s.title.casefold() == name.casefold() and s.title != current for s in self.probe.book):
            raise ValueError('同じ名前のシートがあります（大文字小文字は区別しません）')

    def pack_book(self):
        stream = BytesIO()
        self.probe.book.save(stream)
        checked = load_workbook(BytesIO(stream.getvalue()), data_only=False)
        try:
            if snapshot(checked) != snapshot(self.probe.book):
                raise ValueError('シート操作の復元用データを検証できませんでした')
        finally:
            checked.close()
        return stream.getvalue()

    def sheet_operation(self, operation, name=None):
        if operation not in {'add', 'rename', 'delete'}:
            raise ValueError('未対応のシート操作です')
        reason = self.structure_reason(operation)
        if reason:
            raise ProtectedCellError(reason)
        if operation in {'add', 'rename'}:
            self.validate_sheet_name(name, current=self.sheet_name if operation == 'rename' else None)
        if operation == 'rename' and name == self.sheet_name:
            return False
        before, before_name = self.pack_book(), self.sheet_name
        try:
            if operation == 'add':
                self.probe.book.create_sheet(name)  # Append only; no reordering.
                self.sheet_name = name
            elif operation == 'rename':
                sheet = self.sheet
                if sheet.title.casefold() == name.casefold():
                    index = 1
                    existing = {s.title.casefold() for s in self.probe.book}
                    while f'_SheetDock_{index}'.casefold() in existing:
                        index += 1
                    sheet.title = f'_SheetDock_{index}'
                sheet.title = name
                self.sheet_name = name
            else:
                self.probe.book.remove(self.sheet)
                self.sheet_name = next(s.title for s in self.probe.book if s.sheet_state == 'visible')
            self.probe.book.active = self.probe.book.index(self.sheet)
            after = self.pack_book()
        except Exception:
            self.probe.book = load_workbook(BytesIO(before), data_only=False)
            self.sheet_name = before_name
            raise
        self.undo_stack.append((before_name, {'before': before, 'after': after,
                                             'before_name': before_name, 'after_name': self.sheet_name}))
        del self.undo_stack[:-HISTORY_LIMIT]
        self.redo_stack.clear()
        return True

    def find(self, query, *, all_sheets=True, include_formulas=False):
        if not query:
            raise ValueError('検索文字列を入力してください')
        result = []
        for sheet in self.probe.book if all_sheets else [self.sheet]:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'f' and not include_formulas:
                        continue
                    if query in display_value(cell.value):
                        result.append((sheet.title, cell.row - 1, cell.column - 1,
                                       display_value(cell.value)))
        return result

    def replace_match(self, match, query, replacement):
        if self.probe.read_only:
            raise ProtectedCellError('ブックは読み取り専用です')
        if not query or not isinstance(replacement, str):
            raise ValueError('検索文字列と置換後の文字列を指定してください')
        if match.sheet not in self.probe.book.sheetnames:
            raise ValueError('対象シートが変わりました。再検索してください')
        sheet = self.probe.book[match.sheet]
        if not (0 <= match.row < extent(sheet)[0] and 0 <= match.column < extent(sheet)[1]):
            raise ValueError('検索位置が変わりました。再検索してください')
        cell = sheet.cell(match.row + 1, match.column + 1)
        reason = protection_reason(self.probe, sheet, cell)
        if reason or not isinstance(cell.value, str):
            raise ProtectedCellError(reason or '1件置換は通常の文字列セルだけが対象です')
        if (cell.value != match.value or match.start < 0 or match.end != match.start + len(query)
                or cell.value[match.start:match.end] != query):
            raise ValueError('検索位置または値が変わりました。再検索してください')
        value = cell.value[:match.start] + replacement + cell.value[match.end:]
        if value == cell.value:
            return False
        before = cell_state(cell)
        try:
            self.probe.set_value(match.sheet, cell.coordinate, value)
        except Exception:
            restore_cell(cell, before)
            raise
        self.undo_stack.append((match.sheet, [(match.sheet, cell.coordinate, before, cell_state(cell))]))
        del self.undo_stack[:-HISTORY_LIMIT]
        self.redo_stack.clear()
        self.sheet_name = match.sheet
        return True

    def replace_all(self, query, replacement, *, all_sheets=True):
        if self.probe.read_only:
            raise ProtectedCellError('ブックは読み取り専用です')
        matches = self.find(query, all_sheets=all_sheets, include_formulas=True)
        pending, skipped, occurrences = [], 0, 0
        for name, row, column, value in matches:
            sheet = self.probe.book[name]
            cell = sheet.cell(row + 1, column + 1)
            if protection_reason(self.probe, sheet, cell) or not isinstance(cell.value, str):
                skipped += 1
                continue
            new = cell.value.replace(query, replacement)
            if new != cell.value:
                pending.append((name, cell, cell_state(cell), new))
                occurrences += cell.value.count(query)
        try:
            for name, cell, before, value in pending:
                self.probe.set_value(name, cell.coordinate, value)
        except Exception:
            for name, cell, before, value in pending:
                restore_cell(cell, before)
            raise
        if pending:
            changes = [(name, cell.coordinate, before, cell_state(cell))
                       for name, cell, before, value in pending]
            self.undo_stack.append((self.sheet_name, changes))
            del self.undo_stack[:-HISTORY_LIMIT]
            self.redo_stack.clear()
        return len(pending), occurrences, skipped

    def export_csv(self, path, *, formulas='reject'):
        if formulas not in {'reject', 'text'}:
            raise ValueError('数式の出力方法が不正です')
        sheet = self.sheet
        if any(c.data_type == 'f' and not isinstance(c.value, str) for row in sheet for c in row):
            raise ValueError('特殊数式のCSV出力は未対応です')
        if formulas == 'reject' and any(c.data_type == 'f' for row in sheet for c in row):
            raise ValueError('数式セルがあります。「数式文字列を出力」を明示的に選んでください')
        rows_count, columns_count = data_extent(sheet)
        rows = [[display_value(sheet.cell(r, c).value) for c in range(1, columns_count + 1)]
                for r in range(1, rows_count + 1)]
        target = Path(path).resolve()
        if target.suffix.lower() != '.csv':
            raise ValueError('出力先は.csvを指定してください')
        if target.exists():
            raise FileExistsError(target)
        temporary = None
        with save_lock(target):
            try:
                with tempfile.NamedTemporaryFile(dir=target.parent, prefix=f'.{target.name}.',
                                                 suffix='.csv', delete=False) as stream:
                    temporary = Path(stream.name)
                save_csv(CsvDocument(rows, CsvFormat('utf-8', True, '\r\n', ',')),
                         temporary, backup=False)
                os.link(temporary, target)
            finally:
                if temporary is not None:
                    temporary.unlink(missing_ok=True)

    def save_as(self, target):
        if Path(target).exists():
            raise FileExistsError(target)
        stamp = save_checked(self.probe, target, None, backup=False)
        self.probe.path = Path(target).resolve()
        self.stamp = stamp
        self.saved_values = self.values()

    def save(self, *, backup=True, expected=_DEFAULT_EXPECTED):
        stamp = save_checked(self.probe, self.probe.path,
                             self.stamp if expected is _DEFAULT_EXPECTED else expected,
                             backup=backup)
        self.stamp = stamp
        self.saved_values = self.values()


class XlsxPane(InlineEditing):
    def __init__(self, parent, *, owner=None, on_change=None, on_open=None,
                 on_save=None, on_save_as=None, backup_var=None, check_export=None,
                 on_backup_change=None, on_selection=None):
        root = parent
        self.root, self.session = owner or parent, None
        self._on_change = on_change
        self._on_selection = on_selection
        self._rendering, self._digit_width = False, 7
        self._open_handler, self._save_handler, self._save_as_handler = on_open, on_save, on_save_as
        self._check_export = check_export
        self.backup = backup_var if backup_var is not None else tk.BooleanVar(value=True)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(5, weight=1)
        toolbar = ttk.Frame(root)
        toolbar.grid(row=0, column=0, sticky='ew', padx=8, pady=4)
        ttk.Label(toolbar, text='シート').pack(side='left')
        ttk.Button(toolbar, text='読取専用の理由', command=self.show_reasons).pack(side='right')
        self.sheet_tabs = SheetTabs(toolbar, self.switch_sheet)
        self.sheet_tabs.pack(side='left', fill='x', expand=True, padx=8)
        structure = ttk.Frame(root)
        structure.grid(row=1, column=0, sticky='ew', padx=8, pady=3)
        sheets_row = ttk.Frame(structure)
        sheets_row.pack(fill='x')
        self.structure_buttons = {}
        for operation, label in [('add', 'シート追加'), ('rename', 'シート名変更'), ('delete', 'シート削除')]:
            button = ttk.Button(sheets_row, text=label, command=lambda op=operation: self.sheet_operation(op), state='disabled')
            button.pack(side='left', padx=3)
            self.structure_buttons[operation] = button
        ttk.Label(sheets_row, text='数式を含むブックの削除・名前変更は制限します。').pack(side='left')
        axis_row = ttk.Frame(structure)
        axis_row.pack(fill='x', pady=(3, 0))
        ttk.Label(axis_row, text='行列操作').pack(side='left')
        self.append_buttons = {}
        for axis, label in [('rows', '行追加'), ('columns', '列追加')]:
            button = ttk.Button(axis_row, text=label, state='disabled',
                                command=lambda target=axis: self.append_axis(target))
            button.pack(side='left', padx=3)
            self.append_buttons[axis] = button
        self.axis_buttons = {}
        for operation, label in [('insert-rows', '行挿入'), ('insert-columns', '列挿入'),
                                 ('delete-rows', '行削除'), ('delete-columns', '列削除')]:
            button = ttk.Button(axis_row, text=label, state='disabled',
                                command=lambda op=operation: self.axis_operation(op))
            button.pack(side='left', padx=3)
            self.axis_buttons[operation] = button
        self.height_button = ttk.Button(axis_row, text='行の高さ', command=self.set_row_height, state='disabled')
        self.height_button.pack(side='left', padx=3)
        self.width_button = ttk.Button(axis_row, text='列の幅', command=self.set_column_width, state='disabled')
        self.width_button.pack(side='left', padx=3)
        self.axis_status = tk.StringVar(value='行列操作の対象セルを選択してください')
        ttk.Label(structure, textvariable=self.axis_status, wraplength=950).pack(fill='x')
        tools = ttk.Frame(root)
        tools.grid(row=2, column=0, sticky='ew', padx=8, pady=3)
        search_row = ttk.Frame(tools)
        search_row.pack(fill='x')
        replace_row = ttk.Frame(tools)
        replace_row.pack(fill='x', pady=(3, 0))
        ttk.Label(search_row, text='検索').pack(side='left')
        self.query = tk.StringVar()
        self.query_entry = ttk.Entry(search_row, textvariable=self.query, width=16)
        self.query_entry.pack(side='left')
        self.all_sheets = tk.BooleanVar(value=True)
        ttk.Checkbutton(search_row, text='全シート', variable=self.all_sheets).pack(side='left')
        self.formula_search = tk.BooleanVar(value=False)
        ttk.Checkbutton(search_row, text='数式も検索', variable=self.formula_search).pack(side='left')
        ttk.Button(search_row, text='前へ', command=lambda: self.find_next(backward=True)).pack(side='left', padx=3)
        ttk.Button(search_row, text='次へ', command=self.find_next).pack(side='left')
        self.replacement = tk.StringVar()
        ttk.Label(replace_row, text='置換後').pack(side='left')
        ttk.Entry(replace_row, textvariable=self.replacement, width=16).pack(side='left', padx=4)
        self.replace_one_button = ttk.Button(replace_row, text='1件置換', command=self.replace_one, state='disabled')
        self.replace_one_button.pack(side='left', padx=3)
        self.replace_button = ttk.Button(replace_row, text='文字列セルを一括置換', command=self.replace_all, state='disabled')
        self.replace_button.pack(side='left')
        exporting = ttk.Frame(root)
        exporting.grid(row=3, column=0, sticky='ew', padx=8, pady=3)
        ttk.Label(exporting, text='CSV出力設定').pack(side='left')
        self.formula_export = ttk.Combobox(exporting, state='readonly', width=25,
                                         values=['数式があれば出力中止', '数式文字列を出力'])
        self.formula_export.current(0)
        self.formula_export.pack(side='left', padx=5)
        ttk.Label(exporting, text='UTF-8 BOM / CRLF / カンマ。上部の「CSV出力」で実行。').pack(side='left')
        self.search_tracker = XlsxSearch()
        self.status = tk.StringVar(value='XLSXファイルを開いてください')
        ttk.Label(root, textvariable=self.status, wraplength=950).grid(row=4, column=0, sticky='ew', padx=8)
        self.sheet_widget = Sheet(root, data=[['']], headers=['A'], cell_auto_resize_enabled=False)
        self.sheet_widget.grid(row=5, column=0, sticky='nsew', padx=8, pady=4)
        self.sheet_widget.enable_bindings(*VIEW_BINDINGS)
        self.sheet_widget.bind('<<SheetSelect>>', self.selection_changed)
        self.sheet_widget.extra_bindings('column_width_resize', self.column_resized)
        # tksheet emits resize events for every motion, before its release handler
        # clears currently_resizing_width. Commit once after that handler finishes.
        self._pending_column_widths = {}
        self.sheet_widget.CH.bind('<ButtonRelease-1>', self.finish_column_resize, add='+')
        self._pending_row_sizes = {}
        self.sheet_widget.extra_bindings('row_height_resize', self.row_resized)
        self.sheet_widget.RI.bind('<ButtonRelease-1>', self.finish_row_resize, add='+')
        for keys in ('<Control-v>', '<Control-V>'):
            self.sheet_widget.bind(keys, self.paste)
        self.detail = tk.StringVar(value='')
        ttk.Label(root, textvariable=self.detail, wraplength=950).grid(row=6, column=0, sticky='ew', padx=8)
        ttk.Label(root, text='F2・ダブルクリックで直接編集。Enter／Tabで確定、Escで取消。末尾追加は「行追加」「列追加」。\n'
                  '灰色＝保護セル。数式は文字列表示、結合範囲は分けて表示。先頭1000行×50列まで。',
                  wraplength=950).grid(row=7, column=0, sticky='ew', padx=8, pady=4)
        self.setup_inline()

    def confirm_discard(self):
        return self.finish_inline_edit() and (not self.session or not self.session.dirty or messagebox.askyesno(
            '未保存の変更', '編集内容を破棄して続けますか？', parent=self.root))

    @settled
    def open_file(self):
        if not self.confirm_discard():
            return
        path = filedialog.askopenfilename(parent=self.root, filetypes=[('Excel workbook', '*.xlsx')])
        if not path:
            return
        try:
            session = XlsxSession(path)
            self.session = session
            self.render()
        except Exception as exc:
            messagebox.showerror('読込失敗', str(exc), parent=self.root)

    def switch_sheet(self, name):
        if not self.finish_inline_edit():
            self.sheet_tabs.sync(self.session.probe.book.sheetnames, self.session.sheet_name)
            return
        if self.session and name in self.session.probe.book.sheetnames and name != self.session.sheet_name:
            self.session.sheet_name = name
            self.render()

    def render(self):
        self._rendering = True
        try:
            self._render()
        finally:
            self._rendering = False

    def _render(self):
        session = self.session
        sheet = session.sheet
        rows, columns = min(extent(sheet)[0], VIEW_ROWS), min(extent(sheet)[1], VIEW_COLUMNS)
        data = [[format_value(sheet.cell(r, c).value, sheet.cell(r, c).number_format)[0]
                 for c in range(1, columns + 1)]
                for r in range(1, rows + 1)]
        self.sheet_widget.set_sheet_data(data, delete_options=True)
        self.sheet_widget.headers([get_column_letter(c) for c in range(1, columns + 1)])
        self._digit_width = self.measure_digit_width()
        self.sheet_widget.set_options(min_column_width=20,
            max_column_width=width_to_pixels(MAX_EXCEL_WIDTH, self._digit_width))
        self.sheet_widget.set_column_widths(session.column_widths(self._digit_width))
        self._point_scale = float(self.root.winfo_fpixels('1p'))
        self._row_minimum = self.sheet_widget.MT.min_row_height
        self.sheet_widget.set_options(max_row_height=xlsx_heights.height_to_pixels(xlsx_heights.MAX_POINTS, self._point_scale))
        self.sheet_widget.set_row_heights(session.row_heights(self._point_scale, self._row_minimum))
        if session.probe.read_only:
            self.sheet_widget.disable_bindings('row_height_resize')
            self.sheet_widget.disable_bindings('column_width_resize')
        else:
            self.sheet_widget.enable_bindings('column_width_resize')
            self.sheet_widget.enable_bindings('row_height_resize')
        for r in range(rows):
            for c in range(columns):
                cell = sheet.cell(r + 1, c + 1)
                self.sheet_widget.align_cells(r, c, align=cell_alignment(cell), redraw=False)
                protected = protection_reason(session.probe, sheet, cell)
                bg, fg = None, None
                if cell.fill.patternType == 'solid' and cell.fill.fgColor.type == 'rgb':
                    bg = '#' + cell.fill.fgColor.rgb[-6:]
                if cell.font.color and cell.font.color.type == 'rgb':
                    fg = '#' + cell.font.color.rgb[-6:]
                if protected:
                    bg, fg = '#e5e7eb', '#374151'
                if bg or fg:
                    self.sheet_widget.highlight_cells(r, c, bg=bg or False, fg=fg or False, redraw=False)
        self.sheet_widget.refresh()
        mode = '読み取り専用' if session.probe.read_only else '対応範囲内'
        self.status.set(f'{sheet.title} / {mode} / 表示 {rows}行×{columns}列（表示範囲全体 {extent(sheet)[0]}行×{extent(sheet)[1]}列）')
        self.replace_button.configure(state='disabled' if session.probe.read_only else 'normal')
        self.replace_one_button.configure(state='disabled' if session.probe.read_only else 'normal')
        self.sheet_tabs.sync(session.probe.book.sheetnames, session.sheet_name)
        for operation, button in self.structure_buttons.items():
            button.configure(state='disabled' if session.structure_reason(operation) else 'normal')
        self._axis_block_reason = session.axis_reason()
        self.axis_status.set(self._axis_block_reason or '連続範囲は選択件数、1行・1列だけなら件数入力。挿入は先頭の手前、削除は確認後に実行。')
        self.update_axis_commands()
        self.detail.set('セルを選択してください')
        if getattr(self, '_on_change', None):
            self._on_change()

    def measure_digit_width(self):
        book = self.session.probe.book
        normal = next((s for s in book._named_styles if s.name == 'Normal'), book._named_styles[0])
        font = normal.font
        measured = tkfont.Font(root=self.root, family=font.name or 'Calibri',
                               size=round(font.sz or 11), weight='bold' if font.b else 'normal')
        return max(1, max(measured.measure(str(digit)) for digit in range(10)))

    def column_resized(self, event):
        if self._rendering or self.session is None:
            return
        changes = {column: value['new_size']
                   for column, value in event.get('resized', {}).get('columns', {}).items()}
        header = getattr(getattr(self, 'sheet_widget', None), 'CH', None)
        if header is not None and header.currently_resizing_width:
            self._pending_column_widths.update(changes)
            return
        self.apply_column_widths(changes)

    def finish_column_resize(self, event=None):
        changes, self._pending_column_widths = self._pending_column_widths, {}
        if changes:
            self.apply_column_widths(changes)

    def apply_column_widths(self, changes):
        if not self.finish_inline_edit():
            self.render()
            return
        try:
            if self.session.resize_columns(changes, self._digit_width) and self._on_change:
                self._on_change()
        except Exception as exc:
            self.render()
            messagebox.showerror('列幅変更失敗', str(exc), parent=self.root)

    def row_resized(self, event):
        if self._rendering or self.session is None:
            return
        accumulate_row_resize(self._pending_row_sizes, event)
        if not self.sheet_widget.RI.currently_resizing_height:
            self.finish_row_resize()

    def finish_row_resize(self, event=None):
        if not self.finish_inline_edit():
            self._pending_row_sizes = {}
            self.render()
            return
        if self._rendering or self.session is None:
            return
        pending, self._pending_row_sizes = self._pending_row_sizes, {}
        changes = {row: new for row, (old, new) in pending.items() if old != new}
        try:
            if changes and self.session.resize_rows(changes, self._point_scale, self._row_minimum):
                if self._on_change:
                    self._on_change()
        except Exception as exc:
            self.render()
            messagebox.showerror('行の高さ変更失敗', str(exc), parent=self.root)

    @settled
    def set_row_height(self):
        cell = self.selected_cell()
        if cell is None or self.session.probe.read_only:
            return
        try:
            pixels = simpledialog.askinteger('行の高さ', '選択行の高さ（ピクセル）', parent=self.root,
                initialvalue=self.session.row_heights(self._point_scale, self._row_minimum)[cell.row - 1],
                minvalue=self._row_minimum,
                maxvalue=xlsx_heights.height_to_pixels(xlsx_heights.MAX_POINTS, self._point_scale))
            if pixels is not None and self.session.resize_rows({cell.row - 1: pixels}, self._point_scale, self._row_minimum):
                self.render()
        except Exception as exc:
            messagebox.showerror('行の高さ変更失敗', str(exc), parent=self.root)

    @settled
    def set_column_width(self):
        cell = self.selected_cell()
        if cell is None or self.session.probe.read_only:
            return
        try:
            pixels = simpledialog.askinteger('列の幅', '選択列の幅（ピクセル）', parent=self.root,
                initialvalue=self.session.column_widths(self._digit_width)[cell.column - 1],
                minvalue=MIN_PIXELS, maxvalue=width_to_pixels(MAX_EXCEL_WIDTH, self._digit_width))
            if pixels is not None and self.session.resize_columns({cell.column - 1: pixels}, self._digit_width):
                self.render()
        except Exception as exc:
            messagebox.showerror('列の幅変更失敗', str(exc), parent=self.root)

    def selected_cell(self):
        selected = self.sheet_widget.get_currently_selected()
        if not self.session or not selected:
            return None
        if selected.row >= min(extent(self.session.sheet)[0], VIEW_ROWS) or selected.column >= min(extent(self.session.sheet)[1], VIEW_COLUMNS):
            return None
        return self.session.sheet.cell(selected.row + 1, selected.column + 1)

    def selection_changed(self, event=None):
        self.update_axis_commands()
        cell = self.selected_cell()
        if cell is None:
            if getattr(self, '_on_selection', None):
                self._on_selection()
            return
        reason = protection_reason(self.session.probe, self.session.sheet, cell)
        text, supported = format_value(cell.value, cell.number_format)
        extra = ''
        if not supported:
            extra = ' / 表示書式未対応（元の値を表示）'
        elif text != display_value(cell.value):
            extra = f' / 元の値={display_value(cell.value)[:80]}'
        self.detail.set(f'{cell.coordinate} / 型={cell.data_type} / {reason or "編集可能（既存の型を維持）"}{extra}')

        if getattr(self, '_on_selection', None):
            self._on_selection()

    def update_axis_commands(self):
        editable = self.session is not None and not self.session.probe.read_only and self.selected_cell() is not None
        for name in ('height_button', 'width_button'):
            button = getattr(self, name, None)
            if button is not None:
                button.configure(state='normal' if editable else 'disabled')
        for axis, button in getattr(self, 'append_buttons', {}).items():
            enabled = (self.session is not None and not self.session.probe.read_only
                       and not getattr(self, '_axis_block_reason', '未読込'))
            if enabled:
                rows, columns = extent(self.session.sheet)
                enabled = (rows < VIEW_ROWS if axis == 'rows' else columns < VIEW_COLUMNS)
            button.configure(state='normal' if enabled else 'disabled')
        if not getattr(self, 'axis_buttons', None):
            return
        enabled = (self.session is not None and not getattr(self, '_axis_block_reason', '未読込')
                   and self.selected_cell() is not None)
        for button in self.axis_buttons.values():
            button.configure(state='normal' if enabled else 'disabled')

    @settled
    def append_axis(self, axis):
        if self.session is None:
            return
        try:
            reason = self.session.axis_reason()
            if reason:
                messagebox.showinfo('行列追加の制限', reason, parent=self.root)
                return
            rows, columns = extent(self.session.sheet)
            maximum = min(100, VIEW_ROWS - rows if axis == 'rows' else VIEW_COLUMNS - columns)
            if maximum < 1:
                messagebox.showinfo('行列追加の制限', '表示範囲の上限に達しています', parent=self.root)
                return
            unit = '行' if axis == 'rows' else '列'
            count = simpledialog.askinteger(unit + '追加', '末尾へ追加する' + unit + '数',
                parent=self.root, initialvalue=1, minvalue=1, maxvalue=maximum)
            if count is None:
                return
            if self.session.append_axis(axis, count, self._digit_width):
                self.search_tracker = XlsxSearch()
                self.render()
                row, column = (rows, 0) if axis == 'rows' else (0, columns)
                self.sheet_widget.select_cell(row, column)
                self.sheet_widget.see(row, column)
                self.selection_changed()
        except Exception as exc:
            messagebox.showerror('行列追加失敗', str(exc), parent=self.root)

    @settled
    def axis_operation(self, action):
        if not self.session:
            return
        try:
            reason = self.session.axis_reason()
            if reason:
                messagebox.showinfo('行列操作の制限', reason, parent=self.root)
                return
            cell = self.selected_cell()
            if cell is None:
                messagebox.showinfo('選択', '行列操作の対象セルを選択してください', parent=self.root)
                return
            rows = action.endswith('rows')
            deleting = action.startswith('delete')
            size = extent(self.session.sheet)[0 if rows else 1]
            index, selected_count = selected_axis_range(self.sheet_widget.get_all_selection_boxes(), rows, size)
            maximum = min(100, size - index + 1, size - 1) if deleting else min(
                100, (VIEW_ROWS if rows else VIEW_COLUMNS) - size)
            if maximum < 1:
                messagebox.showinfo('行列操作の制限', '範囲上限、または残す行・列の制限により操作できません', parent=self.root)
                return
            unit = '行' if rows else '列'
            label = ('削除' if deleting else '挿入')
            if selected_count > 1:
                if selected_count > maximum:
                    messagebox.showinfo('行列操作の制限',
                        '選択範囲全体を操作できません。1回100件、範囲上限、残す行・列の制限を確認してください', parent=self.root)
                    return
                count = selected_count
            else:
                count = simpledialog.askinteger(unit + label, f'{index}{unit}目からの{label}件数',
                    minvalue=1, maxvalue=maximum, initialvalue=1, parent=self.root)
            if count is None:
                return
            if deleting and not messagebox.askyesno('行列削除の確認',
                    f'{self.session.sheet_name}の{index}{unit}目から{count}{unit}を削除します。\n'
                    '範囲内のセルと対応する行高・列幅を削除します。続けますか？', parent=self.root):
                return
            row, column = cell.row - 1, cell.column - 1
            if rows:
                row = index - 1
            else:
                column = index - 1
            if self.session.axis_operation(action, index, count):
                self.search_tracker = XlsxSearch()
                self.render()
                row = min(row, extent(self.session.sheet)[0] - 1)
                column = min(column, extent(self.session.sheet)[1] - 1)
                self.sheet_widget.select_cell(row, column)
                self.sheet_widget.see(row, column)
                self.selection_changed()
        except Exception as exc:
            messagebox.showerror('行列操作失敗', str(exc), parent=self.root)

    @settled
    def edit_cell(self):
        cell = self.selected_cell()
        if cell is None:
            return
        reason = protection_reason(self.session.probe, self.session.sheet, cell)
        if reason:
            messagebox.showinfo('編集制限', reason, parent=self.root)
            return
        text = simpledialog.askstring('セル編集', f'{cell.coordinate}（型={cell.data_type}）の新しい値',
                                      initialvalue=display_value(cell.value), parent=self.root)
        if text is None:
            return
        try:
            self.session.edit(cell.row - 1, cell.column - 1, text)
            self.render()
        except Exception as exc:
            messagebox.showerror('編集失敗', str(exc), parent=self.root)

    @settled
    def save_as(self):
        if getattr(self, '_save_as_handler', None):
            return self._save_as_handler()
        if not self.session or self.session.probe.read_only:
            return
        target = filedialog.asksaveasfilename(parent=self.root, defaultextension='.xlsx',
                                             filetypes=[('Excel workbook', '*.xlsx')])
        if not target:
            return
        try:
            self.session.save_as(target)
            self.render()
            messagebox.showinfo('別名保存完了', 'コピーを保存しました。元ファイルは変更していません。', parent=self.root)
        except Exception as exc:
            messagebox.showerror('保存失敗', str(exc), parent=self.root)

    @settled
    def paste(self, event=None):
        cell = self.selected_cell()
        if cell is None:
            return 'break'
        try:
            text = self.root.clipboard_get()
            self.session.paste(cell.row - 1, cell.column - 1, text)
            self.render()
        except Exception as exc:
            messagebox.showerror('貼り付け失敗', str(exc), parent=self.root)
        return 'break'

    @settled
    def save(self, event=None):
        if getattr(self, '_save_handler', None):
            self._save_handler()
            return 'break'
        if not self.session or self.session.probe.read_only:
            return 'break'
        try:
            self.session.save(backup=self.backup.get())
            self.render()
        except ExternalChangeError:
            self.resolve_external_change()
        except Exception as exc:
            messagebox.showerror('保存失敗', str(exc), parent=self.root)
        return 'break'

    @settled
    def resolve_external_change(self):
        try:
            observed = fingerprint(self.session.probe.path)
            answer = messagebox.askyesnocancel(
                '外部変更を検出',
                'ファイルが外部で変更または削除されています。\n'
                'はい：現在の編集内容で上書き\n'
                'いいえ：ファイルを再読込（編集内容は破棄）\n'
                'キャンセル：編集を続ける', parent=self.root)
            if answer is True:
                self.session.save(backup=self.backup.get(), expected=observed)
                self.render()
            elif answer is False and self.confirm_discard():
                candidate = XlsxSession(self.session.probe.path)
                self.session = candidate
                self.render()
        except Exception as exc:
            messagebox.showerror('保存・再読込失敗', str(exc), parent=self.root)

    @settled
    def history(self, redo=False):
        if self.session:
            try:
                if self.session.history(redo):
                    self.render()
            except Exception as exc:
                messagebox.showerror('履歴操作失敗', str(exc), parent=self.root)
        return 'break'

    def search_options(self):
        return dict(all_sheets=self.all_sheets.get(), include_formulas=self.formula_search.get())

    def show_match(self, match):
        if match is None:
            messagebox.showinfo('検索', '見つかりませんでした', parent=self.root)
            return
        self.session.sheet_name = match.sheet
        self.render()
        if match.row < VIEW_ROWS and match.column < VIEW_COLUMNS:
            self.sheet_widget.select_cell(match.row, match.column)
            self.sheet_widget.see(match.row, match.column)
            self.selection_changed()
        else:
            messagebox.showinfo('検索結果',
                f'{match.sheet}!{get_column_letter(match.column+1)}{match.row+1} は表示上限の範囲外です',
                parent=self.root)
        self.detail.set(self.detail.get() +
            f' / 検索 {self.search_tracker.index+1}/{len(self.search_tracker.matches)}'
            f' / セル内 {match.start+1}～{match.end}文字目')

    @settled
    def find_next(self, backward=False):
        if not self.session:
            return
        try:
            match = self.search_tracker.move(self.session, self.query.get(),
                backward=backward, **self.search_options())
            self.show_match(match)
        except Exception as exc:
            messagebox.showerror('検索失敗', str(exc), parent=self.root)

    @settled
    def replace_one(self):
        if not self.session or self.session.probe.read_only:
            return
        try:
            # Require a current, verified search result before replacing.
            query, replacement = self.query.get(), self.replacement.get()
            options = self.search_options()
            match = self.search_tracker.current(self.session, query, **options)
            selected = self.selected_cell()
            if (selected is None or self.session.sheet_name != match.sheet
                    or (selected.row - 1, selected.column - 1) != (match.row, match.column)):
                raise ValueError('検索結果とは別のセルが選択されています。もう一度検索してください')
            self.session.replace_match(match, query, replacement)
            self.render()
            try:
                next_match = self.search_tracker.after_replace(
                    self.session, query, match, replacement, **options)
            except ValueError as exc:
                self.search_tracker = XlsxSearch()
                messagebox.showwarning('1件置換完了',
                    f'置換は完了しましたが、次の検索を中止しました。\n{exc}', parent=self.root)
                return
            if next_match is not None:
                self.show_match(next_match)
            else:
                self.detail.set('1件置換完了 / 一致箇所は残っていません')
        except Exception as exc:
            messagebox.showerror('1件置換失敗', str(exc), parent=self.root)

    @settled
    def replace_all(self):
        if not self.session or self.session.probe.read_only:
            return
        if not messagebox.askyesno('一括置換', '文字列セルだけを置換します。保護セル・数値・日時・真偽値は除外します。実行しますか？', parent=self.root):
            return
        try:
            cells, occurrences, skipped = self.session.replace_all(
                self.query.get(), self.replacement.get(), all_sheets=self.all_sheets.get())
            self.render()
            messagebox.showinfo('置換完了', f'{cells}セル・{occurrences}箇所を置換 / {skipped}セルを除外', parent=self.root)
        except Exception as exc:
            messagebox.showerror('置換失敗', str(exc), parent=self.root)

    @settled
    def export_csv(self):
        if not self.session:
            return
        target = filedialog.asksaveasfilename(parent=self.root, defaultextension='.csv',
                                             filetypes=[('CSV', '*.csv')])
        if not target:
            return
        try:
            if getattr(self, '_check_export', None):
                self._check_export(target)
            mode = 'text' if self.formula_export.current() == 1 else 'reject'
            self.session.export_csv(target, formulas=mode)
            messagebox.showinfo('CSV出力完了', '選択シートの値をCSVへ出力しました。XLSXの保存状態は変わりません。\n数式文字列は、Excel等で開くと数式として解釈される場合があります。', parent=self.root)
        except Exception as exc:
            messagebox.showerror('CSV出力失敗', str(exc), parent=self.root)

    @settled
    def show_reasons(self):
        if self.session:
            reasons = self.session.probe.reasons
            messagebox.showinfo('読取専用の理由', '\n'.join(reasons) if reasons else '現在の検証範囲で未対応機能は検出されていません。', parent=self.root)

    @settled
    def sheet_operation(self, operation):
        if not self.session:
            return
        try:
            reason = self.session.structure_reason(operation)
            if reason:
                raise ProtectedCellError(reason)
            name = None
            if operation in {'add', 'rename'}:
                name = simpledialog.askstring('シート操作', 'シート名を入力してください', parent=self.root,
                    initialvalue=self.session.sheet_name if operation == 'rename' else '')
                if name is None:
                    return
            elif not messagebox.askyesno('シート削除', f'「{self.session.sheet_name}」を削除しますか？\nUndoで元に戻せます。', parent=self.root):
                return
            self.session.sheet_operation(operation, name)
            self.render()
        except Exception as exc:
            messagebox.showerror('シート操作失敗', str(exc), parent=self.root)

    @settled
    def close(self):
        if self.confirm_discard():
            self.root.destroy()

