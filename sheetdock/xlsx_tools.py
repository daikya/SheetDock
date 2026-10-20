"""Typed row sorting and view-only XLSX filters; workbook stays authoritative."""
from copy import copy
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
import math

from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from .xlsx_axis import HISTORY_LIMIT, VIEW_ROWS, VIEW_COLUMNS
from .xlsx_extent import extent, extents
from .xlsx_display import raw_value as display_value
from .xlsx_io import ProtectedCellError, inspect_package, snapshot


def typed_key(cell):
    if cell is None or cell.value is None or cell.value == '':
        return None, None
    if cell.data_type in {'f', 'e'}:
        raise ProtectedCellError('数式・エラー値は並べ替えキーにできません')
    value = cell.value
    if type(value) is str:
        return 'text', value
    if type(value) is bool:
        return 'boolean', value
    if type(value) in (int, float):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError('有限でない数値は比較できません')
        return 'number', Decimal(str(value))
    if type(value) in (datetime, time):
        if value.tzinfo is not None:
            raise ValueError('タイムゾーン付きの日時は対象外です')
        return type(value).__name__, value
    if type(value) in (date, timedelta):
        return type(value).__name__, value
    raise TypeError('この型の並べ替えには対応していません')


def relocated_cells(records, mapping):
    moved = []
    for record in records:
        row, column = coordinate_to_tuple(record[0])
        target = mapping.get(row, row)
        moved.append((target, column, (f'{get_column_letter(column)}{target}', *record[1:])))
    return tuple(record for _, _, record in sorted(moved, key=lambda item: item[:2]))


class SortFilterOperations:
    def check_column(self, column):
        rows, columns = extent(self.sheet)
        if rows > VIEW_ROWS or columns > VIEW_COLUMNS:
            raise ValueError('並べ替え・フィルターは対象シート全体1000行×50列以内です')
        if type(column) is not int or not 1 <= column <= columns:
            raise ValueError('列番号は既存範囲内の1始まりで指定してください')

    def sort_rows(self, column, *, header=True, descending=False):
        self.check_column(column)
        if type(header) is not bool or type(descending) is not bool:
            raise TypeError('見出し・降順はboolで指定してください')
        if self.sheet_name in self.filters:
            raise ValueError('並べ替え前にフィルターを解除してください')
        reason = self.axis_reason()
        if reason:
            raise ProtectedCellError(reason)
        if self.sheet.sheet_format.zeroHeight or self.sheet.sheet_format.defaultRowHeight == 0:
            raise ProtectedCellError('標準行が非表示のシートは対象外です')
        if any(cell.data_type == 'e' for cell in self.sheet._cells.values()):
            raise ProtectedCellError('対象シートにエラー値があるため並べ替えできません')
        count, _ = extent(self.sheet)
        start = 2 if header else 1
        filled, blanks, family = [], [], None
        for row in range(start, count + 1):
            kind, key = typed_key(self.sheet._cells.get((row, column)))
            if kind is None:
                blanks.append(row)
            else:
                if family is not None and kind != family:
                    raise TypeError('キー列の型が混在しています。今回の並べ替えは同じ型だけが対象です')
                family = kind
                filled.append((key, row))
        filled.sort(key=lambda item: item[0], reverse=descending)
        order = [row for _, row in filled] + blanks
        if order == list(range(start, count + 1)):
            return False
        mapping = {source: target for target, source in enumerate(order, start)}
        before = self.pack_book()
        baseline, old_extents = snapshot(self.probe.book), extents(self.probe.book)
        candidate = load_workbook(BytesIO(before), data_only=False)
        try:
            sheet = candidate[self.sheet_name]
            old_cells = dict(sheet._cells)
            sheet._cells.clear()
            for (row, column), cell in old_cells.items():
                moved = copy(cell)
                moved.parent, moved.row, moved.column = sheet, mapping.get(row, row), column
                moved._style = copy(cell._style)
                sheet._cells[(moved.row, column)] = moved
            dimensions = list(sheet.row_dimensions.items())
            sheet.row_dimensions.clear()
            for index, dimension in dimensions:
                moved = copy(dimension)
                moved.parent, moved.index = sheet, mapping.get(index, index)
                moved._style = copy(dimension._style)
                sheet.row_dimensions[moved.index] = moved
            actual = snapshot(candidate)
            active = candidate.sheetnames.index(self.sheet_name)
            for i, (old, new) in enumerate(zip(baseline, actual)):
                if i != active:
                    if old != new:
                        raise ValueError('他シートが変化したため中止しました')
                    continue
                if any(old[j] != new[j] for j in (0, 1, 2, 4)):
                    raise ValueError('列定義・シート情報・結合状態が変化しました')
                expected_dimensions = tuple(sorted((mapping.get(index, index), value) for index, value in old[3]))
                if new[3] != expected_dimensions or new[5] != relocated_cells(old[5], mapping):
                    raise ValueError('値・型・書式・行定義の移動結果が期待と一致しません')
            if extents(candidate) != old_extents:
                raise ValueError('並べ替えで論理的な表示範囲が変化しました')
            stream = BytesIO()
            candidate.save(stream)
            after = stream.getvalue()
            if inspect_package(after):
                raise ValueError('保存後に未対応の構造を検出しました')
            checked = load_workbook(BytesIO(after), data_only=False)
            try:
                if snapshot(checked) != actual or extents(checked) != old_extents:
                    raise ValueError('保存・再読込後の内容が一致しません')
            finally:
                checked.close()
            self.probe.book.close()
            self.probe.book, candidate = candidate, None
            self.undo_stack.append((self.sheet_name, {'before': before, 'after': after,
                'before_name': self.sheet_name, 'after_name': self.sheet_name}))
            del self.undo_stack[:-HISTORY_LIMIT]
            self.redo_stack.clear()
            return True
        finally:
            if candidate is not None:
                candidate.close()

    def set_filter(self, column, *, operator='contains', query='', header=True):
        self.check_column(column)
        if self.sheet.sheet_format.zeroHeight or self.sheet.sheet_format.defaultRowHeight == 0:
            raise ValueError('標準行が非表示のシートはフィルター表示の対象外です')
        if operator not in {'contains', 'equals', 'empty', 'nonempty'} or not isinstance(query, str) or type(header) is not bool:
            raise ValueError('フィルター条件が不正です')
        self.filters[self.sheet_name] = (column, operator, query, header)
        self.filter_headers[self.sheet_name] = header

    def clear_filter(self):
        self.filters.pop(self.sheet_name, None)

    def visible_rows(self, *, header=True):
        if type(header) is not bool:
            raise TypeError('見出し指定はboolです')
        condition = self.filters.get(self.sheet_name)
        if condition:
            column, operator, query, header = condition
        count, _ = extent(self.sheet)
        if count > VIEW_ROWS or extent(self.sheet)[1] > VIEW_COLUMNS:
            raise ValueError('フィルター表示は1000行×50列以内です')
        if self.sheet.sheet_format.zeroHeight or self.sheet.sheet_format.defaultRowHeight == 0:
            raise ValueError('標準行が非表示のシートはフィルター表示の対象外です')
        indices = [row for row in range(2 if header else 1, count + 1)
                   if not ((dimension := self.sheet.row_dimensions.get(row)) is not None
                           and (dimension.hidden or dimension.height == 0))]
        if not condition:
            return indices
        visible = []
        for row in indices:
            cell = self.sheet._cells.get((row, column))
            text = display_value(cell.value if cell is not None else None)
            if {'contains': lambda: query in text, 'equals': lambda: query == text,
                'empty': lambda: text == '', 'nonempty': lambda: text != ''}[operator]():
                visible.append(row)
        return visible
